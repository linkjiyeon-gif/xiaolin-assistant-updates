from __future__ import annotations

import re
from collections import deque
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Sequence

from .log_monitor import parse_log_time


@dataclass
class CrashRecord:
    index: int
    time: str
    package: str
    crash_type: str
    title: str
    summary: str
    stack: str
    raw_snippet: str
    remark: str = ""


JAVA_PATTERNS = ["FATAL EXCEPTION", "java.lang.", "Caused by:"]
NATIVE_PATTERNS = ["signal ", "backtrace", "tombstone", "native crash"]
ANR_PATTERNS = ["ANR in", "Application Not Responding", "Input dispatching timed out"]
UNITY_LUA_PATTERNS = ["LuaException", "stack traceback", "xlua", "tolua", "Unity"]
OOM_PATTERNS = ["OutOfMemoryError", "low memory", "failed to allocate"]
IOS_CRASH_PATTERNS = ["SIGABRT", "SIGSEGV", "Terminating app", "uncaught exception", "terminated due to", "Exception Type:"]


def _contains_any(line: str, patterns: Sequence[str]) -> bool:
    low = line.lower()
    return any(pattern.lower() in low for pattern in patterns)


def classify_line(line: str) -> str:
    if _contains_any(line, OOM_PATTERNS):
        return "OOM"
    if _contains_any(line, IOS_CRASH_PATTERNS):
        low = line.lower()
        if "sigabrt" in low:
            return "iOS SIGABRT"
        if "sigsegv" in low:
            return "iOS SIGSEGV"
        return "iOS Crash"
    if _contains_any(line, ANR_PATTERNS):
        return "ANR"
    if _contains_any(line, UNITY_LUA_PATTERNS):
        low = line.lower()
        if "lua" in low or "xlua" in low or "tolua" in low or "stack traceback" in low:
            return "Lua Error"
        if "unity" in low and any(token in low for token in ["exception", "error", "crash", "stacktrace"]):
            return "Unity Crash"
    if _contains_any(line, NATIVE_PATTERNS):
        return "Native Crash"
    if _contains_any(line, JAVA_PATTERNS):
        if "fatal exception" in line.lower():
            return "Fatal Exception"
        return "Java Crash"
    return ""


def extract_package(line: str, fallback: str = "") -> str:
    text = line or ""
    patterns = [
        r"ANR in ([A-Za-z0-9_.$]+)",
        r"Process:\s*([A-Za-z0-9_.$]+)",
        r"package=([A-Za-z0-9_.$]+)",
        r"pkg=([A-Za-z0-9_.$]+)",
        r"Bundle ID[:=]\s*([A-Za-z0-9_.$-]+)",
        r"bundleID[:=]\s*([A-Za-z0-9_.$-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return fallback


def make_record(index: int, line: str, snippet_lines: Sequence[str], fallback_package: str = "") -> CrashRecord:
    crash_type = classify_line(line) or "未知异常"
    time_text = parse_log_time(line)
    package = extract_package(line, fallback_package)
    stripped = line.strip()
    title = stripped[:180]

    # Prefer exception-like lines as summary.
    summary = stripped
    for item in reversed(list(snippet_lines)):
        item_text = item.strip()
        if any(token in item_text for token in ["Exception", "Error", "Caused by", "ANR in", "signal", "LuaException", "SIGABRT", "SIGSEGV", "Terminating app", "uncaught exception"]):
            summary = item_text
            break

    stack_candidates = []
    for item in snippet_lines[-60:]:
        item_text = item.rstrip("\n")
        if (
            " at " in item_text
            or item_text.strip().startswith("at ")
            or "Caused by:" in item_text
            or "backtrace" in item_text.lower()
            or "#00" in item_text
            or "stack traceback" in item_text.lower()
        ):
            stack_candidates.append(item_text)
    stack = "\n".join(stack_candidates[-30:]) or stripped

    return CrashRecord(
        index=index,
        time=time_text,
        package=package,
        crash_type=crash_type,
        title=title,
        summary=summary[:500],
        stack=stack,
        raw_snippet="\n".join(snippet_lines[-80:]),
    )


class CrashAnalyzer:
    def __init__(self, context_lines: int = 40):
        self.context_lines = max(10, int(context_lines))

    def analyze_lines(self, lines: Iterable[str], fallback_package: str = "") -> List[CrashRecord]:
        buffer = deque(maxlen=self.context_lines)
        records: List[CrashRecord] = []
        seen = set()
        last_type_line = {}
        for line_no, line in enumerate(lines, start=1):
            text = str(line or "").rstrip("\n")
            buffer.append(text)
            crash_type = classify_line(text)
            if not crash_type:
                continue
            signature = self._signature(crash_type, text)
            if signature in seen:
                continue
            # Avoid splitting one crash stack into many records.
            if line_no - last_type_line.get(crash_type, -9999) <= 30:
                continue
            seen.add(signature)
            last_type_line[crash_type] = line_no
            records.append(make_record(len(records) + 1, text, list(buffer), fallback_package=fallback_package))
        return records

    @staticmethod
    def _signature(crash_type: str, line: str) -> str:
        normalized = re.sub(r"\d+", "#", line.lower())[:180]
        return f"{crash_type}:{normalized}"


class CrashStreamDetector:
    def __init__(self, context_lines: int = 40, max_seen: int = 300):
        self.context_lines = max(10, int(context_lines))
        self.buffer = deque(maxlen=self.context_lines)
        self.records: List[CrashRecord] = []
        self.seen: deque[str] = deque(maxlen=max_seen)
        self.seen_set = set()
        self.line_no = 0
        self.last_type_line = {}

    def clear(self) -> None:
        self.buffer.clear()
        self.records.clear()
        self.seen.clear()
        self.seen_set.clear()
        self.line_no = 0
        self.last_type_line.clear()

    def feed_line(self, line: str, fallback_package: str = "") -> List[CrashRecord]:
        text = str(line or "").rstrip("\n")
        self.line_no += 1
        self.buffer.append(text)
        crash_type = classify_line(text)
        if not crash_type:
            return []
        signature = CrashAnalyzer._signature(crash_type, text)
        if signature in self.seen_set:
            return []
        if self.line_no - self.last_type_line.get(crash_type, -9999) <= 30:
            return []
        if len(self.seen) == self.seen.maxlen:
            old = self.seen.popleft()
            self.seen_set.discard(old)
        self.seen.append(signature)
        self.seen_set.add(signature)
        self.last_type_line[crash_type] = self.line_no
        record = make_record(len(self.records) + 1, text, list(self.buffer), fallback_package=fallback_package)
        self.records.append(record)
        return [record]

    def get_records(self) -> List[CrashRecord]:
        return list(self.records)


def format_crash_records(records: Sequence[CrashRecord], max_records: int | None = None) -> str:
    items = list(records)
    if max_records is not None:
        items = items[-max_records:]
    if not items:
        return "暂无崩溃提取结果。"
    rows = []
    for record in items:
        rows.append(
            f"[{record.index}] {record.time or '-'} | {record.crash_type} | {record.package or '-'} | {record.summary}"
        )
    return "\n".join(rows)


def crash_records_to_tsv(records: Sequence[CrashRecord]) -> str:
    rows = ["序号\t时间\t包名\t崩溃类型\t异常标题\t异常摘要\t关键堆栈\t原始日志片段\t备注"]
    for record in records:
        rows.append("\t".join([
            str(record.index),
            record.time,
            record.package,
            record.crash_type,
            record.title.replace("\t", " "),
            record.summary.replace("\t", " "),
            record.stack.replace("\t", " ").replace("\r", ""),
            record.raw_snippet.replace("\t", " ").replace("\r", ""),
            record.remark,
        ]))
    return "\n".join(rows)


def export_crash_records_to_excel(records: Sequence[CrashRecord], output_path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "崩溃日志提取"
    headers = ["序号", "时间", "包名", "崩溃类型", "异常标题", "异常摘要", "关键堆栈", "原始日志片段", "备注"]
    ws.append(headers)
    fill = PatternFill(fill_type="solid", fgColor="FADBD8")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in records:
        ws.append([
            record.index,
            record.time,
            record.package,
            record.crash_type,
            record.title,
            record.summary,
            record.stack,
            record.raw_snippet,
            record.remark,
        ])

    widths = [8, 20, 30, 18, 50, 60, 80, 100, 30]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
