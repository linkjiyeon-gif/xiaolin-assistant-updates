from __future__ import annotations

from collections import deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, List, Sequence


DEFAULT_KEYWORDS = [
    "Exception",
    "Error",
    "Fatal",
    "Crash",
    "ANR",
    "NullPointerException",
    "IndexOutOfBoundsException",
    "OutOfMemoryError",
    "LuaException",
    "Unity",
    "il2cpp",
    "stacktrace",
]


@dataclass
class KeywordHit:
    index: int
    time: str
    keyword: str
    line: str
    before_context: List[str] = field(default_factory=list)
    after_context: List[str] = field(default_factory=list)

    @property
    def context_text(self) -> str:
        parts: List[str] = []
        if self.before_context:
            parts.append("--- 前置上下文 ---")
            parts.extend(self.before_context)
        parts.append("--- 命中日志 ---")
        parts.append(self.line)
        if self.after_context:
            parts.append("--- 后置上下文 ---")
            parts.extend(self.after_context)
        return "\n".join(parts)


def parse_log_time(line: str) -> str:
    text = (line or "").strip()
    # logcat -v time: 05-15 14:20:30.123 ...
    if len(text) >= 18 and text[0:2].isdigit() and text[2:3] == "-" and text[6:8].isdigit():
        return text[:18]
    return ""


def normalize_keywords(keywords: Iterable[str]) -> List[str]:
    result: List[str] = []
    seen = set()
    for item in keywords:
        key = str(item or "").strip()
        if not key or key.startswith("#"):
            continue
        low = key.lower()
        if low not in seen:
            result.append(key)
            seen.add(low)
    return result


class LogMonitor:
    """Keyword monitor with before/after context support.

    This class has no ADB dependency and never writes source logs.
    """

    def __init__(self, keywords: Sequence[str] | None = None, context_lines: int = 20):
        self.keywords = normalize_keywords(keywords or DEFAULT_KEYWORDS)
        self.context_lines = max(0, int(context_lines))
        self.before_buffer = deque(maxlen=self.context_lines)
        self.pending_after: List[list] = []  # [hit, remaining]
        self.hits: List[KeywordHit] = []

    def configure(self, keywords: Sequence[str], context_lines: int = 20) -> None:
        self.keywords = normalize_keywords(keywords)
        self.context_lines = max(0, int(context_lines))
        self.before_buffer = deque(list(self.before_buffer)[-self.context_lines:], maxlen=self.context_lines)
        self.pending_after.clear()

    def clear(self) -> None:
        self.before_buffer.clear()
        self.pending_after.clear()
        self.hits.clear()

    def feed_line(self, line: str) -> List[KeywordHit]:
        text = str(line or "").rstrip("\n")
        new_hits: List[KeywordHit] = []

        # First, append this line as after-context for earlier hits.
        still_pending: List[list] = []
        for item in self.pending_after:
            hit, remaining = item
            if remaining > 0:
                hit.after_context.append(text)
                remaining -= 1
            if remaining > 0:
                still_pending.append([hit, remaining])
        self.pending_after = still_pending

        lower = text.lower()
        for keyword in self.keywords:
            if keyword.lower() in lower:
                hit = KeywordHit(
                    index=len(self.hits) + 1,
                    time=parse_log_time(text),
                    keyword=keyword,
                    line=text,
                    before_context=list(self.before_buffer),
                )
                self.hits.append(hit)
                new_hits.append(hit)
                if self.context_lines > 0:
                    self.pending_after.append([hit, self.context_lines])
                break

        self.before_buffer.append(text)
        return new_hits

    def analyze_lines(self, lines: Iterable[str]) -> List[KeywordHit]:
        self.clear()
        for line in lines:
            self.feed_line(line)
        return list(self.hits)

    def get_hits(self) -> List[KeywordHit]:
        return list(self.hits)


def format_keyword_hits(hits: Sequence[KeywordHit], max_records: int | None = None) -> str:
    items = list(hits)
    if max_records is not None:
        items = items[-max_records:]
    if not items:
        return "暂无关键字命中记录。"
    rows = []
    for hit in items:
        rows.append(f"[{hit.index}] {hit.time or '-'} | {hit.keyword} | {hit.line}")
    return "\n".join(rows)


def keyword_hits_to_tsv(hits: Sequence[KeywordHit]) -> str:
    rows = ["序号\t时间\t关键字\t日志内容\t上下文"]
    for hit in hits:
        rows.append("\t".join([
            str(hit.index),
            hit.time,
            hit.keyword,
            hit.line.replace("\t", " "),
            hit.context_text.replace("\t", " ").replace("\r", ""),
        ]))
    return "\n".join(rows)


def export_keyword_hits_to_txt(hits: Sequence[KeywordHit], output_path: str) -> None:
    path = Path(output_path)
    if path.suffix.lower() != ".txt":
        path = path.with_suffix(".txt")
    path.write_text(keyword_hits_to_tsv(hits), encoding="utf-8")


def export_keyword_hits_to_excel(hits: Sequence[KeywordHit], output_path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "关键字命中"
    headers = ["序号", "时间", "关键字", "日志内容", "上下文"]
    ws.append(headers)
    fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for hit in hits:
        ws.append([hit.index, hit.time, hit.keyword, hit.line, hit.context_text])
    widths = [8, 20, 24, 90, 100]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
