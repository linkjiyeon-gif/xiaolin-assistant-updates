from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from .file_readers import FileReadError, read_text_with_fallback


@dataclass
class LanguageMapping:
    name: str
    source_candidates: List[str]
    config_candidates: List[str]


@dataclass
class TranslationCompareOptions:
    source_header_row: int = 1
    source_data_start_row: int = 2
    source_key_column: str = ""
    source_sheet: str = ""
    config_header_row: int = 1
    config_real_header_row: int = 0
    config_data_start_row: int = 2
    config_key_column: str = ""
    config_sheet: str = ""
    language_mapping_text: str = ""
    enabled_languages: List[str] = None
    ignore_trim: bool = True
    ignore_all_spaces: bool = False
    ignore_newlines: bool = False
    ignore_case: bool = False
    ignore_full_half_width: bool = False
    ignore_punctuation: bool = False
    check_extra_ids: bool = True


@dataclass
class TranslationRecord:
    text_id: str
    row_number: int
    values: Dict[int, str]


@dataclass
class TranslationTable:
    path: str
    sheet_name: str
    headers: List[str]
    key_col_index: int
    records: Dict[str, TranslationRecord]
    duplicate_ids: Dict[str, List[int]]
    invalid_rows: List[int]


DEFAULT_SOURCE_KEY_CANDIDATES = ["Text ID", "字串编号", "ID", "Key", "tid", "string_id", "text_id", "key"]
DEFAULT_CONFIG_KEY_CANDIDATES = ["#tid", "tid", "id", "key", "string_id", "text_id", "Text ID"]

DEFAULT_LANGUAGE_MAPPING_TEXT = """中文简体: CN / cn / zh / zh_cn / 简体中文 / 中文 => cn / zh / zh_cn / CN
中文繁体: TW / tw / zh_tw / 繁体中文 / Traditional CN => tw / zh_tw / TW
英语: EN / en / English / 英文 => en / EN
日语: JP / ja / jp / Japanese / 日文 => ja / jp / JP
韩语: KR / ko / kr / Korean / 韩文 => ko / kr / KR
德语: DE / de / German / 德文 => de / DE
法语: FR / fr / French / 法文 => fr / FR
西班牙语: ES / es / Spanish / 西语 => es / ES
葡萄牙语: PT / pt / Portuguese / 葡语 => pt / PT
巴西葡语: PT_BR / pt_br / Brazilian Portuguese / 巴葡 => pt_br / pt / PT_BR
俄语: RU / ru / Russian / 俄语 => ru / RU
意大利语: IT / it / Italian / 意语 => it / IT
泰语: TH / th / Thai / 泰语 => th / TH
越南语: VN / vi / vn / Vietnamese / 越南语 => vi / vn / VN
印尼语: ID / id / Indonesian / 印尼语 => id / ID
土耳其语: TR / tr / Turkish / 土耳其语 => tr / TR
阿拉伯语: ARB / ar / Arabic / 阿语 => ar / arb / ARB
乌克兰语: UA / uk / Ukrainian / 乌克兰语 => uk / ua / UA"""

PUNCTUATION_PATTERN = re.compile(r"[\u3000-\u303F\uFF00-\uFF65\s!\"#$%&'()*+,\-./:;<=>?@\[\\\]^_`{|}~，。！？、；：‘’“”（）【】《》…—·]")


def safe_int(value, default: int) -> int:
    try:
        parsed = int(str(value).strip())
        return parsed if parsed > 0 else default
    except Exception:
        return default


def normalize_header(value: str) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.strip().lower()
    text = re.sub(r"[\s\-]+", "_", text)
    text = re.sub(r"[^\w#\u4e00-\u9fff]+", "", text)
    return text


def split_candidates(raw: str) -> List[str]:
    return [part.strip() for part in re.split(r"[/|,，;；]", raw or "") if part.strip()]


def parse_language_mappings(raw_text: str) -> List[LanguageMapping]:
    text = raw_text.strip() or DEFAULT_LANGUAGE_MAPPING_TEXT
    mappings: List[LanguageMapping] = []
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if ":" in line:
            name, rest = line.split(":", 1)
        elif "=" in line:
            name, rest = line.split("=", 1)
        else:
            continue
        if "=>" in rest:
            source_raw, config_raw = rest.split("=>", 1)
        else:
            source_raw = config_raw = rest
        source_candidates = split_candidates(source_raw)
        config_candidates = split_candidates(config_raw)
        name = name.strip()
        if name and source_candidates and config_candidates:
            mappings.append(LanguageMapping(name, source_candidates, config_candidates))
    return mappings


def read_table_rows(path: str, sheet_name: str = "") -> Tuple[List[List[str]], str]:
    file_path = Path(path)
    ext = file_path.suffix.lower()
    if ext == ".csv":
        raw = read_text_with_fallback(file_path)
        sample = raw[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(raw), dialect))
        return [["" if cell is None else str(cell) for cell in row] for row in rows], "CSV"

    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise FileReadError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc
        wb = load_workbook(file_path, data_only=True, read_only=True)
        try:
            if sheet_name and sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb[wb.sheetnames[0]]
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if cell is None else str(cell) for cell in row])
            return rows, ws.title
        finally:
            wb.close()

    raise FileReadError(f"翻译比对暂只支持 xlsx / csv：{file_path.name}")


def find_column(headers: Sequence[str], preferred: str, candidates: Sequence[str]) -> int:
    normalized_headers = [normalize_header(header) for header in headers]
    search_items: List[str] = []
    if preferred:
        search_items.append(preferred)
    search_items.extend(candidates)

    for candidate in search_items:
        normalized = normalize_header(candidate)
        if not normalized:
            continue
        for index, header in enumerate(normalized_headers):
            if header == normalized:
                return index

    for candidate in search_items:
        normalized = normalize_header(candidate)
        if not normalized:
            continue
        for index, header in enumerate(normalized_headers):
            if not header:
                continue
            if normalized and (normalized in header or header in normalized):
                return index
    return -1


def first_existing_column(headers: Sequence[str], candidates: Sequence[str]) -> int:
    return find_column(headers, "", candidates)


def parse_translation_table(
    path: str,
    sheet_name: str,
    header_row: int,
    data_start_row: int,
    key_column: str,
    key_candidates: Sequence[str],
) -> TranslationTable:
    rows, actual_sheet = read_table_rows(path, sheet_name)
    if not rows:
        raise FileReadError(f"文件为空：{path}")

    header_index = max(0, header_row - 1)
    data_start_index = max(header_index + 1, data_start_row - 1)
    if header_index >= len(rows):
        raise FileReadError(f"表头行超出文件范围：第 {header_row} 行")

    headers = [cell.strip() for cell in rows[header_index]]
    key_col = find_column(headers, key_column, key_candidates)
    if key_col < 0:
        expected = key_column or " / ".join(key_candidates[:6])
        raise FileReadError(f"表头识别失败：未找到主键列（期望：{expected}）")

    raw_records: Dict[str, TranslationRecord] = {}
    id_rows: Dict[str, List[int]] = {}
    invalid_rows: List[int] = []

    for row_index in range(data_start_index, len(rows)):
        row = rows[row_index]
        row_number = row_index + 1
        text_id = row[key_col].strip() if key_col < len(row) else ""
        if not text_id:
            if any(str(cell).strip() for cell in row):
                invalid_rows.append(row_number)
            continue
        values = {index: (row[index] if index < len(row) else "") for index in range(len(headers))}
        id_rows.setdefault(text_id, []).append(row_number)
        if text_id not in raw_records:
            raw_records[text_id] = TranslationRecord(text_id=text_id, row_number=row_number, values=values)

    duplicate_ids = {text_id: rows for text_id, rows in id_rows.items() if len(rows) > 1}
    return TranslationTable(
        path=path,
        sheet_name=actual_sheet,
        headers=headers,
        key_col_index=key_col,
        records=raw_records,
        duplicate_ids=duplicate_ids,
        invalid_rows=invalid_rows,
    )


def normalize_translation_text(value: str, options: TranslationCompareOptions) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if options.ignore_trim:
        text = text.strip()
    if options.ignore_full_half_width:
        text = unicodedata.normalize("NFKC", text)
    if options.ignore_newlines:
        text = text.replace("\n", "")
    if options.ignore_all_spaces:
        text = re.sub(r"\s+", "", text)
    if options.ignore_punctuation:
        text = PUNCTUATION_PATTERN.sub("", text)
    if options.ignore_case:
        text = text.lower()
    return text
