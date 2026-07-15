from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
import posixpath
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Sequence, Tuple

from .file_readers import FileReadError
from .translation_compare import (
    DEFAULT_LANGUAGE_MAPPING_TEXT,
    DEFAULT_SOURCE_KEY_CANDIDATES,
    find_column,
    normalize_header,
    parse_language_mappings,
    read_table_rows,
    safe_int,
)

from .localization_sheet_parser import (
    LANGUAGE_PRIORITY,
    MULTI_FILE_DICTIONARY_STRUCTURE,
    MULTI_TABLE_SHEETS_STRUCTURE,
    MULTI_SHEET_STRUCTURE,
    SINGLE_SHEET_STRUCTURE,
    UNKNOWN_STRUCTURE,
    NormalizedLocalizationDataset,
    detect_localization_structure,
    normalize_localization_mode,
    parse_multi_sheet_localization,
)


ProgressCallback = Callable[[Dict[str, object]], None]


def _localization_cell_text(cell_or_value) -> str:
    """Convert an Excel cell/value without leaking date serials or midnight text.

    This is intentionally scoped to localization checks so the generic file
    comparison reader keeps its historical behavior.
    """
    cell = cell_or_value if hasattr(cell_or_value, "value") else None
    value = cell.value if cell is not None else cell_or_value
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time.min:
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    return str(value)


def _worksheet_rows_as_text(ws, *, max_rows: int | None = None, strip: bool = False) -> List[List[str]]:
    rows: List[List[str]] = []
    kwargs = {"min_row": 1}
    if max_rows is not None:
        kwargs["max_row"] = max_rows
    for row in ws.iter_rows(**kwargs):
        values = [_localization_cell_text(cell) for cell in row]
        rows.append([value.strip() for value in values] if strip else values)
    return rows


def _emit_progress(
    progress_callback: ProgressCallback | None,
    stage: str,
    message: str,
    *,
    current: int = 0,
    total: int = 0,
    ratio: float | None = None,
    sheet_name: str = "",
    detail: str = "",
) -> None:
    if progress_callback is None:
        return
    payload: Dict[str, object] = {
        "stage": stage,
        "message": message,
        "current": current,
        "total": total,
        "sheet_name": sheet_name,
        "detail": detail,
    }
    if ratio is not None:
        payload["ratio"] = min(max(float(ratio), 0.0), 1.0)
    try:
        progress_callback(payload)
    except Exception:
        # 进度回调只用于 UI 反馈，不能影响实际检查结果。
        pass


@dataclass
class LocalizationColumn:
    name: str
    header: str
    index: int
    display_name: str = ""
    confidence: str = "高"
    excluded_reason: str = ""
    code: str = ""
    raw_header: str = ""


@dataclass
class LocalizationRow:
    text_id: str
    row_number: int
    values: Dict[int, str]


@dataclass
class LocalizationTable:
    path: str
    sheet_name: str
    headers: List[str]
    key_col_index: int
    records: Dict[str, LocalizationRow]
    duplicate_ids: Dict[str, List[int]]
    invalid_rows: List[int]
    display_headers: List[str] = field(default_factory=list)
    type_headers: List[str] = field(default_factory=list)
    header_row: int = 1
    data_start_row: int = 2
    structure_type: str = "单行表头"
    excluded_columns: Dict[str, str] = field(default_factory=dict)
    diagnostics: List[str] = field(default_factory=list)
    row_records: List[LocalizationRow] = field(default_factory=list)


@dataclass
class LocalizationCheckOptions:
    header_row: int = 1
    data_start_row: int = 2
    key_column: str = ""
    sheet_name: str = ""
    source_language_column: str = ""
    enabled_language_columns: List[str] = field(default_factory=list)
    ignore_ids_text: str = ""
    allowed_chinese_languages_text: str = "中文简体, 中文繁体, CN, cn, zh, zh_cn, TW, tw, zh_tw, 中文, 简体中文, 繁体中文, 繁體中文, 日本語, 日文, 日语, 韓国語, 한국어"
    symbols_text: str = "% , : , ： , + , - , / , \\ , () , [] , {} , \\n"
    max_length: int = 0
    length_ratio: float = 0.0
    check_empty: bool = True
    check_chinese: bool = True
    check_placeholders: bool = True
    check_tags: bool = True
    check_symbols: bool = True
    check_length: bool = False
    check_duplicate_id: bool = True
    check_invalid_id: bool = True
    ui_limit: int = 3000
    table_structure: str = "自动识别"  # 自动识别 / 单 Sheet 多语言列 / 多 Sheet 语言页
    expected_languages_text: str = "EN,DE,FR,RU,ES,JP,ID,KR,TH,IT,PT,TR,VN,ARB"
    check_missing_id: bool = True
    check_source_consistency: bool = True
    check_numbers: bool = True
    check_unfinished: bool = True
    check_spelling: bool = True
    check_terms: bool = True
    check_duplicate_translation: bool = True
    aggregate_empty_language_columns: bool = True
    precomputed_multi_table_sheets: List[Tuple[str, int, str]] = field(default_factory=list)
    empty_requires_source_text: bool = False


@dataclass
class LocalizationIssue:
    index: int
    issue_type: str
    item_id: str
    language: str
    sheet_name: str
    row_number: str
    source_text: str
    current_text: str
    remark: str
    suggestion: str
    issue_level: str = "普通问题"
    rule_name: str = ""
    source_language: str = ""
    is_false_positive: bool = False  # 兼容旧 UI 变量名；导出报告中不再显示为“疑似误报”。
    requires_manual_confirm: bool = False
    count_in_error_stats: bool = True
    rule_confidence: str = "中"
    is_degraded: bool = False
    module: str = ""
    source_file: str = ""
    target_file: str = ""
    result_category: str = "明确问题"
    affected_languages: str = ""
    coverage_total: int = 0
    coverage_filled: int = 0
    coverage_blank: int = 0
    coverage_status: str = ""


@dataclass
class DictionaryLanguageFile:
    code: str
    name: str
    path: str
    file_name: str
    sheet_name: str = ""


@dataclass
class DictionaryRecord:
    text_id: str
    module: str
    text: str
    row_number: int


@dataclass
class DictionaryFileData:
    info: DictionaryLanguageFile
    records: Dict[str, DictionaryRecord]
    duplicate_ids: Dict[str, List[int]]
    invalid_rows: List[int]
    diagnostics: List[str] = field(default_factory=list)


# ===== 表结构识别 / 语言列识别底座 =====

LANGUAGE_META: Dict[str, Tuple[str, List[str]]] = {
    "CN": ("中文简体", ["cn", "zh", "zh_cn", "simplified_cn", "source_cn", "简中", "简体", "简体中文", "中文", "中文简体", "原文", "cn_source_text", "source_text"]),
    "TW": ("中文繁体", ["tw", "zh_tw", "traditional_cn", "traditional", "繁体", "繁體", "繁体中文", "繁體中文", "繁體文本內容", "繁体中文翻译", "traditional_cn_translation"]),
    "EN": ("英语", ["en", "english", "english_text", "en_translation", "英文", "英语", "英語"]),
    "DE": ("德语", ["de", "german", "de_translation", "德语", "德語", "德語文本內容"]),
    "FR": ("法语", ["fr", "french", "fr_translation", "法语", "法語", "法語文本內容"]),
    "RU": ("俄语", ["ru", "russian", "ru_translation", "俄语", "俄語", "俄羅斯文本內容", "俄罗斯文本内容"]),
    "ES": ("西班牙语", ["es", "spanish", "es_translation", "西语", "西班牙语", "西班牙語"]),
    "JP": ("日语", ["jp", "ja", "japanese", "jp_translation", "ja_translation", "日语", "日語", "日文"]),
    "KR": ("韩语", ["kr", "ko", "korean", "kr_translation", "ko_translation", "韩语", "韓語", "韩文", "韓文"]),
    "TH": ("泰语", ["th", "thai", "th_translation", "泰语", "泰語"]),
    "IT": ("意大利语", ["it", "italian", "it_translation", "意大利语", "意大利語", "意语", "意語"]),
    "PT": ("葡萄牙语", ["pt", "portuguese", "pt_translation", "葡语", "葡語", "葡萄牙语", "葡萄牙語"]),
    "PT_BR": ("巴西葡语", ["pt_br", "ptbr", "brazilian_portuguese", "pt_br_translation", "巴葡", "巴西葡语", "巴西葡語"]),
    "TR": ("土耳其语", ["tr", "turkish", "tr_translation", "土耳其语", "土耳其語"]),
    "VN": ("越南语", ["vn", "vi", "vietnamese", "vn_translation", "vi_translation", "越南语", "越南語"]),
    "UA": ("乌克兰语", ["ua", "uk", "ukrainian", "ua_translation", "uk_translation", "乌克兰语", "烏克蘭語"]),
    "ARB": ("阿拉伯语", ["arb", "ar", "arabic", "arb_translation", "ar_translation", "阿语", "阿語", "阿拉伯语", "阿拉伯語"]),
    "ID": ("印尼语", ["id", "indonesian", "id_translation", "印尼语", "印尼語", "印尼文本內容", "印尼文本内容"]),
    "HIN": ("印地语", ["hin", "hi", "hindi", "hin_translation", "hi_translation", "印地语", "印地語", "印度语", "印度語"]),
    "UR": ("乌尔都语", ["ur", "urdu", "ur_translation", "乌尔都语", "烏爾都語", "乌尔都", "烏爾都"]),
    "FA": ("波斯语", ["fa", "farsi", "persian", "fa_translation", "farsi_translation", "persian_translation", "波斯语", "波斯語"]),
    "MS": ("马来语", ["ms", "malay", "ms_translation", "malay_translation", "马来语", "馬來語", "马来西亚语", "馬來西亞語"]),
}

PROGRAM_LANGUAGE_CODES = {alias for _code, (_name, aliases) in LANGUAGE_META.items() for alias in aliases if re.fullmatch(r"[a-z_]{2,5}", alias)}
KEY_FIELD_ALIASES = ["#tid", "tid", "text_id", "string_id", "key", "#id", "id", "文本id", "文本编号", "文本編號", "字串id", "字串编号", "字串編號"]
TYPE_TOKENS = {"uint", "int", "float", "double", "string", "str", "bool", "boolean", "long", "text"}
AUXILIARY_KEYWORDS = {
    "status", "state", "状态", "狀態", "翻译状态", "翻譯狀態", "勿改", "date", "日期", "modified", "last_modified",
    "batch", "批次", "round", "trans_round", "change", "change_log", "状态变化", "狀態變化", "备注", "備註",
    "comment", "note", "notes", "tdesc", "desc", "文本注释", "文本註釋", "注释", "註釋", "负责人", "負責人",
    "owner", "in_charge", "incharge", "in-charge", "char_limit", "字符限制", "字数限制", "字數限制", "length_limit",
    "id_check", "check_id", "检查_id", "檢查_id", "检查id", "檢查id", "检查 ID", "ID Check",
}

LANGUAGE_ORDER = ["CN", "EN", "DE", "FR", "ID", "RU", "TW", "ES", "IT", "JP", "KR", "PT_BR", "PT", "TH", "TR", "VN", "UA", "ARB", "HIN", "UR", "FA", "MS"]

DICTIONARY_FILE_LANGUAGE_MAP = {
    "chinesesimplified": "CN",
    "simplifiedchinese": "CN",
    "zhcn": "CN",
    "cn": "CN",
    "chinesetraditional": "TW",
    "traditionalchinese": "TW",
    "zhtw": "TW",
    "tw": "TW",
    "english": "EN",
    "en": "EN",
    "french": "FR",
    "fr": "FR",
    "german": "DE",
    "de": "DE",
    "indonesian": "ID",
    "id": "ID",
    "italian": "IT",
    "it": "IT",
    "japanese": "JP",
    "jp": "JP",
    "ja": "JP",
    "korean": "KR",
    "kr": "KR",
    "ko": "KR",
    "portuguese": "PT",
    "pt": "PT",
    "russian": "RU",
    "ru": "RU",
    "spanish": "ES",
    "es": "ES",
    "turkish": "TR",
    "tr": "TR",
}


def _split_config_values(raw: str) -> List[str]:
    return [part.strip() for part in re.split(r"[,，;；|\n]+", raw or "") if part.strip()]


def _format_identifier(value) -> str:
    """Format table primary keys as stable text.

    Excel sometimes stores ID columns as numeric cells. When those cells are
    read as float-like strings (for example ``103659.0``), they should still be
    displayed as the integer identifier ``103659`` in reports. Only integer
    decimal forms are normalized; real decimal IDs such as ``1.5`` are kept.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip()
    if isinstance(value, Decimal):
        return str(int(value)) if value == value.to_integral_value() else str(value).strip()
    text = str(value).strip()
    if re.fullmatch(r"[+-]?\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _language_name_for_code(code: str) -> str:
    normalized = str(code or "").upper()
    return LANGUAGE_META.get(normalized, (normalized, []))[0] or normalized


def _dictionary_language_code_from_name(name: str) -> str:
    stem = Path(str(name or "")).stem
    if stem.lower().startswith("dictionary_"):
        suffix = stem[len("dictionary_"):]
    else:
        suffix = stem
    key = re.sub(r"[^A-Za-z0-9]+", "", suffix).lower()
    return DICTIONARY_FILE_LANGUAGE_MAP.get(key, key.upper() if key else "")


def _dictionary_root(path: str) -> Path:
    file_path = Path(path)
    if file_path.is_dir():
        return file_path
    return file_path.parent


def detect_dictionary_language_files(path: str) -> Tuple[List[DictionaryLanguageFile], List[LocalizationColumn]]:
    """Detect dictionary_*.xlsx language files under a folder.

    This mode is for projects where every language is stored in a separate
    workbook, for example dictionary_ChineseSimplified.xlsx as CN and
    dictionary_English.xlsx as EN.
    """
    root = _dictionary_root(path)
    if not root.exists():
        return [], []
    files: List[DictionaryLanguageFile] = []
    seen_codes: set[str] = set()
    for file_path in sorted(root.glob("dictionary_*.xlsx")):
        if file_path.name.startswith("~$"):
            continue
        code = _dictionary_language_code_from_name(file_path.name)
        if not code:
            continue
        # If duplicate language files exist, keep the first one but still avoid
        # ambiguous UI columns.
        if code in seen_codes:
            continue
        seen_codes.add(code)
        files.append(DictionaryLanguageFile(
            code=code,
            name=_language_name_for_code(code),
            path=str(file_path),
            file_name=file_path.name,
        ))
    files.sort(key=lambda info: LANGUAGE_ORDER.index(info.code) if info.code in LANGUAGE_ORDER else 999)
    columns = [
        LocalizationColumn(
            name=info.name,
            header=info.code,
            index=index,
            display_name=info.file_name,
            code=info.code,
            raw_header=info.file_name,
        )
        for index, info in enumerate(files)
    ]
    return files, columns


def _find_dictionary_columns(sample_rows: Sequence[Sequence[str]]) -> Tuple[int, int, int, int, int]:
    header_index = -1
    id_col = -1
    content_col = -1
    module_col = -1
    for row_index, row in enumerate(sample_rows[:12]):
        normalized = [_normalize_column_key(str(cell or "")) for cell in row]
        for index, value in enumerate(normalized):
            if id_col < 0 and value in {"id", "text_id", "#tid", "tid", "key"}:
                id_col = index
            if content_col < 0 and value in {"contents", "content", "text", "translation", "target_text"}:
                content_col = index
        if id_col >= 0 and content_col >= 0:
            header_index = row_index
            break
        id_col = -1
        content_col = -1
    if header_index < 0:
        raise FileReadError("字典表表头识别失败：未找到 ID / Contents 列")
    for row in sample_rows[:12]:
        for index, cell in enumerate(row):
            value = str(cell or "").strip()
            norm = _normalize_column_key(value)
            if value == "所属模块" or norm in {"module", "module_name", "所属模块"}:
                module_col = index
                break
        if module_col >= 0:
            break
    if module_col < 0 and content_col - id_col >= 2:
        module_col = id_col + 1
    data_start_index = header_index + 1
    return header_index, data_start_index, id_col, module_col if module_col >= 0 else -1, content_col


def _xlsx_first_sheet_info(zf: zipfile.ZipFile) -> Tuple[str, str]:
    workbook = ET.fromstring(zf.read("xl/workbook.xml"))
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    sheet = workbook.find(".//main:sheets/main:sheet", ns)
    if sheet is None:
        raise FileReadError("字典表读取失败：工作簿中未找到 Sheet")
    sheet_name = sheet.attrib.get("name", "") or "Sheet1"
    relationship_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "")
    if not relationship_id:
        raise FileReadError("字典表读取失败：Sheet 关系 ID 缺失")

    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target = ""
    for rel in rels:
        if rel.attrib.get("Id") == relationship_id:
            target = rel.attrib.get("Target", "")
            break
    if not target:
        raise FileReadError("字典表读取失败：未找到 Sheet XML 路径")
    target = target.replace("\\", "/")
    if target.startswith("/"):
        sheet_path = target.lstrip("/")
    elif target.startswith("xl/"):
        sheet_path = target
    else:
        sheet_path = posixpath.normpath(posixpath.join("xl", target))
    return sheet_name, sheet_path


def _xlsx_read_shared_strings(zf: zipfile.ZipFile) -> List[str]:
    try:
        handle = zf.open("xl/sharedStrings.xml")
    except KeyError:
        return []
    strings: List[str] = []
    with handle:
        for _event, elem in ET.iterparse(handle, events=("end",)):
            if not elem.tag.endswith("}si"):
                continue
            parts: List[str] = []
            for child in elem.iter():
                if child.tag.endswith("}t") and child.text:
                    parts.append(child.text)
            strings.append("".join(parts))
            elem.clear()
    return strings


def _xlsx_cell_col_index(ref: str, fallback: int) -> int:
    match = re.match(r"([A-Za-z]+)", ref or "")
    if not match:
        return fallback
    value = 0
    for char in match.group(1).upper():
        value = value * 26 + (ord(char) - ord("A") + 1)
    return max(value - 1, 0)


def _xlsx_cell_text(cell: ET.Element, shared_strings: Sequence[str]) -> str:
    cell_type = cell.attrib.get("t", "")
    if cell_type == "inlineStr":
        parts: List[str] = []
        for child in cell.iter():
            if child.tag.endswith("}t") and child.text:
                parts.append(child.text)
        return "".join(parts)
    value = ""
    for child in cell:
        if child.tag.endswith("}v"):
            value = child.text or ""
            break
    if cell_type == "s":
        try:
            index = int(value)
            return shared_strings[index] if 0 <= index < len(shared_strings) else ""
        except Exception:
            return ""
    if cell_type == "b":
        return "TRUE" if value == "1" else "FALSE"
    return value


def _iter_xlsx_rows(zf: zipfile.ZipFile, sheet_path: str, shared_strings: Sequence[str]) -> Iterable[Tuple[int, List[str]]]:
    with zf.open(sheet_path) as handle:
        for _event, elem in ET.iterparse(handle, events=("end",)):
            if not elem.tag.endswith("}row"):
                continue
            try:
                row_number = int(elem.attrib.get("r", "0") or 0)
            except Exception:
                row_number = 0
            values: List[str] = []
            fallback_index = 0
            for cell in elem:
                if not cell.tag.endswith("}c"):
                    continue
                col_index = _xlsx_cell_col_index(cell.attrib.get("r", ""), fallback_index)
                fallback_index = max(fallback_index + 1, col_index + 1)
                if col_index >= len(values):
                    values.extend([""] * (col_index - len(values) + 1))
                values[col_index] = _xlsx_cell_text(cell, shared_strings)
            if values:
                yield row_number or 1, values
            elem.clear()


def _build_dictionary_data_from_rows(
    info: DictionaryLanguageFile,
    sample_rows: Sequence[Sequence[str]],
    rows: Iterable[Tuple[int, Sequence[str]]],
    parser_name: str,
) -> DictionaryFileData:
    _header_index, _data_start_index, id_col, module_col, content_col = _find_dictionary_columns(sample_rows)

    records: Dict[str, DictionaryRecord] = {}
    id_rows: Dict[str, List[int]] = defaultdict(list)
    invalid_rows: List[int] = []
    diagnostics: List[str] = [
        f"{info.file_name}：Sheet={info.sheet_name or info.file_name}；ID列={id_col + 1}；模块列={module_col + 1 if module_col >= 0 else '未识别'}；文本列={content_col + 1}；读取器={parser_name}"
    ]

    for row_number, row in rows:
        raw_id = row[id_col] if id_col < len(row) else ""
        text_id = _format_identifier(raw_id)
        raw_id_norm = _normalize_column_key(text_id)
        # Skip type/comment/config rows such as string / #唯一功能... / 0.
        if not text_id or raw_id_norm in {"id", "string", "0"} or text_id.startswith("#"):
            if row_number > 4 and any(str(cell or "").strip() for cell in row):
                invalid_rows.append(row_number)
            continue
        module = str(row[module_col]).strip() if module_col >= 0 and module_col < len(row) and row[module_col] is not None else ""
        text = str(row[content_col]).strip() if content_col < len(row) and row[content_col] is not None else ""
        id_rows[text_id].append(row_number)
        if text_id not in records:
            records[text_id] = DictionaryRecord(text_id=text_id, module=module, text=text, row_number=row_number)
    duplicate_ids = {text_id: rows_ for text_id, rows_ in id_rows.items() if len(rows_) > 1}
    return DictionaryFileData(info=info, records=records, duplicate_ids=duplicate_ids, invalid_rows=invalid_rows, diagnostics=diagnostics)


def _read_dictionary_file_fast(info: DictionaryLanguageFile) -> DictionaryFileData:
    with zipfile.ZipFile(Path(info.path)) as zf:
        sheet_name, sheet_path = _xlsx_first_sheet_info(zf)
        info.sheet_name = sheet_name
        shared_strings = _xlsx_read_shared_strings(zf)
        sample_rows: List[List[str]] = []
        for _row_number, row in _iter_xlsx_rows(zf, sheet_path, shared_strings):
            sample_rows.append(["" if cell is None else str(cell) for cell in row])
            if len(sample_rows) >= 12:
                break
        if not sample_rows:
            raise FileReadError("字典表读取失败：Sheet 中没有可读取的数据")
        return _build_dictionary_data_from_rows(
            info,
            sample_rows,
            _iter_xlsx_rows(zf, sheet_path, shared_strings),
            "fast-xlsx",
        )


def _read_dictionary_file_openpyxl(info: DictionaryLanguageFile, fallback_reason: str = "") -> DictionaryFileData:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FileReadError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    wb = load_workbook(Path(info.path), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        info.sheet_name = ws.title
        sample_rows = []
        for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
            sample_rows.append(["" if cell is None else str(cell) for cell in row])
        data = _build_dictionary_data_from_rows(
            info,
            sample_rows,
            ((row_number, row) for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1)),
            "openpyxl",
        )
        if fallback_reason:
            data.diagnostics.append(f"{info.file_name}：fast-xlsx 回退到 openpyxl，原因：{fallback_reason}")
        return data
    finally:
        wb.close()


def _read_dictionary_file(info: DictionaryLanguageFile) -> DictionaryFileData:
    try:
        return _read_dictionary_file_fast(info)
    except Exception as exc:
        return _read_dictionary_file_openpyxl(info, str(exc))


def _normalize_column_key(value: str) -> str:
    return normalize_header(value).lstrip("^$")


def _header_tokens(value: str) -> set[str]:
    normalized = _normalize_column_key(value)
    if not normalized:
        return set()
    tokens = {token for token in re.split(r"[_\s]+", normalized) if token}
    # 保留组合 token，方便匹配 pt_br / text_id / id_check / source_text。
    tokens.add(normalized)
    if "pt" in tokens and "br" in tokens:
        tokens.add("pt_br")
        tokens.add("ptbr")
    if "source" in tokens and "text" in tokens:
        tokens.add("source_text")
    if "text" in tokens and "id" in tokens:
        tokens.add("text_id")
    if "id" in tokens and "check" in tokens:
        tokens.add("id_check")
    if "check" in tokens and "id" in tokens:
        tokens.add("check_id")
    if "char" in tokens and "limit" in tokens:
        tokens.add("char_limit")
    if "in" in tokens and "charge" in tokens:
        tokens.add("in_charge")
    compact = normalized.replace("_", "")
    # 兼容“简体CN / 韩语KR / 泰语TH / 波斯语FA”这类中文展示名与
    # 语言代码粘连的表头。这里只在 token 结尾命中，避免 en 误命中 content。
    compact_codes = {
        "cn", "tw", "en", "de", "fr", "ru", "es", "jp", "ja", "kr", "ko", "th", "it",
        "pt", "ptbr", "tr", "vn", "vi", "ua", "uk", "arb", "ar", "id", "hin", "hi", "ur", "fa", "ms",
    }
    for code in compact_codes:
        if len(compact) > len(code) and compact.endswith(code):
            tokens.add(code)
            if code == "ptbr":
                tokens.add("pt_br")
    return tokens


def _has_translation_marker(value: str) -> bool:
    normalized = _normalize_column_key(value)
    tokens = _header_tokens(value)
    markers = {
        "translation", "translate", "target", "target_text", "译文", "譯文", "翻译", "翻譯", "文本內容", "文本内容",
        "source", "source_text", "原文", "源文本",
    }
    return any(marker in tokens or marker in normalized for marker in markers)


def _is_auxiliary_header(header: str, display_header: str = "") -> bool:
    joined = f"{header} {display_header}".strip()
    normalized = _normalize_column_key(joined)
    if not normalized:
        return False
    if normalized.startswith("_"):
        normalized = normalized[1:]
    tokens = _header_tokens(joined)
    if {"id", "check"}.issubset(tokens) or "id_check" in tokens or "check_id" in tokens:
        return True
    if {"char", "limit"}.issubset(tokens) or "char_limit" in tokens:
        return True
    if {"in", "charge"}.issubset(tokens) or "in_charge" in tokens:
        return True
    return any(_normalize_column_key(keyword) in normalized for keyword in AUXILIARY_KEYWORDS)


def _is_type_row(row: Sequence[str]) -> bool:
    values = [_normalize_column_key(cell) for cell in row if str(cell).strip()]
    if not values:
        return False
    hits = sum(1 for value in values if value in TYPE_TOKENS or value.endswith("[]") or value.rstrip("0123456789") in TYPE_TOKENS)
    return hits >= max(2, int(len(values) * 0.6))


def _looks_like_body_text(value: str) -> bool:
    text = str(value or "").strip()
    if not text:
        return False
    normalized = _normalize_column_key(text)
    if re.fullmatch(r"\d+(?:\.\d+)?", normalized):
        return True
    sentence_marks = sum(text.count(ch) for ch in "，。！？；,.!?;。")
    rich_hits = len(re.findall(r"\[[^\]]+\]|<[^>]+>|\{[^{}]+\}|%", text))
    cjk_count = len(re.findall(r"[\u4e00-\u9fff]", text))
    alpha_words = len(re.findall(r"[A-Za-z]{3,}", text))
    if len(text) >= 28 and (sentence_marks >= 1 or rich_hits >= 1 or cjk_count >= 8 or alpha_words >= 5):
        return True
    if rich_hits >= 2:
        return True
    return False


def _row_looks_like_data(row: Sequence[str]) -> bool:
    values = [str(cell).strip() for cell in row]
    if not any(values):
        return False
    first_non_empty = next((value for value in values if value), "")
    numeric_id_like = bool(re.fullmatch(r"\d{3,}", first_non_empty))
    body_cells = sum(1 for value in values if _looks_like_body_text(value))
    return numeric_id_like and body_cells >= 1


@lru_cache(maxsize=4096)
def _language_from_header(header: str, display_header: str = "") -> Tuple[str, str, str]:
    """Return (code, name, confidence).

    支持纯程序字段（cn/en/tw/de/fr/id/ru）、三行表头展示名，以及
    “EN 翻译 / EN Translation”这类单行混合展示名；同时明确排除
    Text ID / ID Check / Notes / Char Limit 等辅助列。
    """
    raw = str(header or "").strip()
    display = str(display_header or "").strip()
    if not raw and not display:
        return "", "", ""
    if _is_auxiliary_header(raw, display):
        return "", "", "排除"

    norm = _normalize_column_key(raw)
    display_norm = _normalize_column_key(display)
    joined = f"{raw} {display}".strip()
    joined_norm = _normalize_column_key(joined)
    tokens = _header_tokens(joined)
    raw_tokens = _header_tokens(raw)
    display_tokens = _header_tokens(display)

    # Text ID / ID Check 这类字段不能被 ID 语言误伤。
    if "text_id" in tokens or "id_check" in tokens or "check_id" in tokens:
        return "", "", "排除"

    # 程序字段行：cn/en/tw/de/fr/id/ru 等精确命中。ID 允许命中，但主键列会在上层排除。
    for code in LANGUAGE_ORDER:
        name, aliases = LANGUAGE_META[code]
        alias_norms = {_normalize_column_key(alias) for alias in aliases}
        if norm in alias_norms:
            # 单行表头里的 id 必须是“ID 翻译/ID Translation”或纯程序字段 id。
            if code == "ID" and norm != "id" and not _has_translation_marker(raw):
                continue
            return code, name, "高"
        if display_norm in alias_norms:
            return code, name, "高"

    # Traditional CN / 繁体中文类字段必须优先判为繁体，避免其中的 CN 被误判为简中。
    if ("traditional" in tokens or "traditional_cn" in tokens or "繁体" in joined_norm or "繁體" in joined_norm) and _has_translation_marker(joined):
        return "TW", LANGUAGE_META["TW"][0], "高"

    # 单行混合表头：EN 翻译 / EN Translation、PT_BR 翻译、UA 翻译等。
    code_tokens = {
        "CN": {"cn", "zh", "zh_cn", "simplified_cn"},
        "TW": {"tw", "zh_tw", "traditional", "traditional_cn"},
        "EN": {"en", "english"},
        "DE": {"de", "german"},
        "FR": {"fr", "french"},
        "ID": {"id", "indonesian"},
        "RU": {"ru", "russian"},
        "ES": {"es", "spanish"},
        "IT": {"it", "italian"},
        "JP": {"jp", "ja", "japanese"},
        "KR": {"kr", "ko", "korean"},
        "PT_BR": {"pt_br", "ptbr", "brazilian_portuguese"},
        "PT": {"pt", "portuguese"},
        "TH": {"th", "thai"},
        "TR": {"tr", "turkish"},
        "VN": {"vn", "vi", "vietnamese"},
        "UA": {"ua", "uk", "ukrainian"},
        "ARB": {"arb", "ar", "arabic"},
        "HIN": {"hin", "hi", "hindi"},
        "UR": {"ur", "urdu"},
        "FA": {"fa", "farsi", "persian"},
        "MS": {"ms", "malay"},
    }
    marker = _has_translation_marker(joined)
    for code in LANGUAGE_ORDER:
        if tokens & code_tokens.get(code, set()):
            if code == "ID" and not marker and norm != "id":
                continue
            # 双表头结构中程序字段常常都叫 t，语言信息只写在上一行展示名里，
            # 例如“韩语KR / t”。这种情况下即使没有 Translation 标记，也应按
            # 展示名中的语言代码识别。
            display_language_hint = bool(display and norm in {"t", "text", "target", "translation", "译文", "翻译"})
            if marker or code in {"CN", "TW"} or display_language_hint:
                name = LANGUAGE_META[code][0]
                return code, name, "高" if marker else "中"

    # 中文/英文长别名包含匹配，避免 en 误命中 content。
    for code in LANGUAGE_ORDER:
        name, aliases = LANGUAGE_META[code]
        for alias in aliases:
            alias_norm = _normalize_column_key(alias)
            if len(alias_norm) >= 4 and joined_norm and (alias_norm in joined_norm or joined_norm in alias_norm):
                if code == "ID" and "text_id" in tokens:
                    continue
                return code, name, "中"
    return "", "", ""


def _find_key_column(headers: Sequence[str], preferred: str = "") -> int:
    """Find a stable text-id key column without letting empty cells or ID Check win."""
    normalized = [_normalize_column_key(header) for header in headers]
    if preferred:
        preferred_norm = _normalize_column_key(preferred)
        for index, header_norm in enumerate(normalized):
            if header_norm and header_norm == preferred_norm and not _is_auxiliary_header(headers[index]):
                return index

    exact_priority = ["#tid", "tid", "text_id", "string_id", "key", "#id"]
    for token in exact_priority:
        for index, header_norm in enumerate(normalized):
            if not header_norm or _is_auxiliary_header(headers[index]):
                continue
            if header_norm == token:
                return index

    contains_priority = ["字串編號", "字串编号", "字串id", "文本編號", "文本编号", "文本id", "text_id"]
    for token in contains_priority:
        token_norm = _normalize_column_key(token)
        for index, header_norm in enumerate(normalized):
            if not header_norm or _is_auxiliary_header(headers[index]):
                continue
            if token_norm in header_norm or header_norm in token_norm:
                return index

    # 最后才允许单独 id 作为主键，且必须不是“ID 翻译/ID Translation”这类语言列。
    for index, header_norm in enumerate(normalized):
        if not header_norm or _is_auxiliary_header(headers[index]):
            continue
        if header_norm == "id" and not _has_translation_marker(headers[index]):
            return index
    return -1


def _header_row_score(row: Sequence[str], next_row: Sequence[str] | None = None, prev_row: Sequence[str] | None = None) -> int:
    values = [str(cell).strip() for cell in row]
    if not any(values):
        return 0
    non_empty = sum(1 for value in values if value)
    key_col = _find_key_column(values, "")
    key_hit = 1 if key_col >= 0 else 0
    lang_hits = 0
    aux_hits = 0
    mixed_header_hits = 0
    for cell in values:
        code, _name, _confidence = _language_from_header(cell)
        if code:
            lang_hits += 1
        if _is_auxiliary_header(cell):
            aux_hits += 1
        cell_norm = _normalize_column_key(cell)
        if any(token in cell_norm for token in ["translation", "source_text", "text_id", "翻译", "翻譯", "原文", "traditional_cn"]):
            mixed_header_hits += 1

    normalized = [_normalize_column_key(value) for value in values if value]
    program_hits = sum(1 for value in normalized if value in PROGRAM_LANGUAGE_CODES or value in {"#tid", "tid", "text_id", "key"})
    type_penalty = 25 if _is_type_row(values) else 0
    data_row_penalty = 80 if _row_looks_like_data(values) else 0
    long_text_penalty = sum(1 for value in values if _looks_like_body_text(value)) * 8
    placeholder_penalty = sum(1 for value in values if re.search(r"\[[^\]]+\]|\{[^{}]+\}|<[^>]+>|%", value)) * 4
    numeric_first_penalty = 18 if values and re.fullmatch(r"\d{3,}", values[0] or "") else 0
    empty_penalty = max(0, len(values) - non_empty) // 8
    data_bonus = 0
    if next_row:
        if key_col >= 0 and key_col < len(next_row) and re.fullmatch(r"\d{2,}", str(next_row[key_col]).strip()):
            data_bonus += 12
        # 下一行存在正文且本行有语言列，说明本行很可能是真表头。
        if lang_hits >= 2 and any(_looks_like_body_text(cell) for cell in next_row):
            data_bonus += 10
    return (
        key_hit * 35
        + lang_hits * 12
        + program_hits * 8
        + mixed_header_hits * 6
        + aux_hits * 2
        + min(non_empty, 24)
        + data_bonus
        - type_penalty
        - data_row_penalty
        - long_text_penalty
        - placeholder_penalty
        - numeric_first_penalty
        - empty_penalty
    )

def _detect_header_layout(rows: List[List[str]], forced_header_row: int = 0, forced_data_start_row: int = 0) -> Tuple[int, int, List[str], List[str], str, List[str]]:
    diagnostics: List[str] = []
    if forced_header_row and forced_header_row > 1 and forced_header_row <= len(rows):
        header_index = forced_header_row - 1
        diagnostics.append(f"使用用户指定表头行：第 {forced_header_row} 行")
    else:
        candidates = []
        for idx, row in enumerate(rows[: min(len(rows), 20)]):
            prev_row = rows[idx - 1] if idx > 0 else []
            next_row = rows[idx + 1] if idx + 1 < len(rows) else []
            candidates.append((idx, _header_row_score(row, next_row, prev_row)))
        candidates.sort(key=lambda item: item[1], reverse=True)
        header_index = candidates[0][0] if candidates and candidates[0][1] > 0 else 0
        diagnostics.append("自动识别表头行：第 {} 行（评分 {}）".format(header_index + 1, candidates[0][1] if candidates else 0))

    type_headers: List[str] = []
    display_headers: List[str] = []
    structure_type = "单行表头"
    if header_index >= 1 and _is_type_row(rows[header_index - 1]):
        type_headers = [str(cell).strip() for cell in rows[header_index - 1]]
        if header_index >= 2:
            display_headers = [str(cell).strip() for cell in rows[header_index - 2]]
            structure_type = "三行表头"
        else:
            structure_type = "双行表头"
    elif header_index >= 1:
        prev = [str(cell).strip() for cell in rows[header_index - 1]]
        # 上一行像展示名而非说明行时，保留为展示名。
        prev_lang_hits = sum(1 for cell in prev if _language_from_header("", cell)[0])
        if prev_lang_hits >= 2:
            display_headers = prev
            structure_type = "双行表头"

    if forced_data_start_row and forced_data_start_row > header_index + 1:
        data_start_index = forced_data_start_row - 1
        diagnostics.append(f"使用用户指定数据起始行：第 {forced_data_start_row} 行")
    else:
        key_col = _find_key_column(rows[header_index], "")
        data_start_index = header_index + 1
        for idx in range(header_index + 1, len(rows)):
            row = rows[idx]
            if not any(str(cell).strip() for cell in row):
                continue
            if _is_type_row(row):
                continue
            if key_col >= 0 and key_col < len(row) and str(row[key_col]).strip():
                data_start_index = idx
                break
    diagnostics.append(f"自动识别数据起始行：第 {data_start_index + 1} 行")
    return header_index, data_start_index, display_headers, type_headers, structure_type, diagnostics


SKIP_SHEET_NAME_KEYWORDS = {"首页", "说明", "readme", "README", "history", "changelog", "版本", "目录", "封面"}


def _analyze_localization_sheet_sample(sheet_name: str, rows: List[List[str]], forced_header_row: int = 0, forced_data_start_row: int = 0) -> Tuple[int, str, bool, int, int]:
    if not rows:
        return -10000, "空 Sheet", False, 0, 0
    try:
        header_index, data_start_index, display_headers, type_headers, _structure, _diag = _detect_header_layout(rows, forced_header_row, forced_data_start_row)
    except Exception as exc:
        return -10000, f"表头识别异常：{exc}", False, 0, 0
    if header_index >= len(rows):
        return -10000, "表头超出范围", False, 0, 0
    headers = [str(cell).strip() for cell in rows[header_index]]
    key_col = _find_key_column(headers, "")
    columns = detect_language_columns(headers, key_col, display_headers, type_headers) if key_col >= 0 else []
    record_hits = 0
    for row in rows[data_start_index:]:
        if key_col >= 0 and key_col < len(row) and str(row[key_col]).strip():
            record_hits += 1
    name_norm = str(sheet_name or "").strip().lower()
    name_penalty = 220 if any(keyword.lower() in name_norm for keyword in SKIP_SHEET_NAME_KEYWORDS) else 0
    # 真实单 Sheet 多语言表通常具有“主键列 + 多语言列 + 若干数据行”。
    score = (120 if key_col >= 0 else 0) + len(columns) * 45 + min(record_hits, 40) * 3 - name_penalty
    reason = f"主键列={'有' if key_col >= 0 else '无'}；语言列={len(columns)}；样本记录={record_hits}"
    is_candidate = key_col >= 0 and len(columns) >= 2 and record_hits > 0
    return score, reason, is_candidate, len(columns), record_hits


def _score_localization_sheet_sample(sheet_name: str, rows: List[List[str]], forced_header_row: int = 0, forced_data_start_row: int = 0) -> Tuple[int, str]:
    score, reason, _is_candidate, _language_count, _record_hits = _analyze_localization_sheet_sample(sheet_name, rows, forced_header_row, forced_data_start_row)
    return score, reason


def _sample_worksheet_rows(ws, max_rows: int = 80) -> List[List[str]]:
    return _worksheet_rows_as_text(ws, max_rows=max_rows, strip=True)


def list_localization_table_sheets(
    path: str,
    options: LocalizationCheckOptions | None = None,
    progress_callback: ProgressCallback | None = None,
) -> List[Tuple[str, int, str]]:
    """返回工作簿中符合“单 Sheet 多语言列”结构的 Sheet。

    用于“多 Sheet 翻译表”模式：一个 Excel 里包含多个版本 Sheet，
    每个 Sheet 自己都是 Text ID + CN/EN/DE... 多语言列。
    """
    file_path = Path(path)
    if file_path.suffix.lower() != ".xlsx":
        return []
    options = options or LocalizationCheckOptions()
    forced_header = safe_int(options.header_row, 1) if safe_int(options.header_row, 1) > 1 else 0
    forced_start = safe_int(options.data_start_row, 2) if safe_int(options.data_start_row, 2) > 2 else 0
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FileReadError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    candidates: List[Tuple[str, int, str]] = []
    _emit_progress(progress_callback, "scan_open", "正在打开工作簿并扫描 Sheet 结构…", ratio=0.03)
    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        worksheets = list(wb.worksheets)
        total = len(worksheets)
        for idx, ws in enumerate(worksheets, start=1):
            _emit_progress(
                progress_callback,
                "scan_sheet",
                f"正在扫描 Sheet {idx}/{total}：{ws.title}",
                current=idx,
                total=total,
                ratio=0.03 + (idx / max(total, 1)) * 0.10,
                sheet_name=ws.title,
            )
            sample_rows = _sample_worksheet_rows(ws)
            score, reason, is_candidate, _language_count, _record_hits = _analyze_localization_sheet_sample(ws.title, sample_rows, forced_header, forced_start)
            if is_candidate:
                candidates.append((ws.title, score, reason))
    finally:
        wb.close()
    return candidates


def detect_multi_table_sheet_columns(path: str, options: LocalizationCheckOptions | None = None) -> Tuple[List[Tuple[str, int, str]], List[LocalizationColumn]]:
    """轻量识别“多 Sheet 翻译表”的候选 Sheet 和语言并集。

    只读取每个 Sheet 前 80 行，用于 UI 预览和语言勾选；正式检查时再读取完整数据。
    """
    file_path = Path(path)
    if file_path.suffix.lower() != ".xlsx":
        return [], []
    options = options or LocalizationCheckOptions()
    forced_header = safe_int(options.header_row, 1) if safe_int(options.header_row, 1) > 1 else 0
    forced_start = safe_int(options.data_start_row, 2) if safe_int(options.data_start_row, 2) > 2 else 0
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FileReadError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    candidates: List[Tuple[str, int, str]] = []
    columns_by_header: Dict[str, LocalizationColumn] = {}
    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            sample_rows = _sample_worksheet_rows(ws)
            score, reason, is_candidate, _language_count, _record_hits = _analyze_localization_sheet_sample(ws.title, sample_rows, forced_header, forced_start)
            if not is_candidate:
                continue
            candidates.append((ws.title, score, reason))
            try:
                header_index, _data_start_index, display_headers, type_headers, _structure, _diag = _detect_header_layout(sample_rows, forced_header, forced_start)
                headers = [str(cell).strip() for cell in sample_rows[header_index]]
                key_col = _find_key_column(headers, options.key_column)
                columns = detect_language_columns(headers, key_col, display_headers, type_headers) if key_col >= 0 else []
                _merge_localization_columns(columns_by_header, columns)
            except Exception:
                continue
    finally:
        wb.close()

    columns = sorted(
        columns_by_header.values(),
        key=lambda col: LANGUAGE_ORDER.index(col.code or col.header) if (col.code or col.header) in LANGUAGE_ORDER else 999,
    )
    return candidates, columns


def _read_localization_table_rows(path: str, sheet_name: str, forced_header_row: int = 0, forced_data_start_row: int = 0) -> Tuple[List[List[str]], str, List[str]]:
    file_path = Path(path)
    if file_path.suffix.lower() != ".xlsx":
        rows, actual_sheet = read_table_rows(path, sheet_name)
        return rows, actual_sheet, []

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FileReadError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            return _worksheet_rows_as_text(ws), ws.title, []
        scored: List[Tuple[int, str, str]] = []
        for ws in wb.worksheets:
            sample_rows = _sample_worksheet_rows(ws)
            score, reason = _score_localization_sheet_sample(ws.title, sample_rows, forced_header_row, forced_data_start_row)
            scored.append((score, ws.title, reason))
        scored.sort(key=lambda item: item[0], reverse=True)
        best_score, best_sheet, best_reason = scored[0] if scored else (-10000, "", "")
        if best_sheet:
            ws = wb[best_sheet]
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            rows = _worksheet_rows_as_text(ws)
        else:
            rows = []
    finally:
        wb.close()

    if not best_sheet:
        rows, actual_sheet = read_table_rows(path, "")
        return rows, actual_sheet, []

    diagnostics = [f"自动选择数据 Sheet：{best_sheet}（评分 {best_score}，{best_reason}）"]
    return rows, best_sheet, diagnostics


def parse_localization_table(options: LocalizationCheckOptions, path: str) -> LocalizationTable:
    forced_header = safe_int(options.header_row, 1) if safe_int(options.header_row, 1) > 1 else 0
    forced_start = safe_int(options.data_start_row, 2) if safe_int(options.data_start_row, 2) > 2 else 0
    rows, actual_sheet, auto_sheet_diagnostics = _read_localization_table_rows(path, options.sheet_name, forced_header, forced_start)
    if not rows:
        raise FileReadError(f"文件为空：{path}")

    header_index, data_start_index, display_headers, type_headers, structure_type, diagnostics = _detect_header_layout(rows, forced_header, forced_start)
    diagnostics = list(auto_sheet_diagnostics) + diagnostics
    if header_index >= len(rows):
        raise FileReadError(f"表头行超出文件范围：第 {header_index + 1} 行")

    headers = [str(cell).strip() for cell in rows[header_index]]
    key_col = _find_key_column(headers, options.key_column)
    if key_col < 0:
        expected = options.key_column or "#tid / tid / text_id / key / id"
        raise FileReadError(f"表头识别失败：未找到主键列（期望：{expected}）。已识别表头行：第 {header_index + 1} 行")

    excluded_columns: Dict[str, str] = {}
    max_len = max(len(headers), len(display_headers), len(type_headers))
    for idx in range(max_len):
        header = headers[idx] if idx < len(headers) else ""
        display = display_headers[idx] if idx < len(display_headers) else ""
        if idx == key_col:
            continue
        if not str(header).strip() and not str(display).strip():
            excluded_columns[f"列{idx + 1}"] = "空列，默认排除"
            continue
        if _is_auxiliary_header(header, display):
            excluded_columns[header or display or f"列{idx + 1}"] = "辅助列/状态列/备注列，默认不作为语言列"

    records: Dict[str, LocalizationRow] = {}
    row_records: List[LocalizationRow] = []
    id_rows: Dict[str, List[int]] = {}
    invalid_rows: List[int] = []
    for row_index in range(data_start_index, len(rows)):
        row = rows[row_index]
        row_number = row_index + 1
        if not any(str(cell).strip() for cell in row):
            continue
        text_id = _format_identifier(row[key_col]) if key_col < len(row) else ""
        if not text_id:
            # 说明行/备注行不计入无效 ID；有多个语言列文本但没有主键才提示。
            language_like_values = 0
            for idx, value in enumerate(row):
                if idx == key_col:
                    continue
                if str(value).strip():
                    language_like_values += 1
            if language_like_values >= 2:
                invalid_rows.append(row_number)
            continue
        values = {index: (str(row[index]) if index < len(row) and row[index] is not None else "") for index in range(len(headers))}
        id_rows.setdefault(text_id, []).append(row_number)
        row_record = LocalizationRow(text_id=text_id, row_number=row_number, values=values)
        row_records.append(row_record)
        if text_id not in records:
            records[text_id] = row_record

    if not records:
        diagnostics.append("当前记录数为 0，请检查表头所在行、数据起始行和主键列是否正确")
    lang_preview = []
    for idx, header in enumerate(headers):
        display = display_headers[idx] if idx < len(display_headers) else ""
        code, name, conf = _language_from_header(header, display)
        if code:
            lang_preview.append(f"{header}={name}({conf})")
    if len(lang_preview) <= 2:
        diagnostics.append("当前识别到的语言列较少，请确认是否存在三行表头结构或辅助列误识别")
    if any(_normalize_column_key(h) in {"cn", "en", "tw", "de", "fr", "id", "ru"} for h in headers):
        diagnostics.append("检测到 cn/en/tw/de/fr/id/ru 等程序字段，已按程序字段行作为真实表头")
    if excluded_columns:
        diagnostics.append("检测到辅助列并默认排除：" + "、".join(list(excluded_columns.keys())[:8]))

    return LocalizationTable(
        path=path,
        sheet_name=actual_sheet,
        headers=headers,
        key_col_index=key_col,
        records=records,
        duplicate_ids={text_id: rows_ for text_id, rows_ in id_rows.items() if len(rows_) > 1},
        invalid_rows=invalid_rows,
        display_headers=display_headers,
        type_headers=type_headers,
        header_row=header_index + 1,
        data_start_row=data_start_index + 1,
        structure_type=structure_type,
        excluded_columns=excluded_columns,
        diagnostics=diagnostics,
        row_records=row_records,
    )


def _localization_table_from_rows(
    options: LocalizationCheckOptions,
    path: str,
    rows: List[List[str]],
    actual_sheet: str,
    auto_sheet_diagnostics: Sequence[str] | None = None,
) -> LocalizationTable:
    forced_header = safe_int(options.header_row, 1) if safe_int(options.header_row, 1) > 1 else 0
    forced_start = safe_int(options.data_start_row, 2) if safe_int(options.data_start_row, 2) > 2 else 0
    if not rows:
        raise FileReadError(f"文件为空：{path}")

    header_index, data_start_index, display_headers, type_headers, structure_type, diagnostics = _detect_header_layout(rows, forced_header, forced_start)
    diagnostics = list(auto_sheet_diagnostics or []) + diagnostics
    if header_index >= len(rows):
        raise FileReadError(f"表头行超出文件范围：第 {header_index + 1} 行")

    headers = [str(cell).strip() for cell in rows[header_index]]
    key_col = _find_key_column(headers, options.key_column)
    if key_col < 0:
        expected = options.key_column or "#tid / tid / text_id / key / id"
        raise FileReadError(f"表头识别失败：未找到主键列（期望：{expected}）。已识别表头行：第 {header_index + 1} 行")

    excluded_columns: Dict[str, str] = {}
    max_len = max(len(headers), len(display_headers), len(type_headers))
    for idx in range(max_len):
        header = headers[idx] if idx < len(headers) else ""
        display = display_headers[idx] if idx < len(display_headers) else ""
        if idx == key_col:
            continue
        if not str(header).strip() and not str(display).strip():
            excluded_columns[f"列{idx + 1}"] = "空列，默认排除"
            continue
        if _is_auxiliary_header(header, display):
            excluded_columns[header or display or f"列{idx + 1}"] = "辅助列/状态列/备注列，默认不作为语言列"

    records: Dict[str, LocalizationRow] = {}
    row_records: List[LocalizationRow] = []
    id_rows: Dict[str, List[int]] = {}
    invalid_rows: List[int] = []
    for row_index in range(data_start_index, len(rows)):
        row = rows[row_index]
        row_number = row_index + 1
        if not any(str(cell).strip() for cell in row):
            continue
        text_id = _format_identifier(row[key_col]) if key_col < len(row) else ""
        if not text_id:
            language_like_values = 0
            for idx, value in enumerate(row):
                if idx == key_col:
                    continue
                if str(value).strip():
                    language_like_values += 1
            if language_like_values >= 2:
                invalid_rows.append(row_number)
            continue
        values = {index: (str(row[index]) if index < len(row) and row[index] is not None else "") for index in range(len(headers))}
        id_rows.setdefault(text_id, []).append(row_number)
        row_record = LocalizationRow(text_id=text_id, row_number=row_number, values=values)
        row_records.append(row_record)
        if text_id not in records:
            records[text_id] = row_record

    if not records:
        diagnostics.append("当前记录数为 0，请检查表头所在行、数据起始行和主键列是否正确")
    lang_preview = []
    for idx, header in enumerate(headers):
        display = display_headers[idx] if idx < len(display_headers) else ""
        code, name, conf = _language_from_header(header, display)
        if code:
            lang_preview.append(f"{header}={name}({conf})")
    if len(lang_preview) <= 2:
        diagnostics.append("当前识别到的语言列较少，请确认是否存在三行表头结构或辅助列误识别")
    if any(_normalize_column_key(h) in {"cn", "en", "tw", "de", "fr", "id", "ru"} for h in headers):
        diagnostics.append("检测到 cn/en/tw/de/fr/id/ru 等程序字段，已按程序字段行作为真实表头")
    if excluded_columns:
        diagnostics.append("检测到辅助列并默认排除：" + "、".join(list(excluded_columns.keys())[:8]))

    return LocalizationTable(
        path=path,
        sheet_name=actual_sheet,
        headers=headers,
        key_col_index=key_col,
        records=records,
        duplicate_ids={text_id: rows_ for text_id, rows_ in id_rows.items() if len(rows_) > 1},
        invalid_rows=invalid_rows,
        display_headers=display_headers,
        type_headers=type_headers,
        header_row=header_index + 1,
        data_start_row=data_start_index + 1,
        structure_type=structure_type,
        excluded_columns=excluded_columns,
        diagnostics=diagnostics,
        row_records=row_records,
    )


def detect_language_columns(headers: Sequence[str], key_column_index: int = -1, display_headers: Sequence[str] | None = None, type_headers: Sequence[str] | None = None) -> List[LocalizationColumn]:
    display_headers = list(display_headers or [])
    detected: List[LocalizationColumn] = []
    used_indexes = set()
    for index, header in enumerate(headers):
        if index == key_column_index or index in used_indexes:
            continue
        display = display_headers[index] if index < len(display_headers) else ""
        code, lang_name, confidence = _language_from_header(header, display)
        if not code:
            continue
        # header 字段用于 UI 勾选和后续回传，必须唯一稳定。双表头游戏表里
        # 多个语言列的程序字段都可能叫 t，因此不能继续把原始 header 当 key。
        detected.append(LocalizationColumn(name=lang_name, header=code, index=index, display_name=display, confidence=confidence or "中", code=code, raw_header=str(header).strip()))
        used_indexes.add(index)

    # 保留旧映射兜底，支持单行展示名表头。
    if not detected:
        mappings = parse_language_mappings(DEFAULT_LANGUAGE_MAPPING_TEXT)
        for mapping in mappings:
            candidates = list(dict.fromkeys(mapping.source_candidates + mapping.config_candidates + [mapping.name]))
            col = _find_language_candidate_column(headers, candidates, {key_column_index} | used_indexes)
            if col >= 0 and col != key_column_index and col not in used_indexes:
                code, _name, _confidence = _language_from_header(mapping.name)
                detected.append(LocalizationColumn(name=mapping.name, header=code or mapping.name, index=col, display_name=str(headers[col]).strip(), confidence="中", code=code, raw_header=str(headers[col]).strip()))
                used_indexes.add(col)

    # 如果仍失败，不再把所有非空列都当语言列，避免状态/日期/备注误报；只返回空让 UI 明确提示。
    detected.sort(key=lambda col: LANGUAGE_ORDER.index(col.code or _language_from_header(col.header, col.display_name)[0]) if (col.code or _language_from_header(col.header, col.display_name)[0]) in LANGUAGE_ORDER else 999)
    return detected


def _find_language_candidate_column(headers: Sequence[str], candidates: Sequence[str], excluded_indexes=None) -> int:
    excluded_indexes = set(excluded_indexes or [])
    normalized_headers = [_normalize_column_key(header) for header in headers]
    normalized_candidates = [_normalize_column_key(candidate) for candidate in candidates if str(candidate).strip()]
    for candidate in normalized_candidates:
        for index, header in enumerate(normalized_headers):
            if index in excluded_indexes or not header:
                continue
            if header == candidate:
                return index
    for candidate in normalized_candidates:
        if not candidate:
            continue
        for index, header in enumerate(normalized_headers):
            if index in excluded_indexes or not header:
                continue
            if len(candidate) <= 3 and not re.search(r"[\u4e00-\u9fff]", candidate):
                padded = f"_{header}_"
                if f"_{candidate}_" in padded:
                    return index
                continue
            if candidate in {"中文", "翻译", "文本"}:
                continue
            if len(candidate) >= 4 and (candidate in header or (len(header) >= 4 and header in candidate)):
                return index
    return -1


def auto_detect_source_language(headers: Sequence[str], enabled_headers: Sequence[str]) -> str:
    available = list(enabled_headers) or list(headers)
    normalized_to_header = {_normalize_column_key(header): header for header in available}
    for candidate in ["cn", "zh_cn", "中文简体", "简中", "中文", "en", "english"]:
        key = _normalize_column_key(candidate)
        if key in normalized_to_header:
            return normalized_to_header[key]
    for header in available:
        code, _name, _confidence = _language_from_header(header)
        if code in {"CN", "TW", "EN"}:
            return header
    return available[0] if available else ""


def _localization_column_label(column: LocalizationColumn | None, fallback: str = "") -> str:
    if column is None:
        return fallback
    raw_header = (column.raw_header or "").strip()
    raw_key = _normalize_column_key(raw_header)
    if raw_header and raw_key not in {"t", "text", "target", "translation", "译文", "翻译"}:
        return raw_header
    return column.code or column.header or column.name or column.display_name or fallback


def _is_repeated_text_program_table(table: LocalizationTable) -> bool:
    normalized_headers = [_normalize_column_key(header) for header in table.headers]
    repeated_t = sum(1 for header in normalized_headers if header == "t")
    has_display_language = sum(1 for display in table.display_headers if _language_from_header("", display)[0])
    return repeated_t >= 3 and has_display_language >= 2


def _resolve_enabled_language_columns(
    table: LocalizationTable,
    detected_columns: Sequence[LocalizationColumn],
    requested_headers: Sequence[str],
) -> List[LocalizationColumn]:
    by_index = {column.index: column for column in detected_columns}
    selected: List[LocalizationColumn] = []

    if not requested_headers:
        selected = list(detected_columns)
    else:
        for header in requested_headers:
            idx = _find_detected_language_index(detected_columns, header)
            if idx < 0:
                idx = _find_header_index(table.headers, header)
            if idx < 0 or idx == table.key_col_index:
                continue
            display = table.display_headers[idx] if idx < len(table.display_headers) else ""
            if _is_auxiliary_header(table.headers[idx] if idx < len(table.headers) else "", display):
                continue
            column = by_index.get(idx)
            if column is None:
                code, name, confidence = _language_from_header(table.headers[idx] if idx < len(table.headers) else "", display)
                column = LocalizationColumn(
                    name=name or header,
                    header=code or header,
                    index=idx,
                    display_name=display,
                    confidence=confidence or "中",
                    code=code,
                    raw_header=table.headers[idx] if idx < len(table.headers) else "",
                )
            selected.append(column)

    result: List[LocalizationColumn] = []
    seen = set()
    for column in selected:
        if column.index in seen or column.index == table.key_col_index:
            continue
        display = table.display_headers[column.index] if column.index < len(table.display_headers) else ""
        raw_header = table.headers[column.index] if column.index < len(table.headers) else ""
        if _is_auxiliary_header(raw_header, display):
            continue
        result.append(column)
        seen.add(column.index)
    return result


def _find_header_index(headers: Sequence[str], header_or_name: str) -> int:
    if not header_or_name:
        return -1
    preferred = _normalize_column_key(header_or_name)
    if not preferred:
        return -1
    for index, header in enumerate(headers):
        normalized = _normalize_column_key(header)
        if normalized and normalized == preferred:
            return index
    for index, header in enumerate(headers):
        normalized = _normalize_column_key(header)
        if not normalized:
            continue
        # 两位语言代码必须边界/精确，避免 en 命中 content。
        if len(preferred) <= 3 and re.fullmatch(r"[a-z_]+", preferred):
            if normalized == preferred:
                return index
            continue
        if preferred in normalized or normalized in preferred:
            return index
    return -1


def _find_detected_language_index(detected_columns: Sequence[LocalizationColumn], header_or_name: str) -> int:
    if not header_or_name:
        return -1
    preferred = _normalize_column_key(header_or_name)
    if not preferred:
        return -1
    for column in detected_columns:
        if (
            _normalize_column_key(column.header) == preferred
            or _normalize_column_key(column.name) == preferred
            or _normalize_column_key(column.code) == preferred
            or _normalize_column_key(column.display_name) == preferred
            or _normalize_column_key(column.raw_header) == preferred
        ):
            return column.index
    for column in detected_columns:
        header_key = _normalize_column_key(column.header)
        name_key = _normalize_column_key(column.name)
        code_key = _normalize_column_key(column.code)
        display_key = _normalize_column_key(column.display_name)
        raw_key = _normalize_column_key(column.raw_header)
        if header_key and (preferred == header_key or (len(preferred) > 3 and (preferred in header_key or header_key in preferred))):
            return column.index
        if name_key and (preferred == name_key or (len(preferred) > 3 and (preferred in name_key or name_key in preferred))):
            return column.index
        if code_key and preferred == code_key:
            return column.index
        if display_key and (preferred == display_key or (len(preferred) > 3 and (preferred in display_key or display_key in preferred))):
            return column.index
        if raw_key and (preferred == raw_key or (len(preferred) > 3 and (preferred in raw_key or raw_key in preferred))):
            return column.index
    return -1


# ===== 文本检查规则底座 =====

CHINESE_RE = re.compile(r"[\u4e00-\u9fff]")
PRINTF_PLACEHOLDER_RE = re.compile(
    r"(?<![A-Za-z0-9_])%(?:\d+\$)?[+#\-0,(]*(?:\d+)?(?:\.\d+)?[sdifumfFeEgGxXoOcC](?:\d+)?(?![A-Za-z0-9_])"
)
BRACKETED_PERCENT_PLACEHOLDER_RE = re.compile(r"\[%[A-Za-z][A-Za-z0-9_]{0,40}\]")
POSITIONAL_PERCENT_PLACEHOLDER_RE = re.compile(r"(?<![A-Za-z0-9_])%\d+(?![A-Za-z0-9_$])")
HASH_PLACEHOLDER_RE = re.compile(r"(?<![A-Za-z0-9_])#[A-Z][A-Z0-9_]{1,40}(?![A-Za-z0-9_])")
BRACE_PLACEHOLDER_RE = re.compile(r"\{[^{}\r\n]{1,80}\}")
SQUARE_PLACEHOLDER_RE = re.compile(r"\[(?:\d+|(?:value|text|param|var|num|arg|count|amount|id)[A-Za-z0-9_]{0,40})\]", re.IGNORECASE)
ANGLE_PLACEHOLDER_RE = re.compile(r"<\s*(?:value|text|param|var|num|arg)[A-Za-z0-9_]{0,80}\s*>", re.IGNORECASE)
ESCAPED_NEWLINE_RE = re.compile(r"\\n")
ANGLE_TAG_RE = re.compile(r"<\s*(/)?\s*([A-Za-z][\w\-]*)(?:\s+([^<>]*?))?\s*(/)?>")
BRACKET_TAG_RE = re.compile(r"\[\s*(/)?\s*([A-Za-z][\w\-]*)(?:(?:\s+|=)([^\]]*?))?\s*\]")
ANY_ANGLE_TAG_RE = re.compile(r"<[^<>]*>")
ANY_BRACKET_TAG_RE = re.compile(r"\[(?:/?(?:color|size|font|b|i|u|url|ref|link|br)(?:(?:\s+|=)[^\]]*)?)\]", re.IGNORECASE)
PLACEHOLDER_SIGNAL_RE = re.compile(r"[%$#\\{}\[\]<>]")
TAG_SIGNAL_RE = re.compile(r"[\[\]<>]")
NUMBER_SIGNAL_RE = re.compile(r"[0-9零〇一二两兩俩倆三四五六七八九十百千万萬ⅠⅡⅢⅣⅤ%％]")
HEX_COLOR_RE = re.compile(r"(?:#[0-9A-Fa-f]{3,8}\b|0x[0-9A-Fa-f]{3,8}\b)")
UNICODE_ESCAPE_RE = re.compile(r"\\(?:u[0-9A-Fa-f]{4}|x[0-9A-Fa-f]{2})")
HTML_ENTITY_RE = re.compile(r"&\#(?:\d+|x[0-9A-Fa-f]+);")
DOLLAR_PLACEHOLDER_RE = re.compile(r"\$\{[A-Za-z_][A-Za-z0-9_]*\}")
IDENTIFIER_NUMBER_RE = re.compile(r"(?<![A-Za-z0-9_])(?:[A-Za-z_][A-Za-z0-9_]*\d+[A-Za-z0-9_]*(?:\.\d+)*)(?![A-Za-z0-9_])")
DATE_RE = re.compile(r"(?<!\d)(\d{4})[/-](\d{1,2})[/-](\d{1,2})(?!\d)")
CN_DATE_RE = re.compile(r"([一二三四五六七八九十两兩俩倆零〇0-9]{1,4})\s*月\s*([一二三四五六七八九十两兩俩倆零〇0-9]{1,4})\s*[日号號]")
TIME_RE = re.compile(r"(?<!\d)([01]?\d|2[0-3]):([0-5]\d)(?!\d)")
CN_HOUR_RE = re.compile(r"(?<!\d)([01]?\d|2[0-3])\s*(?:点|點|时|時)(?!\d)")
LV_STAGE_RE = re.compile(
    r"(?<![A-Za-z0-9_])((?:lv|level|tier|niv(?:eau)?|st(?:ufe)?|ур(?:овень)?|r|t)\s*[.\-]?\s*\d+)(?![A-Za-z0-9_])",
    re.IGNORECASE,
)
CN_LEVEL_RE = re.compile(r"(?:等级|等級)\s*([一二三四五六七八九十两兩俩倆百千万萬零〇0-9]+)")
ARABIC_CN_LEVEL_RE = re.compile(r"(?<![A-Za-z0-9_])(\d+)\s*(?:级|級)(?![A-Za-z0-9_])")
THOUSAND_SEP_RE = re.compile(r"(?<!\d)(\d{1,3}(?:[,\.\s]\d{3})+)(?!\d)")
MULTIPLIER_X_RE = re.compile(r"(?<![A-Za-z0-9_])([+-]?\d+(?:\.\d+)?)\s*[xX倍](?![A-Za-z0-9_])")
CN_NUMBER_UNIT_RE = re.compile(r"(?<!每)([一二三四五六七八九十两兩俩倆百千万萬零〇]+)\s*(?:次|层|層|秒|毫秒|分钟|分鐘|小时|小時|天|日|个|個|名|倍|级|級|星|輪|轮|回|點|点)")
EN_WORD_NUMBER_RE = re.compile(r"\b(?:(once|twice)\s+(?:every|each|per|a|an|daily|weekly|monthly|time|times|second|seconds|sec|secs|minute|minutes|min|mins|hour|hours|day|days)|(?:one|two|three|four|five|six|seven|eight|nine|ten)\s+(?:time|times|second|seconds|sec|secs|minute|minutes|min|mins|hour|hours|day|days|layer|layers|level|levels|star|stars))\b", re.IGNORECASE)
EN_STRONG_WORD_NUMBER_RE = re.compile(r"\b(?:one|two|three|four|five|six|seven|eight|nine|ten)\s+(seconds?|secs?|minutes?|mins?|hours?|days?|levels?|stars?)\b", re.IGNORECASE)
EN_NATURAL_ONE_RE = re.compile(r"\b(?:a|an|one|single|once|each|every)\b", re.IGNORECASE)
EN_NATURAL_QUANTITY_RE = re.compile(r"\b(?:1|one|a|an|single)\s+(?:time|times|instance|instances|item|items|stack|stacks|layer|layers|shield|shields|container|containers|piece|pieces|card|cards|punch|punches)\b", re.IGNORECASE)
LIST_PREFIX_RE = re.compile(r"(?m)^\s*\d+\s*[\.、\)]\s*(?![\-–—~～])")
EN_ORDINAL_RE = re.compile(r"\b(?:(first|second|third|fourth|fifth|sixth|seventh|eighth|ninth|tenth)|([1-9]|10)(?:st|nd|rd|th))\s+(?:time|times|round|rounds|stage|stages|chapter|chapters)\b", re.IGNORECASE)
EN_PREFIX_NUMBER_RE = re.compile(
    r"\b(?:chapter|stage|round|tier|level|goal)\s+(one|two|three|four|five|six|seven|eight|nine|ten)\b",
    re.IGNORECASE,
)
CN_STRONG_NUMBER_UNIT_RE = re.compile(r"([一二三四五六七八九十两兩俩倆百千万萬零〇]+)\s*(?:秒|毫秒|分钟|分鐘|小时|小時|天|日|倍|级|級|星|輪|轮|回)")
CN_NATURAL_ONE_RE = re.compile(r"(?<!第)(?:一|壹|１)\s*(?:个|個|名|张|張|拳|次|层|層|件|份|本|支|枚|项|項|位|条|條|颗|顆|只|隻)")
CN_NATURAL_ARABIC_ONE_RE = re.compile(r"(?<!第)1\s*(?:次|层|層)")
CN_ORDINAL_RE = re.compile(r"第\s*([一二三四五六七八九十两兩俩倆百千万萬零〇0-9]+)\s*(?:次|层|層|轮|輪|回|阶段|階段|章|章节|章節|节|節)")
ROMAN_NUMERAL_RE = re.compile(r"(?<![A-Za-z0-9_])([ⅠⅡⅢⅣⅤ])(?!(?:[A-Za-z0-9_]))")
ASCII_ROMAN_RE = re.compile(r"(?<![A-Za-z0-9_])((?:IV|V|III|II|I))(?![A-Za-z0-9_])")
PERCENT_RE = re.compile(r"(?<![A-Za-z0-9_])([+-]?\d+(?:[\.,]\d+)?)\s*(?:%|％)(?![A-Za-z0-9_])")
PERCENT_WORD_RE = re.compile(r"(?<![A-Za-z0-9_])(\d+(?:[\.,]\d+)?)\s*(?:percent|percentage|prozent|pour\s+cent|persen|por\s+ciento|percento|процент(?:а|ов)?)(?![A-Za-z0-9_])", re.IGNORECASE)
LOCALIZED_DECIMAL_RE = re.compile(r"(?<!\d)(\d+),(\d{1,2})(?!\d)")
RANGE_RE = re.compile(r"(?<![A-Za-z0-9_])([0-9]+)(?:\.)?\s*[\-–—~～]\s*([0-9]+)(?:\.)?(?![A-Za-z0-9_])")
NEGATIVE_NUMBER_RE = re.compile(r"(?<![A-Za-z0-9_.])-(\d+(?:\.\d+)?)(?![A-Za-z0-9_])")
NUMBER_UNIT_PATTERN = r"(?:milliseconds?|msecs?|ms|seconds?|secs?|sec|minutes?|mins?|min|hours?|hrs?|hr|days?|day|m|s|h|d|秒|毫秒|分钟|分鐘|小时|小時|天|日|次|名|个|個|层|層|级|級|點|点|倍|回|轮|輪|星)"
NUMBER_WITH_UNIT_RE = re.compile(r"(?<![A-Za-z0-9_])([+-]?(?:(?:\d{1,3}(?:,\d{3})+)|\d+)(?:\.\d+)?)\s*" + NUMBER_UNIT_PATTERN + r"(?![A-Za-z0-9_])", re.IGNORECASE)
NUMBER_RE = re.compile(r"(?<![A-Za-z0-9_])((?:(?:\d{1,3}(?:,\d{3})+)|\d+)(?:\.\d+)?)(?![A-Za-z0-9_])")

RICH_TAG_NAMES = {"color", "b", "i", "u", "size", "font", "url", "ref", "link", "br"}
SELF_CLOSING_TAG_NAMES = {"br"}
STYLE_TAG_NAMES = {"color", "size", "font"}
FUNCTIONAL_TAG_NAMES = {"url", "ref", "link"}
NATURAL_PUNCT_EQUIVALENTS = {
    "：": {"：", ":"}, ":": {"：", ":"},
    "！": {"！", "!"}, "!": {"！", "!"},
    "？": {"？", "?"}, "?": {"？", "?"},
    "（": {"（", "("}, "(": {"（", "("},
    "）": {"）", ")"}, ")": {"）", ")"},
    "，": {"，", ","}, ",": {"，", ","},
    "。": {"。", "."}, ".": {"。", "."},
}
CRITICAL_SYMBOLS = {"{}", "[]", "<>", "%", "/", "\\", "\\n"}
INVISIBLE_TEXT_RE = re.compile(r"[\s\u200b\u200c\u200d\u2060\ufeff]+")

UNFINISHED_TEXT_PATTERNS = [
    "翻译未完成", "翻譯未完成", "翻译错误，修改中", "翻譯錯誤，修改中", "待翻译", "待翻譯",
    "未翻译", "未翻譯", "todo", "tbd", "untranslated", "need translation", "needs translation",
]

SPELLING_SUSPECT_PATTERNS = [
    (re.compile(r"\bcritica\s+hitl\b", re.IGNORECASE), "疑似应为 Critical Hit"),
    (re.compile(r"\bcritical\s+hitl\b", re.IGNORECASE), "疑似应为 Critical Hit"),
    (re.compile(r"\bcritica\s+hit\b", re.IGNORECASE), "疑似应为 Critical Hit"),
]

TERM_CONSISTENCY_GROUPS = [
    ("自动打野/自动狩猎", ["auto hunt", "auto jungle"]),
    ("相册/图鉴", ["album", "album collection"]),
]



@lru_cache(maxsize=4096)
def _language_code_from_label(language: str) -> str:
    code, _name, _confidence = _language_from_header(language)
    if code:
        return code
    sheet_code = _language_code_from_sheet_name(language)
    if sheet_code:
        return sheet_code
    normalized = _normalize_column_key(language).upper()
    return normalized if normalized in LANGUAGE_META else ""


def _contains_unfinished_marker(text: str) -> str:
    value = ("" if text is None else str(text)).strip()
    normalized = value.lower()
    for pattern in UNFINISHED_TEXT_PATTERNS:
        if pattern.lower() in normalized:
            return pattern
    return ""


def _is_japanese_compatible_cjk(text: str) -> bool:
    value = "" if text is None else str(text)
    # 日文含假名时，汉字通常是正常日文表达；不能直接按 CJK 判残留中文。
    return bool(re.search(r"[\u3040-\u30ff]", value))


JP_SIMPLIFIED_ONLY_RE = re.compile(r"[们这门龙级图过进发个吗奖开关双数备导资统电显]")
JP_CHINESE_SENTENCE_MARKERS_RE = re.compile(r"(翻译|翻譯|未完成|修改中|待翻译|待翻譯|未翻译|未翻譯|简体中文|簡體中文|点击|领取|获得|奖励|可以|无法|是否|当前|需要|请先|已完成)")


def _jp_chinese_residual_level(text: str) -> Tuple[str, str]:
    value = "" if text is None else str(text)
    marker = _contains_unfinished_marker(value)
    if marker:
        return "unfinished", marker
    chinese_fragments = _extract_chinese_fragments(value)
    if not chinese_fragments:
        return "", ""
    # JP 不能按“出现汉字”直接报残留中文。正常日文汉字词、无假名短词都默认兼容。
    # 只有命中明显简体字、中文占位文本或强中文句式时才提示；不确定时降级为疑似。
    simplified_hits = JP_SIMPLIFIED_ONLY_RE.findall(value)
    if JP_CHINESE_SENTENCE_MARKERS_RE.search(value) and simplified_hits:
        return "high", ", ".join(chinese_fragments[:8])
    if JP_CHINESE_SENTENCE_MARKERS_RE.search(value) or len(set(simplified_hits)) >= 2:
        return "suspected", ", ".join(chinese_fragments[:8])
    return "", ""


_CN_NUM = {"零":0,"〇":0,"一":1,"二":2,"两":2,"兩":2,"俩":2,"倆":2,"三":3,"四":4,"五":5,"六":6,"七":7,"八":8,"九":9}
_EN_NUM = {"once":1,"one":1,"twice":2,"two":2,"three":3,"four":4,"five":5,"six":6,"seven":7,"eight":8,"nine":9,"ten":10,"first":1,"second":2,"third":3,"fourth":4,"fifth":5,"sixth":6,"seventh":7,"eighth":8,"ninth":9,"tenth":10}
_ROMAN_NUM = {"Ⅰ":1,"Ⅱ":2,"Ⅲ":3,"Ⅳ":4,"Ⅴ":5,"I":1,"II":2,"III":3,"IV":4,"V":5}

def _cn_number_to_int(text: str) -> int | None:
    raw = (text or "").strip()
    if not raw:
        return None
    if raw.isdigit():
        return int(raw)
    total = 0
    section = 0
    number = 0
    unit_map = {"十":10, "百":100, "千":1000, "万":10000, "萬":10000}
    for char in raw:
        if char in _CN_NUM:
            number = _CN_NUM[char]
        elif char in unit_map:
            unit = unit_map[char]
            if unit >= 10000:
                section = (section + (number or 0)) * unit
                total += section
                section = 0
            else:
                section += (number or 1) * unit
            number = 0
        else:
            return None
    return total + section + number


def _normalize_thousand_token(token: str) -> str:
    return _normalize_number_token(re.sub(r"[,\.\s]", "", token or ""))


def _normalize_duplicate_text(text: str) -> str:
    value = ("" if text is None else str(text)).strip().lower()
    value = re.sub(r"\s+", " ", value)
    return value


def _is_translation_duplicate_candidate(text: str) -> bool:
    value = ("" if text is None else str(text)).strip()
    if not value or _contains_unfinished_marker(value) or _is_pure_tag_fragment(value):
        return False
    if len(value) < 8:
        return False
    if re.fullmatch(r"[\d\s,\.:%+\-]+", value):
        return False
    return True


def _spelling_issue(language: str, text: str) -> Tuple[str, str]:
    if _language_code_from_label(language) != "EN":
        return "", ""
    value = "" if text is None else str(text)
    for regex, suggestion in SPELLING_SUSPECT_PATTERNS:
        m = regex.search(value)
        if m:
            return m.group(0), suggestion
    return "", ""


def _term_spans(text: str, term: str) -> List[Tuple[int, int, str]]:
    pattern = r"(?<![A-Za-z0-9])" + re.escape(term.lower()) + r"(?![A-Za-z0-9])"
    return [(m.start(), m.end(), term) for m in re.finditer(pattern, text.lower())]


def _remove_spans_inside_longer(spans: List[Tuple[int, int, str]]) -> List[Tuple[int, int, str]]:
    ordered = sorted(spans, key=lambda item: (-(item[1] - item[0]), item[0]))
    kept: List[Tuple[int, int, str]] = []
    for span in ordered:
        start, end, term = span
        if any(start >= ks and end <= ke and (ke - ks) > (end - start) for ks, ke, _kt in kept):
            continue
        kept.append(span)
    return sorted(kept, key=lambda item: item[0])


def _term_issue(language: str, text: str) -> Tuple[str, str]:
    if _language_code_from_label(language) != "EN":
        return "", ""
    lowered = ("" if text is None else str(text)).lower()
    for group_name, terms in TERM_CONSISTENCY_GROUPS:
        spans: List[Tuple[int, int, str]] = []
        for term in sorted(terms, key=len, reverse=True):
            spans.extend(_term_spans(lowered, term))
        spans = _remove_spans_inside_longer(spans)
        found_terms = sorted({term for _s, _e, term in spans}, key=len, reverse=True)
        # 短语包含关系（如 Album Collection 内含 Album）不算混用；只有独立命中多个术语才提示。
        if len(found_terms) >= 2:
            return group_name, ", ".join(found_terms)
    return "", ""

def _is_allowed_chinese_language(header: str, allowed_chinese: set[str]) -> bool:
    normalized = _normalize_column_key(header)
    if normalized in allowed_chinese:
        return True
    code, _name, _confidence = _language_from_header(header)
    if code in {"CN", "TW"}:
        return True
    padded = f"_{normalized}_"
    for token in allowed_chinese:
        if not token:
            continue
        if len(token) <= 3 and not re.search(r"[\u4e00-\u9fff]", token):
            if f"_{token}_" in padded:
                return True
        elif token in normalized:
            return True
    return False


def _extract_chinese_fragments(text: str) -> List[str]:
    return list(dict.fromkeys(re.findall(r"[\u4e00-\u9fff]+", text or "")))


def _extract_unallowed_chinese_fragments(text: str, allowed_terms: set[str]) -> List[str]:
    """Return real residual-CJK fragments after removing configured whitelist terms.

    The same whitelist backing the UI's "允许中文" configuration is reused for
    localized language names and project proper nouns. Removing a whitelisted
    term from the candidate text still leaves surrounding Chinese text visible,
    so a sentence such as "请选择简体中文" is not accidentally ignored.
    """
    candidate = "" if text is None else str(text)
    for term in sorted(allowed_terms, key=len, reverse=True):
        if term and CHINESE_RE.search(term):
            candidate = re.sub(re.escape(term), "", candidate, flags=re.IGNORECASE)
    return _extract_chinese_fragments(candidate)


def _extract_placeholders(text: str) -> List[str]:
    value = "" if text is None else str(text)
    found: List[Tuple[int, str]] = []
    occupied: List[Tuple[int, int]] = []

    def add_match(match) -> None:
        span = match.span()
        if any(span[0] < end and span[1] > start for start, end in occupied):
            return
        occupied.append(span)
        found.append((match.start(), match.group(0)))

    # Match the largest complete structures first so [%s1] is one placeholder,
    # instead of a printf token plus a business number. Escaped newlines are
    # formatting markers and intentionally do not participate here.
    for regex in (
        BRACKETED_PERCENT_PLACEHOLDER_RE,
        DOLLAR_PLACEHOLDER_RE,
        BRACE_PLACEHOLDER_RE,
        ANGLE_PLACEHOLDER_RE,
        HASH_PLACEHOLDER_RE,
        PRINTF_PLACEHOLDER_RE,
        POSITIONAL_PERCENT_PLACEHOLDER_RE,
    ):
        for match in regex.finditer(value):
            add_match(match)
    for match in SQUARE_PLACEHOLDER_RE.finditer(value):
        token = match.group(0)
        inner = token.strip("[]").split("=", 1)[0].strip().lower().lstrip("/")
        if inner in RICH_TAG_NAMES:
            continue
        add_match(match)
    found.sort(key=lambda item: item[0])
    return [item[1] for item in found]


def _suffix_or_word_percent_values(text: str) -> set[str]:
    value = "" if text is None else str(text)
    result = {_normalize_number_token(match.group(1).replace(",", ".")) for match in PERCENT_RE.finditer(value)}
    result.update(_normalize_number_token(match.group(1).replace(",", ".")) for match in PERCENT_WORD_RE.finditer(value))
    return result


def _placeholder_issue(source_text: str, current_text: str) -> Tuple[str, str, str, str, bool]:
    source = _extract_placeholders(source_text)
    target = _extract_placeholders(current_text)
    # In languages such as Turkish, percentages are commonly written before
    # the number (%50). If the counterpart contains the same suffix/word
    # percentage, that prefix token is numeric formatting rather than %1-style
    # business placeholder syntax.
    percentage_values = _suffix_or_word_percent_values(source_text) | _suffix_or_word_percent_values(current_text)
    if percentage_values:
        source = [token for token in source if not (re.fullmatch(r"%\d+", token) and _normalize_number_token(token[1:]) in percentage_values)]
        target = [token for token in target if not (re.fullmatch(r"%\d+", token) and _normalize_number_token(token[1:]) in percentage_values)]
    if not source and not target:
        return "", "", "", "", False
    if Counter(source) != Counter(target):
        missing = list((Counter(source) - Counter(target)).elements())
        extra = list((Counter(target) - Counter(source)).elements())
        if missing and not extra:
            return "占位符缺失", "高风险错误", f"译文缺少源文本占位符：{missing}；源={source}；译文={target}", "补齐缺失占位符，保持占位符名称和数量一致", False
        if extra and not missing:
            return "占位符新增", "高风险错误", f"译文存在源文本没有的占位符：{extra}；源={source}；译文={target}", "删除多余占位符，或确认源文本是否需要同步", False
        return "占位符不一致", "高风险错误", f"占位符名称或数量不一致：源={source}；译文={target}", "保持占位符名称、数量一致", False
    # Named placeholders are compared by complete name and quantity. A pure
    # order change is normal across languages and must not be treated as an
    # error when the same complete tokens are retained.
    return "", "", "", "", False


def _placeholder_syntax_errors(text: str) -> List[str]:
    """Return raw placeholder syntax errors for one text only."""
    value = "" if text is None else str(text)
    errors: List[str] = []
    if value.count("{") != value.count("}"):
        errors.append(f"花括号数量不匹配：{{ {value.count('{')} 个，}} {value.count('}')} 个")
    if re.search(r"\{[^{}\r\n]{81,}\}", value):
        errors.append("存在过长占位符片段")
    return errors


def _malformed_placeholder_issue(text: str, *, target: bool) -> Tuple[str, str, str, str, bool]:
    errors = _placeholder_syntax_errors(text)
    if not errors:
        return "", "", "", "", False
    issue_type = "目标文本占位符结构异常" if target else "源文本占位符结构异常"
    level = "高风险错误" if target else "源文本问题"
    return issue_type, level, "；".join(errors), "修正占位符格式，确保变量括号完整闭合", False


def _normalize_number_token(token: str) -> str:
    raw = (token or "").strip().replace(",", "")
    if raw.startswith("+"):
        raw = raw[1:]
    try:
        dec = Decimal(raw)
    except (InvalidOperation, ValueError):
        return raw.lower().replace(" ", "")
    if dec == dec.to_integral_value():
        return str(dec.quantize(Decimal(1)))
    normalized = format(dec.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return normalized or "0"


def _consume_regex_tokens(value: str, regex: re.Pattern, token_func) -> Tuple[str, List[Tuple[int, str]]]:
    tokens: List[Tuple[int, str]] = []
    chars = list(value)
    for match in regex.finditer(value):
        tokens.append((match.start(), token_func(match)))
        for idx in range(match.start(), match.end()):
            chars[idx] = " "
    return "".join(chars), tokens


def _mask_natural_quantity_words(value: str) -> str:
    """屏蔽自然语言数量词，避免把“一个/a/an/once”等当强配置数值。"""
    value = LIST_PREFIX_RE.sub(" ", value)
    value = CN_NATURAL_ONE_RE.sub(" ", value)
    value = CN_NATURAL_ARABIC_ONE_RE.sub(" ", value)
    value = EN_NATURAL_QUANTITY_RE.sub(" ", value)
    # a/an/one/single/once/each/every 默认不是强业务数值；如需严格检查可后续做成 UI 开关。
    value = EN_NATURAL_ONE_RE.sub(" ", value)
    return value


def extract_business_numbers(text: str, include_ascii_roman: bool = False) -> List[str]:
    value = "" if text is None else str(text)
    tokens: List[Tuple[int, str]] = []

    # 先屏蔽明确不是业务配置数值的自然语言数量词。
    value = _mask_natural_quantity_words(value)

    # 先提取业务复合数值，再屏蔽不应作为普通数字重复提取的结构。
    value, found = _consume_regex_tokens(value, DATE_RE, lambda m: f"date:{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}")
    tokens.extend(found)
    value, found = _consume_regex_tokens(
        value,
        CN_DATE_RE,
        lambda m: "date:{:02d}-{:02d}".format(
            _cn_number_to_int(m.group(1)) or int(m.group(1)),
            _cn_number_to_int(m.group(2)) or int(m.group(2)),
        ),
    )
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, TIME_RE, lambda m: str(int(m.group(1))) if m.group(2) == "00" else f"{int(m.group(1))}:{m.group(2)}")
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, CN_HOUR_RE, lambda m: str(int(m.group(1))))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, LV_STAGE_RE, lambda m: "lv" + _normalize_number_token(re.search(r"\d+", m.group(1)).group(0)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, CN_LEVEL_RE, lambda m: "lv" + str(_cn_number_to_int(m.group(1)) if _cn_number_to_int(m.group(1)) is not None else m.group(1)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, ARABIC_CN_LEVEL_RE, lambda m: "lv" + _normalize_number_token(m.group(1)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, THOUSAND_SEP_RE, lambda m: _normalize_thousand_token(m.group(1)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, PERCENT_RE, lambda m: _normalize_number_token(m.group(1).replace(",", ".")))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, PERCENT_WORD_RE, lambda m: _normalize_number_token(m.group(1).replace(",", ".")))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, MULTIPLIER_X_RE, lambda m: _normalize_number_token(m.group(1)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, CN_ORDINAL_RE, lambda m: str(_cn_number_to_int(m.group(1)) if _cn_number_to_int(m.group(1)) is not None else m.group(1)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, EN_ORDINAL_RE, lambda m: str(_EN_NUM.get((m.group(1) or m.group(2)).lower(), m.group(2))))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, EN_PREFIX_NUMBER_RE, lambda m: str(_EN_NUM.get(m.group(1).lower(), m.group(1))))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, ROMAN_NUMERAL_RE, lambda m: str(_ROMAN_NUM.get(m.group(1), m.group(1))))
    tokens.extend(found)
    if include_ascii_roman:
        value, found = _consume_regex_tokens(value, ASCII_ROMAN_RE, lambda m: str(_ROMAN_NUM.get(m.group(1).upper(), m.group(1))))
        tokens.extend(found)
    value, found = _consume_regex_tokens(value, CN_STRONG_NUMBER_UNIT_RE, lambda m: str(_cn_number_to_int(m.group(1)) if _cn_number_to_int(m.group(1)) is not None else m.group(1)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, EN_STRONG_WORD_NUMBER_RE, lambda m: str(_EN_NUM.get(re.match(r"[A-Za-z]+", m.group(0)).group(0).lower(), re.match(r"[A-Za-z]+", m.group(0)).group(0).lower())))
    tokens.extend(found)
    # 保留旧规则兜底，但前面已经屏蔽 one/a/once 这类自然数量词。
    value, found = _consume_regex_tokens(value, EN_WORD_NUMBER_RE, lambda m: str(_EN_NUM.get((m.group(1) or re.match(r"[A-Za-z]+", m.group(0)).group(0)).lower(), (m.group(1) or m.group(0)).lower())))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, NUMBER_WITH_UNIT_RE, lambda m: _normalize_number_token(m.group(1)))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, RANGE_RE, lambda m: _normalize_number_token(m.group(1)) + "|" + _normalize_number_token(m.group(2)))
    for pos, pair in found:
        left, right = pair.split("|", 1)
        tokens.extend([(pos, left), (pos + 1, right)])
    value, found = _consume_regex_tokens(value, LOCALIZED_DECIMAL_RE, lambda m: _normalize_number_token(f"{m.group(1)}.{m.group(2)}"))
    tokens.extend(found)
    value, found = _consume_regex_tokens(value, NEGATIVE_NUMBER_RE, lambda m: "-" + _normalize_number_token(m.group(1)))
    tokens.extend(found)

    for regex in (
        UNICODE_ESCAPE_RE, HTML_ENTITY_RE, HEX_COLOR_RE, BRACKETED_PERCENT_PLACEHOLDER_RE,
        PRINTF_PLACEHOLDER_RE, POSITIONAL_PERCENT_PLACEHOLDER_RE, DOLLAR_PLACEHOLDER_RE, HASH_PLACEHOLDER_RE,
        BRACE_PLACEHOLDER_RE, SQUARE_PLACEHOLDER_RE, ANGLE_PLACEHOLDER_RE, ANY_ANGLE_TAG_RE, ANY_BRACKET_TAG_RE,
        IDENTIFIER_NUMBER_RE,
    ):
        value = regex.sub(" ", value)
    for match in NUMBER_RE.finditer(value):
        tokens.append((match.start(), _normalize_number_token(match.group(1))))
    tokens.sort(key=lambda item: item[0])
    return [token for _pos, token in tokens]


def _number_list_text(values: Sequence[str]) -> str:
    return "[" + ", ".join(values) + "]"


def _number_diff_only_natural_one(source_numbers: Sequence[str], current_numbers: Sequence[str]) -> bool:
    sc = Counter(source_numbers)
    cc = Counter(current_numbers)
    diff = (sc - cc) + (cc - sc)
    return bool(diff) and set(diff.keys()) == {"1"}


def _number_values_only_sign_different(source_numbers: Sequence[str], current_numbers: Sequence[str]) -> bool:
    def absolute_values(values: Sequence[str]) -> List[str] | None:
        result: List[str] = []
        for value in values:
            try:
                result.append(_normalize_number_token(str(abs(Decimal(value)))))
            except (InvalidOperation, ValueError):
                return None
        return result

    source_abs = absolute_values(source_numbers)
    current_abs = absolute_values(current_numbers)
    return bool(source_abs is not None and current_abs is not None and source_abs == current_abs)


def _canonical_number_token(token: str) -> str:
    value = str(token or "")
    return value[2:] if value.startswith("lv") else value


def _date_token_parts(token: str) -> Tuple[str, str] | None:
    if not str(token).startswith("date:"):
        return None
    value = str(token)[5:]
    parts = value.split("-")
    if len(parts) == 3:
        return parts[0], f"{parts[1]}-{parts[2]}"
    if len(parts) == 2:
        return "", f"{parts[0]}-{parts[1]}"
    return None


def _date_tokens_equivalent(source_numbers: Sequence[str], current_numbers: Sequence[str]) -> bool:
    if len(source_numbers) != len(current_numbers) or not source_numbers:
        return False
    source_parts = [_date_token_parts(value) for value in source_numbers]
    current_parts = [_date_token_parts(value) for value in current_numbers]
    if any(value is None for value in source_parts + current_parts):
        return False
    return Counter(value[1] for value in source_parts if value) == Counter(value[1] for value in current_parts if value)


def _has_date_context(text: str) -> bool:
    value = "" if text is None else str(text)
    return bool(DATE_RE.search(value) or CN_DATE_RE.search(value) or re.search(r"\b(?:date|datum|fecha|data|tanggal|дата)\b", value, re.IGNORECASE))


def _plausible_excel_serials(values: Sequence[str]) -> bool:
    if not values:
        return False
    try:
        return all(20000 <= int(Decimal(value)) <= 80000 for value in values)
    except (InvalidOperation, ValueError):
        return False


def _number_consistency_error(source_text: str, current_text: str) -> Tuple[str, str, str, bool]:
    source_percent_values = _suffix_or_word_percent_values(source_text)
    current_percent_values = _suffix_or_word_percent_values(current_text)

    def normalize_prefix_percent(text: str, counterpart_values: set[str]) -> str:
        value = "" if text is None else str(text)
        if not counterpart_values:
            return value
        def replace_match(match) -> str:
            number = _normalize_number_token(match.group(1).replace(",", "."))
            return f"{match.group(1)}%" if number in counterpart_values else match.group(0)
        return re.sub(r"%(\d+(?:[\.,]\d+)?)", replace_match, value)

    source_for_numbers = normalize_prefix_percent(source_text, current_percent_values)
    current_for_numbers = normalize_prefix_percent(current_text, source_percent_values)
    include_ascii_roman = bool(ROMAN_NUMERAL_RE.search(source_text or "") or ROMAN_NUMERAL_RE.search(current_text or ""))
    source_numbers = extract_business_numbers(source_for_numbers, include_ascii_roman=include_ascii_roman)
    current_numbers = extract_business_numbers(current_for_numbers, include_ascii_roman=include_ascii_roman)
    if not source_numbers and not current_numbers:
        return "", "", "", False
    if _date_tokens_equivalent(source_numbers, current_numbers):
        return "", "", "", False
    source_canonical = [_canonical_number_token(value) for value in source_numbers]
    current_canonical = [_canonical_number_token(value) for value in current_numbers]
    if Counter(source_canonical) == Counter(current_canonical):
        # 等级前缀和跨语言语序均属于正常本地化表现。
        return "", "", "", False
    if (_has_date_context(source_text) or _has_date_context(current_text)) and (
        _plausible_excel_serials(source_canonical) or _plausible_excel_serials(current_canonical)
    ):
        return "数值本地化转换待确认", "需人工确认", f"文本处于日期上下文，检测到疑似 Excel 日期序列：源={_number_list_text(source_numbers)}；译文={_number_list_text(current_numbers)}", True
    if len(source_canonical) != len(current_canonical):
        # 差异仅由“一/one/a/单个/一层/1 stack”等自然数量词造成时，默认降噪，不计入错误。
        if _number_diff_only_natural_one(source_canonical, current_canonical):
            return "", "", "", False
        return "数值数量不一致", "疑似问题", f"源文本包含 {len(source_numbers)} 个业务数值 {_number_list_text(source_numbers)}，当前译文包含 {len(current_numbers)} 个业务数值 {_number_list_text(current_numbers)}；可能涉及中文数字、单位省略或本地化表达转换", True
    if _number_diff_only_natural_one(source_canonical, current_canonical):
        return "", "", "", False
    if _number_values_only_sign_different(source_canonical, current_canonical):
        return "数值表现形式差异", "疑似问题", f"源文本与译文仅数值正负号表现不同：源={_number_list_text(source_numbers)}；译文={_number_list_text(current_numbers)}", True
    if re.search(r"\b(?:GMT|UTC)\b", source_text or "", re.IGNORECASE) and re.search(r"\b(?:GMT|UTC)\b", current_text or "", re.IGNORECASE):
        return "数值本地化转换待确认", "疑似问题", f"文本包含时区信息，数值可能经过本地化转换：源={_number_list_text(source_numbers)}；译文={_number_list_text(current_numbers)}", True
    return "数值内容不一致", "高风险错误", f"源文本数值为 {_number_list_text(source_numbers)}，当前译文数值为 {_number_list_text(current_numbers)}", False


def _rich_tag_items(text: str) -> List[Tuple[int, str, str, str]]:
    value = "" if text is None else str(text)
    items: List[Tuple[int, str, str, str]] = []
    for match in ANGLE_TAG_RE.finditer(value):
        slash, name, attrs, self_close = match.group(1), match.group(2).lower(), match.group(3) or "", match.group(4)
        if name not in RICH_TAG_NAMES:
            continue
        if slash and attrs.strip():
            continue
        if name in {"b", "i", "u", "br"} and attrs.strip():
            continue
        kind = "self" if self_close or name in SELF_CLOSING_TAG_NAMES else ("close" if slash else "open")
        items.append((match.start(), kind, name, attrs.strip()))
    for match in BRACKET_TAG_RE.finditer(value):
        slash, name, param = match.group(1), match.group(2).lower(), match.group(3) or ""
        raw_tag = match.group(0)
        if name.startswith("color") and name != "color":
            # 兼容 [colorFFFFFF] 这类少了空格/等号的游戏标签，按 color 标签处理，
            # 参数保留在 attrs 中用于后续数量对比和人工定位。
            param = name[len("color"):] or param
            name = "color"
        if name not in RICH_TAG_NAMES:
            continue
        # Registered short tags are only valid as complete structures. In
        # particular, ordinary bracketed prose such as [I Scream] must not be
        # truncated to an [i] tag with a fake attribute named "Scream".
        if slash and param.strip():
            continue
        if name in {"b", "i", "u", "br"} and param.strip():
            continue
        if name in {"url", "link"} and param.strip() and "=" not in raw_tag:
            continue
        if slash:
            kind = "close"
        elif name in SELF_CLOSING_TAG_NAMES:
            kind = "self"
        else:
            kind = "open"
        items.append((match.start(), kind, name, param.strip()))
    items.sort(key=lambda item: item[0])
    return items


def _is_pure_tag_fragment(text: str) -> bool:
    value = ("" if text is None else str(text)).strip()
    if not value:
        return False
    if value in {"\\n", "\n"}:
        return True
    items = _rich_tag_items(value)
    if items and all(kind == "self" for _pos, kind, _name, _attrs in items):
        return True
    return False


def _normalize_tag_attrs(attrs: str) -> str:
    return re.sub(r"\s+", " ", (attrs or "").strip()).lower()


LONG_TAG_START_RE = re.compile(r"[\[<]\s*/?\s*(?:color|size|font|url|ref|link)\b", re.IGNORECASE)
SHORT_TAG_START_RE = re.compile(r"[\[<]\s*/?\s*(?:b|i|u|br)\s*(?=[\]>])", re.IGNORECASE)
INCOMPLETE_SHORT_TAG_RE = re.compile(r"[\[<]\s*/?\s*(?:b|i|u|br)\s*(?=$|[\r\n])", re.IGNORECASE)


def _tag_bracket_error(text: str) -> str:
    value = "" if text is None else str(text)
    for regex in (LONG_TAG_START_RE, SHORT_TAG_START_RE, INCOMPLETE_SHORT_TAG_RE):
        for match in regex.finditer(value):
            opener = value[match.start()]
            closer = "]" if opener == "[" else ">"
            end = value.find(closer, match.end())
            next_same = value.find(opener, match.end())
            newline = value.find("\n", match.end())
            boundaries = [pos for pos in (next_same, newline) if pos >= 0]
            boundary = min(boundaries) if boundaries else len(value)
            if end < 0 or end > boundary:
                fragment = value[match.start():min(boundary, match.start() + 80)]
                return f"标签括号不完整：{fragment}"
    return ""


def _tag_error(text: str) -> str:
    if _is_pure_tag_fragment(text):
        return ""
    bracket_error = _tag_bracket_error(text)
    if bracket_error:
        return bracket_error
    stack: List[str] = []
    for _pos, kind, name, _attrs in _rich_tag_items(text):
        if kind == "self":
            continue
        if kind == "open":
            if name in STYLE_TAG_NAMES and name in stack:
                return f"同类型标签异常嵌套：{name}"
            stack.append(name)
        else:
            if not stack:
                return f"标签 </{name}> 缺少开始标签"
            last = stack.pop()
            if last != name:
                return f"标签闭合顺序异常：期望 </{last}>，实际 </{name}>"
    if stack:
        return f"标签未闭合：{', '.join(stack)}"
    return ""


def _tag_compare_issue(source_text: str, current_text: str) -> Tuple[str, str, str, str, bool]:
    if _is_pure_tag_fragment(source_text) or _is_pure_tag_fragment(current_text):
        return "纯标签片段", "可忽略/规则兼容", "检测到纯标签行，已按配置型文本片段兼容处理", "如项目不允许纯标签拆分，可人工确认该 ID", True
    source_items = [(kind, name, _normalize_tag_attrs(attrs)) for _pos, kind, name, attrs in _rich_tag_items(source_text)]
    current_items = [(kind, name, _normalize_tag_attrs(attrs)) for _pos, kind, name, attrs in _rich_tag_items(current_text)]
    source_structure = [(kind, name) for kind, name, _attrs in source_items]
    current_structure = [(kind, name) for kind, name, _attrs in current_items]
    if source_structure and not current_structure:
        return "标签缺失", "高风险错误", f"源文本包含富文本标签 {source_structure}，译文未保留", "补齐对应富文本标签", False
    if current_structure and not source_structure:
        if any(name in FUNCTIONAL_TAG_NAMES for _kind, name in current_structure):
            return "功能标签新增", "高风险错误", f"译文新增源文本没有的功能型标签 {current_structure}", "确认链接、引用或跳转标签是否应同步到源文本", False
        return "标签新增", "需人工确认", f"译文新增源文本没有的表现型标签 {current_structure}", "人工确认是否为目标语言额外样式需求", True
    if Counter(source_structure) != Counter(current_structure):
        missing = list((Counter(source_structure) - Counter(current_structure)).elements())
        extra = list((Counter(current_structure) - Counter(source_structure)).elements())
        if extra and not missing and not any(name in FUNCTIONAL_TAG_NAMES for _kind, name in extra):
            return "标签新增", "需人工确认", f"译文新增表现型标签：{extra}", "人工确认是否为目标语言额外样式需求", True
        return "标签结构不一致", "高风险错误", f"标签名称或数量不一致；缺少={missing}；新增={extra}", "保持标签名称、数量和闭合结构一致", False
    if source_structure != current_structure:
        return "标签顺序变化", "需人工确认", f"标签结构相同但顺序变化：源={source_structure}；译文={current_structure}", "人工确认标签嵌套和样式范围是否正确", True
    source_attrs = [(name, attrs) for _kind, name, attrs in source_items]
    current_attrs = [(name, attrs) for _kind, name, attrs in current_items]
    if source_attrs != current_attrs:
        changed_names = {name for (name, s_attr), (_other, t_attr) in zip(source_attrs, current_attrs) if s_attr != t_attr}
        if changed_names & FUNCTIONAL_TAG_NAMES:
            return "功能标签属性差异", "高风险错误", f"功能型标签属性发生变化：源={source_attrs}；译文={current_attrs}", "确认链接、引用或跳转目标是否必须保持一致", False
        return "标签属性差异", "需人工确认", f"标签结构一致，仅表现属性变化：源={source_attrs}；译文={current_attrs}", "人工确认颜色、字号或字体是否为目标语言排版需求", True
    return "", "", "", "", False


def _symbol_tokens(raw: str) -> List[str]:
    return [token for token in _split_config_values(raw) if token]


def _contains_symbol_equivalent(text: str, token: str) -> bool:
    value = "" if text is None else str(text)
    if token == "()":
        return ("(" in value or "（" in value) and (")" in value or "）" in value)
    if token == "[]":
        return "[" in value and "]" in value
    if token == "{}":
        return "{" in value and "}" in value
    if token == "<>":
        return "<" in value and ">" in value
    equivalents = NATURAL_PUNCT_EQUIVALENTS.get(token)
    if equivalents:
        return any(eq in value for eq in equivalents)
    return token in value


def _missing_critical_symbols(source_text: str, current_text: str, symbols: Sequence[str]) -> Tuple[List[str], List[str]]:
    missing: List[str] = []
    compatible: List[str] = []
    newline_count_differs = _newline_marker_count(source_text) != _newline_marker_count(current_text)
    source_tags = _rich_tag_items(source_text)
    source_placeholders = _extract_placeholders(source_text)
    for symbol in symbols:
        # Newline markers have their own format rule. Do not duplicate them as
        # business symbols (and do not separately report the backing slash).
        if symbol == "\\n" or (symbol == "\\" and newline_count_differs):
            continue
        if source_tags and symbol in {"[]", "<>"}:
            continue
        if symbol == "%" and any("%" in token for token in source_placeholders):
            continue
        if symbol == "[]" and any(token.startswith("[") and token.endswith("]") for token in source_placeholders):
            continue
        if symbol == "{}" and any(token.startswith("{") and token.endswith("}") for token in source_placeholders):
            continue
        if symbol == "<>" and any(token.startswith("<") and token.endswith(">") for token in source_placeholders):
            continue
        # Brackets are meaningful only when a registered tag/placeholder has
        # already claimed them. Ordinary prose such as [I Scream] or [活动]
        # must not participate in symbol consistency checks.
        if symbol in {"[]", "{}", "<>"}:
            continue
        if symbol == "/":
            source_has_context = bool(
                re.search(r"(?:https?://|ftp://|[A-Za-z]:[\\/]|\d+\s*/\s*\d+|\b[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+\b)", source_text or "", re.IGNORECASE)
            )
            if not source_has_context:
                continue
        if symbol == "\\":
            source_has_context = bool(re.search(r"(?:[A-Za-z]:\\|\\[A-Za-z0-9_.-]+\\)", source_text or ""))
            if not source_has_context:
                continue
        if symbol == "%" and PERCENT_RE.search(source_text or ""):
            source_percent_values = Counter(_normalize_number_token(m.group(1).replace(",", ".")) for m in PERCENT_RE.finditer(source_text or ""))
            target_percent_values = Counter(_normalize_number_token(m.group(1).replace(",", ".")) for m in PERCENT_RE.finditer(current_text or ""))
            target_percent_values.update(_normalize_number_token(m.group(1).replace(",", ".")) for m in PERCENT_WORD_RE.finditer(current_text or ""))
            if source_percent_values == target_percent_values:
                continue
        if not _contains_symbol_equivalent(source_text, symbol):
            continue
        if _contains_symbol_equivalent(current_text, symbol):
            if symbol in NATURAL_PUNCT_EQUIVALENTS:
                compatible.append(symbol)
            continue
        if symbol in CRITICAL_SYMBOLS:
            missing.append(symbol)
        # 普通自然语言标点缺失不作为问题，避免“设备：/Device:”类误报。
    return missing, compatible


def _empty(text: str) -> bool:
    return not INVISIBLE_TEXT_RE.sub("", "" if text is None else str(text))


def _infer_issue_level(issue_type: str, explicit: str = "") -> str:
    if explicit:
        return explicit
    if issue_type in {"空翻译", "残留中文", "占位符缺失", "占位符新增", "占位符不一致", "标签缺失", "标签闭合异常", "数值内容不一致"}:
        return "高风险错误"
    if "顺序" in issue_type or "变化" in issue_type or issue_type in {"标签新增", "待确认", "纯标签片段", "JP 日文汉字疑似残留", "术语一致性疑似问题", "英文疑似拼写错误", "不同源文本译文重复"}:
        return "疑似问题"
    if issue_type in {"格式差异但语义一致", "纯标签片段"}:
        return "提醒"
    return "普通问题"


def _issue_requires_manual_confirm(level: str, issue_type: str) -> bool:
    return level in {"疑似问题", "需人工确认", "提醒", "可忽略/规则兼容"} or issue_type.startswith("疑似") or "疑似" in issue_type


def _issue_count_in_error_stats(level: str) -> bool:
    return level in {"高风险错误", "普通问题"}


def _issue_confidence(level: str) -> str:
    if level in {"高风险错误", "源文本问题"}:
        return "高"
    if level == "普通问题":
        return "中"
    return "低"


def _result_category(level: str, issue_type: str, rule_name: str = "") -> str:
    if level == "源文本问题" or "源文本" in issue_type:
        return "源文本问题"
    if level == "语言覆盖范围" or rule_name == "整列空翻译":
        return "语言覆盖范围"
    if level in {"可忽略/规则兼容", "忽略项或白名单"}:
        return "忽略项或白名单"
    if level in {"疑似问题", "需人工确认", "提醒"}:
        return "需人工确认"
    return "明确问题"


def _make_issue(index: int, issue_type: str, text_id: str, language: str, row_number: int | str, source_text: str, current_text: str, remark: str, suggestion: str, sheet_name: str = "", issue_level: str = "", rule_name: str = "", source_language: str = "", is_false_positive: bool = False, module: str = "", source_file: str = "", target_file: str = "", result_category: str = "", count_in_error_stats: bool | None = None, affected_languages: str = "", coverage_total: int = 0, coverage_filled: int = 0, coverage_blank: int = 0, coverage_status: str = "") -> LocalizationIssue:
    level = _infer_issue_level(issue_type, issue_level)
    manual_confirm = _issue_requires_manual_confirm(level, issue_type)
    category = result_category or _result_category(level, issue_type, rule_name)
    count_in_stats = _issue_count_in_error_stats(level) if count_in_error_stats is None else bool(count_in_error_stats)
    if category != "明确问题":
        count_in_stats = False
    degraded = is_false_positive or manual_confirm or level == "提醒"
    return LocalizationIssue(
        index=index,
        issue_type=issue_type,
        item_id=_format_identifier(text_id),
        language="" if language is None else str(language),
        sheet_name="" if sheet_name is None else str(sheet_name),
        row_number="" if row_number is None else str(row_number),
        source_text="" if source_text is None else str(source_text),
        current_text="" if current_text is None else str(current_text),
        remark=remark,
        suggestion=suggestion,
        issue_level=level,
        rule_name=rule_name or issue_type,
        source_language=source_language,
        is_false_positive=degraded,
        requires_manual_confirm=manual_confirm,
        count_in_error_stats=count_in_stats,
        rule_confidence=_issue_confidence(level),
        is_degraded=degraded,
        module="" if module is None else str(module),
        source_file="" if source_file is None else str(source_file),
        target_file="" if target_file is None else str(target_file),
        result_category=category,
        affected_languages=affected_languages or ("" if language is None else str(language)),
        coverage_total=max(0, int(coverage_total or 0)),
        coverage_filled=max(0, int(coverage_filled or 0)),
        coverage_blank=max(0, int(coverage_blank or 0)),
        coverage_status="" if coverage_status is None else str(coverage_status),
    )


def _expected_language_codes(options: LocalizationCheckOptions) -> List[str]:
    values = _split_config_values(options.expected_languages_text)
    return [value.strip().upper() for value in values if value.strip()]


def _enabled_language_codes(options: LocalizationCheckOptions) -> List[str]:
    return [value.strip().upper() for value in (options.enabled_language_columns or []) if value.strip()]




def _language_code_from_sheet_name(sheet_name: str) -> str:
    norm = _normalize_column_key(sheet_name).replace("_", "")
    for code, (_name, aliases) in LANGUAGE_META.items():
        if norm == code.lower():
            return code
        for alias in aliases:
            alias_norm = _normalize_column_key(alias).replace("_", "")
            if alias_norm and norm == alias_norm:
                return code
    return ""


def _detect_localization_structure_fast(path: str) -> str:
    """轻量结构识别：自动模式只通过 xlsx 元数据判断是否为多 Sheet 语言页。

    不使用 openpyxl 预读内容，避免大文件自动识别阶段和正式解析阶段重复加载造成卡顿。
    """
    file_path = Path(path)
    ext = file_path.suffix.lower()
    if ext == ".xlsx":
        try:
            with zipfile.ZipFile(file_path) as zf:
                xml = zf.read("xl/workbook.xml")
            root = ET.fromstring(xml)
            ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            names = [elem.attrib.get("name", "") for elem in root.findall(".//main:sheet", ns)]
            language_sheets = [name for name in names if _language_code_from_sheet_name(name)]
            if len(language_sheets) >= 2:
                return MULTI_SHEET_STRUCTURE
            return SINGLE_SHEET_STRUCTURE if names else UNKNOWN_STRUCTURE
        except Exception:
            return SINGLE_SHEET_STRUCTURE
    if ext == ".csv":
        return SINGLE_SHEET_STRUCTURE
    return UNKNOWN_STRUCTURE

def _resolve_localization_structure(
    path: str,
    options: LocalizationCheckOptions,
    progress_callback: ProgressCallback | None = None,
) -> str:
    mode = normalize_localization_mode(options.table_structure or "自动识别")
    if mode == MULTI_FILE_DICTIONARY_STRUCTURE:
        return MULTI_FILE_DICTIONARY_STRUCTURE
    if mode == MULTI_TABLE_SHEETS_STRUCTURE:
        return MULTI_TABLE_SHEETS_STRUCTURE
    if mode == MULTI_SHEET_STRUCTURE:
        return MULTI_SHEET_STRUCTURE
    if mode == SINGLE_SHEET_STRUCTURE:
        return SINGLE_SHEET_STRUCTURE
    # 自动模式先通过 workbook 元数据判断多 Sheet 语言页；否则再判断是否为
    # “一个工作簿多个版本 Sheet、每个 Sheet 都是多语言列”的整合表。
    try:
        if len(detect_dictionary_language_files(path)[0]) >= 2:
            return MULTI_FILE_DICTIONARY_STRUCTURE
    except Exception:
        pass
    structure = _detect_localization_structure_fast(path)
    if structure == SINGLE_SHEET_STRUCTURE and not str(options.sheet_name or "").strip():
        try:
            candidates = list_localization_table_sheets(path, options, progress_callback)
            if len(candidates) >= 2:
                options.precomputed_multi_table_sheets = candidates
                return MULTI_TABLE_SHEETS_STRUCTURE
        except Exception:
            pass
    return structure


def _merge_localization_columns(columns_by_header: Dict[str, LocalizationColumn], columns: Sequence[LocalizationColumn]) -> None:
    for column in columns:
        key = (column.code or column.header or column.name).strip()
        if not key:
            continue
        if key not in columns_by_header:
            columns_by_header[key] = LocalizationColumn(
                name=column.name,
                header=column.header,
                index=column.index,
                display_name=column.display_name,
                confidence=column.confidence,
                excluded_reason=column.excluded_reason,
                code=column.code,
                raw_header=column.raw_header,
            )


def _add_aggregated_empty_language_column_issues(
    issues: List[LocalizationIssue],
    index: int,
    table: LocalizationTable,
    target_columns: Sequence[LocalizationColumn],
    source_col: int,
    source_language: str,
    options: LocalizationCheckOptions,
    records_to_check,
) -> Tuple[int, set[int]]:
    aggregated_empty_columns: set[int] = set()
    if not (options.aggregate_empty_language_columns and options.check_empty and records_to_check):
        return index, aggregated_empty_columns
    effective_records = [record for record in records_to_check if not _empty(record.values.get(source_col, ""))]
    if not effective_records:
        return index, aggregated_empty_columns
    for column in target_columns:
        col_index = column.index
        if col_index == table.key_col_index or col_index == source_col:
            continue
        filled = sum(1 for record in effective_records if not _empty(record.values.get(col_index, "")))
        blank = len(effective_records) - filled
        header = _localization_column_label(column, table.headers[col_index] if col_index < len(table.headers) else f"列{col_index + 1}")
        status = "参与逐行检查" if filled else "整列为空，未逐行判为空翻译"
        remark = (
            f"{table.sheet_name} 的 {header}：有效数据 {len(effective_records)} 条，已填写 {filled} 条，空白 {blank} 条"
        )
        issues.append(_make_issue(
            index,
            "语言覆盖状态",
            "",
            header,
            "整列",
            "",
            "",
            remark,
            "整列为空时按未配置语言处理；活跃语言列中的零散空白继续逐行检查",
            table.sheet_name,
            "语言覆盖范围",
            "整列空翻译",
            source_language,
            result_category="语言覆盖范围",
            count_in_error_stats=False,
            coverage_total=len(effective_records),
            coverage_filled=filled,
            coverage_blank=blank,
            coverage_status=status,
        ))
        index += 1
        if not filled:
            aggregated_empty_columns.add(col_index)
    return index, aggregated_empty_columns


def _check_single_table_quality(
    table: LocalizationTable,
    detected_columns: Sequence[LocalizationColumn],
    options: LocalizationCheckOptions,
    stop_event=None,
    progress_callback: ProgressCallback | None = None,
    progress_start: float = 0.15,
    progress_end: float = 0.95,
) -> List[LocalizationIssue]:
    enabled_headers = [header.strip() for header in (options.enabled_language_columns or []) if header.strip()]
    enabled_columns = _resolve_enabled_language_columns(table, detected_columns, enabled_headers)

    if not enabled_columns:
        raise FileReadError(_build_no_language_columns_error(table, detected_columns))

    detected_by_index = {column.index: column for column in detected_columns}
    source_col = _find_detected_language_index(detected_columns, options.source_language_column)
    if source_col < 0:
        source_col = _find_header_index(table.headers, options.source_language_column)
    if source_col < 0:
        auto_source = auto_detect_source_language(table.headers, [_localization_column_label(column) for column in enabled_columns])
        source_col = _find_detected_language_index(detected_columns, auto_source)
        if source_col < 0:
            source_col = _find_header_index(table.headers, auto_source)
    if source_col < 0:
        raise FileReadError("未识别到源语言列，请手动选择源语言列")

    source_column = detected_by_index.get(source_col)
    source_language = _localization_column_label(
        source_column,
        table.display_headers[source_col] if source_col < len(table.display_headers) else (table.headers[source_col] if source_col < len(table.headers) else "源语言"),
    )
    target_columns = [column for column in enabled_columns if column.index != source_col]
    enabled_indexes = [column.index for column in target_columns]
    if not enabled_indexes:
        raise FileReadError("未识别到目标语言列：当前只识别到源语言列，请确认语言列映射或手动勾选目标语言")
    english_col = next((column.index for column in detected_columns if (column.code or column.header).upper() == "EN"), -1)
    game_text_mode = _is_repeated_text_program_table(table)
    allowed_chinese = {_normalize_column_key(value) for value in _split_config_values(options.allowed_chinese_languages_text)}
    ignore_ids = {value for value in _split_config_values(options.ignore_ids_text)}
    symbols = _symbol_tokens(options.symbols_text)
    max_length = safe_int(options.max_length, 0) if options.max_length else 0
    try:
        length_ratio = float(options.length_ratio or 0)
    except Exception:
        length_ratio = 0.0

    issues: List[LocalizationIssue] = []
    index = 1

    for diag in table.diagnostics:
        if "记录数为 0" in diag or "语言列较少" in diag:
            issues.append(_make_issue(index, "解析诊断", "", "", "", "", "", diag, "检查表头行、数据起始行、主键列和语言列识别结果", table.sheet_name, "疑似问题", "解析诊断", source_language, True))
            index += 1

    if options.check_invalid_id:
        for row_number in table.invalid_rows:
            issues.append(_make_issue(index, "无效 ID", "", "", row_number, "", "", "主键 ID 为空但该行存在内容", "补充 ID 或删除无效行", table.sheet_name, "普通问题", "主键", source_language))
            index += 1

    if options.check_duplicate_id:
        for text_id, rows_ in sorted(table.duplicate_ids.items()):
            issues.append(_make_issue(index, "重复 ID", text_id, "", ",".join(map(str, rows_)), "", "", "同一个 ID 在文件中出现多次", "确认是否需要保留唯一配置", table.sheet_name, "普通问题", "重复 ID", source_language))
            index += 1

    records_to_check = table.row_records or list(table.records.values())
    index, aggregated_empty_columns = _add_aggregated_empty_language_column_issues(
        issues, index, table, target_columns, source_col, source_language, options, records_to_check
    )
    total_records = len(records_to_check)
    _emit_progress(
        progress_callback,
        "check_records",
        f"正在检查 {table.sheet_name}：0/{total_records} 条记录",
        current=0,
        total=total_records,
        ratio=progress_start,
        sheet_name=table.sheet_name,
    )
    progress_span = max(progress_end - progress_start, 0.0)
    for position, record in enumerate(records_to_check, start=1):
        text_id = record.text_id
        if stop_event is not None and getattr(stop_event, "is_set", lambda: False)():
            _emit_progress(
                progress_callback,
                "stopping",
                "正在停止多语言检查…",
                current=position,
                total=total_records,
                ratio=progress_start + (position / max(total_records, 1)) * progress_span,
                sheet_name=table.sheet_name,
            )
            break
        if position == 1 or position == total_records or position % 200 == 0:
            _emit_progress(
                progress_callback,
                "check_records",
                f"正在检查 {table.sheet_name}：{position}/{total_records} 条记录",
                current=position,
                total=total_records,
                ratio=progress_start + (position / max(total_records, 1)) * progress_span,
                sheet_name=table.sheet_name,
            )
        if text_id in ignore_ids:
            continue
        source_text = record.values.get(source_col, "")
        if options.check_empty and _empty(source_text):
            issues.append(_make_issue(
                index,
                f"{source_language}源文本为空",
                text_id,
                source_language,
                record.row_number,
                source_text,
                source_text,
                f"ID 已存在，但 {source_language} 源文本为空或仅包含空格/换行/不可见字符",
                "先补充或确认源语言文本，再进行目标语言质量比对",
                table.sheet_name,
                "源文本问题",
                "源文本空文本",
                source_language,
            ))
            index += 1
            continue
        index = _add_source_self_format_issues(
            issues,
            index,
            text_id=text_id,
            language=source_language,
            sheet_name=table.sheet_name,
            row_number=record.row_number,
            source_text=source_text,
            options=options,
        )
        source_tag_invalid = bool(options.check_tags and TAG_SIGNAL_RE.search(source_text or "") and _tag_error(source_text))
        english_text = record.values.get(english_col, "") if english_col >= 0 else ""
        for column in target_columns:
            col_index = column.index
            if col_index == table.key_col_index:
                continue
            if col_index in aggregated_empty_columns:
                continue
            header = _localization_column_label(column, table.headers[col_index] if col_index < len(table.headers) else f"列{col_index + 1}")
            current_text = record.values.get(col_index, "")
            index = _add_text_quality_issues(
                issues, index,
                text_id=text_id,
                language=header,
                sheet_name=table.sheet_name,
                row_number=record.row_number,
                source_text=source_text,
                current_text=current_text,
                source_language=source_language,
                allowed_chinese=allowed_chinese,
                symbols=symbols,
                options=options,
                enable_layout_heuristics=game_text_mode,
                source_tag_invalid=source_tag_invalid,
            )
            if game_text_mode and options.check_terms and english_text:
                index = _add_english_copy_issue(
                    issues,
                    index,
                    text_id=text_id,
                    language=header,
                    sheet_name=table.sheet_name,
                    row_number=record.row_number,
                    source_text=source_text,
                    current_text=current_text,
                    english_text=english_text,
                    source_language=source_language,
                )
            if options.check_length:
                text_len = len(current_text or "")
                source_len = len(source_text or "")
                if max_length > 0 and text_len > max_length:
                    issues.append(_make_issue(index, "文本过长", text_id, header, record.row_number, source_text, current_text, f"当前长度 {text_len}，超过最大长度 {max_length}", "精简文本或确认 UI 容量", table.sheet_name, "普通问题", "长度", source_language))
                    index += 1
                if length_ratio > 0 and source_len > 0 and text_len > source_len * length_ratio:
                    issues.append(_make_issue(index, "文本过长", text_id, header, record.row_number, source_text, current_text, f"当前长度 {text_len}，超过源文本长度 {source_len} 的 {length_ratio:g} 倍", "检查该语言是否可能超框", table.sheet_name, "普通问题", "长度", source_language))
                    index += 1

    if options.check_duplicate_translation:
        index = _add_duplicate_translation_issues_for_single_sheet(issues, index, table, target_columns, source_col, source_language)
    if game_text_mode and options.check_terms:
        index = _add_same_source_translation_variant_issues(issues, index, table, target_columns, source_col, source_language)

    return _reindex(issues)


def _check_multi_table_sheets_quality(
    path: str,
    options: LocalizationCheckOptions,
    stop_event=None,
    progress_callback: ProgressCallback | None = None,
) -> Tuple[List[LocalizationIssue], LocalizationTable, List[LocalizationColumn]]:
    candidates = list(options.precomputed_multi_table_sheets or [])
    if not candidates:
        candidates = list_localization_table_sheets(path, options, progress_callback)
    if not candidates:
        raise FileReadError("未识别到可批量检查的 Sheet：请确认每个版本 Sheet 都包含 Text ID、CN 原文和至少 2 个语言翻译列")

    issues: List[LocalizationIssue] = []
    columns_by_header: Dict[str, LocalizationColumn] = {}
    summary_records: Dict[str, LocalizationRow] = {}
    diagnostics = [
        f"多 Sheet 翻译表模式：识别到 {len(candidates)} 个可检查 Sheet",
        "批量大表模式已自动跳过低收益高成本规则：拼写疑似、术语一致、不同源文本译文重复；保留空翻译、残留中文、占位符、标签、特殊符号和数值检查。",
    ]
    checked_count = 0
    skipped_count = 0

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FileReadError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    _emit_progress(progress_callback, "open_workbook", "正在打开工作簿准备批量检查…", ratio=0.13)
    wb = load_workbook(Path(path), data_only=True, read_only=True)
    try:
        total_sheets = len(candidates)
        for sheet_position, (sheet_name, score, reason) in enumerate(candidates, start=1):
            if stop_event is not None and getattr(stop_event, "is_set", lambda: False)():
                diagnostics.append("用户已请求停止，批量检查提前结束")
                _emit_progress(progress_callback, "stopping", "正在停止多语言检查…", current=sheet_position, total=total_sheets, ratio=0.98, sheet_name=sheet_name)
                break
            sheet_start = 0.14 + ((sheet_position - 1) / max(total_sheets, 1)) * 0.82
            sheet_end = 0.14 + (sheet_position / max(total_sheets, 1)) * 0.82
            _emit_progress(
                progress_callback,
                "read_sheet",
                f"正在读取 Sheet {sheet_position}/{total_sheets}：{sheet_name}",
                current=sheet_position,
                total=total_sheets,
                ratio=sheet_start,
                sheet_name=sheet_name,
            )
            local_options = replace(
                options,
                sheet_name=sheet_name,
                table_structure="单 Sheet 多语言列",
                check_spelling=False,
                check_terms=False,
                check_duplicate_translation=False,
                aggregate_empty_language_columns=True,
            )
            try:
                ws = wb[sheet_name]
                try:
                    ws.reset_dimensions()
                except Exception:
                    pass
                rows = _worksheet_rows_as_text(ws)
                table = _localization_table_from_rows(local_options, path, rows, sheet_name, [f"批量模式选择 Sheet：{sheet_name}（评分 {score}，{reason}）"])
                columns = detect_language_columns(table.headers, table.key_col_index, table.display_headers, table.type_headers)
                _emit_progress(
                    progress_callback,
                    "check_sheet",
                    f"正在检查 Sheet {sheet_position}/{total_sheets}：{sheet_name}（记录 {len(table.row_records or table.records)} 条）",
                    current=sheet_position,
                    total=total_sheets,
                    ratio=sheet_start,
                    sheet_name=sheet_name,
                )
                sheet_issues = _check_single_table_quality(
                    table,
                    columns,
                    local_options,
                    stop_event,
                    progress_callback,
                    sheet_start,
                    sheet_end,
                )
            except Exception as exc:
                skipped_count += 1
                issues.append(_make_issue(
                    len(issues) + 1,
                    "Sheet 解析失败",
                    "",
                    "",
                    "",
                    "",
                    "",
                    f"{sheet_name}：{exc}",
                    "检查该 Sheet 的表头行、数据起始行、主键列和语言列结构",
                    sheet_name,
                    "普通问题",
                    "结构识别",
                ))
                continue
            checked_count += 1
            diagnostics.append(f"{sheet_name}：评分 {score}，{reason}，记录 {len(table.row_records or table.records)} 条，语言列 {len(columns)} 个")
            issues.extend(sheet_issues)
            _merge_localization_columns(columns_by_header, columns)
            for record in table.row_records or list(table.records.values()):
                summary_key = f"{sheet_name}!{record.text_id}#{record.row_number}"
                summary_records[summary_key] = record
            _emit_progress(
                progress_callback,
                "sheet_done",
                f"已完成 Sheet {sheet_position}/{total_sheets}：{sheet_name}",
                current=sheet_position,
                total=total_sheets,
                ratio=sheet_end,
                sheet_name=sheet_name,
            )
    finally:
        wb.close()

    if checked_count == 0:
        raise FileReadError("批量检查未成功读取任何 Sheet，请检查工作簿结构")

    summary_table = LocalizationTable(
        path=path,
        sheet_name=f"全部 Sheet（已检查 {checked_count} 个）",
        headers=[],
        key_col_index=-1,
        records=summary_records,
        duplicate_ids={},
        invalid_rows=[],
        header_row=safe_int(options.header_row, 1),
        data_start_row=safe_int(options.data_start_row, 2),
        structure_type="多 Sheet 翻译表",
        diagnostics=diagnostics + ([f"跳过/失败 Sheet：{skipped_count} 个"] if skipped_count else []),
    )
    columns = sorted(
        columns_by_header.values(),
        key=lambda col: LANGUAGE_ORDER.index(col.code or col.header) if (col.code or col.header) in LANGUAGE_ORDER else 999,
    )
    return _reindex(issues), summary_table, columns


def _choose_dictionary_source(files: Sequence[DictionaryLanguageFile], options: LocalizationCheckOptions) -> str:
    requested = str(options.source_language_column or "").strip()
    available = {info.code for info in files}
    if requested in available:
        return requested
    normalized = _dictionary_language_code_from_name(requested)
    if normalized in available:
        return normalized
    for candidate in ("CN", "EN"):
        if candidate in available:
            return candidate
    return files[0].code if files else ""


def _check_dictionary_files_quality(
    path: str,
    options: LocalizationCheckOptions,
    stop_event=None,
    progress_callback: ProgressCallback | None = None,
) -> Tuple[List[LocalizationIssue], LocalizationTable, List[LocalizationColumn]]:
    files, columns = detect_dictionary_language_files(path)
    if len(files) < 2:
        raise FileReadError("未识别到多文件字典：请选择包含 dictionary_*.xlsx 的文件夹")

    files_by_code = {info.code: info for info in files}
    source_code = _choose_dictionary_source(files, options)
    if not source_code or source_code not in files_by_code:
        raise FileReadError("未识别到源语言文件，请在源语言下拉框中选择可用语言")

    enabled = [value.strip() for value in (options.enabled_language_columns or []) if value.strip()]
    enabled_codes = [code for code in enabled if code in files_by_code and code != source_code]
    if not enabled_codes:
        enabled_codes = [info.code for info in files if info.code != source_code]
    if not enabled_codes:
        raise FileReadError("未识别到目标语言文件：源语言之外至少需要 1 份 dictionary_*.xlsx")

    read_codes = [source_code] + [code for code in enabled_codes if code != source_code]
    file_data: Dict[str, DictionaryFileData] = {}
    _emit_progress(progress_callback, "read_dictionary", f"正在读取字典文件：0/{len(read_codes)}", current=0, total=len(read_codes), ratio=0.05)
    for position, code in enumerate(read_codes, start=1):
        if stop_event is not None and getattr(stop_event, "is_set", lambda: False)():
            break
        info = files_by_code[code]
        _emit_progress(
            progress_callback,
            "read_dictionary",
            f"正在读取字典文件 {position}/{len(read_codes)}：{info.file_name}",
            current=position,
            total=len(read_codes),
            ratio=0.05 + (position / max(len(read_codes), 1)) * 0.25,
            sheet_name=info.file_name,
        )
        file_data[code] = _read_dictionary_file(info)

    source_data = file_data.get(source_code)
    if source_data is None:
        raise FileReadError("源语言文件读取失败，无法继续检查")

    headers = ["ID", "所属模块", source_code] + enabled_codes
    display_headers = ["ID", "所属模块", source_data.info.file_name] + [files_by_code[code].file_name for code in enabled_codes]
    detected_columns = [
        LocalizationColumn(
            name=_language_name_for_code(source_code),
            header=source_code,
            index=2,
            display_name=source_data.info.file_name,
            code=source_code,
            raw_header="",
        )
    ]
    for offset, code in enumerate(enabled_codes, start=3):
        detected_columns.append(LocalizationColumn(
            name=_language_name_for_code(code),
            header=code,
            index=offset,
            display_name=files_by_code[code].file_name,
            code=code,
            raw_header="",
        ))

    row_records: List[LocalizationRow] = []
    records: Dict[str, LocalizationRow] = {}
    metadata_by_id: Dict[str, Dict[str, object]] = {}
    target_row_by_code: Dict[str, Dict[str, int]] = {}
    target_file_by_code = {code: files_by_code[code].file_name for code in enabled_codes}
    target_sheet_by_code = {code: file_data.get(code).info.sheet_name if code in file_data else files_by_code[code].file_name for code in enabled_codes}

    issues: List[LocalizationIssue] = []
    issue_index = 1

    if options.check_duplicate_id:
        for code, data in file_data.items():
            for text_id, rows_ in sorted(data.duplicate_ids.items()):
                issues.append(_make_issue(
                    issue_index,
                    "重复 ID",
                    text_id,
                    code,
                    ",".join(map(str, rows_)),
                    "",
                    "",
                    f"{data.info.file_name} 中同一个 ID 出现多次",
                    "确认是否需要保留唯一文本 ID",
                    data.info.sheet_name or data.info.file_name,
                    "普通问题",
                    "重复 ID",
                    source_code,
                    source_file=source_data.info.file_name,
                    target_file=data.info.file_name,
                ))
                issue_index += 1

    source_ids = set(source_data.records.keys())
    for code in enabled_codes:
        data = file_data.get(code)
        if data is None:
            continue
        target_ids = set(data.records.keys())
        target_row_by_code[code] = {text_id: record.row_number for text_id, record in data.records.items()}
        if options.check_missing_id:
            for text_id in sorted(source_ids - target_ids):
                source_record = source_data.records[text_id]
                issues.append(_make_issue(
                    issue_index,
                    "ID 缺失",
                    text_id,
                    code,
                    source_record.row_number,
                    source_record.text,
                    "",
                    f"{data.info.file_name} 缺少源语言 {source_code} 中存在的 ID",
                    "补充该语言对应 ID 的翻译行，或确认该 ID 是否无需翻译",
                    data.info.sheet_name or data.info.file_name,
                    "普通问题",
                    "ID 缺失/多出",
                    source_code,
                    module=source_record.module,
                    source_file=source_data.info.file_name,
                    target_file=data.info.file_name,
                ))
                issue_index += 1
            for text_id in sorted(target_ids - source_ids):
                target_record = data.records[text_id]
                issues.append(_make_issue(
                    issue_index,
                    "ID 多出",
                    text_id,
                    code,
                    target_record.row_number,
                    "",
                    target_record.text,
                    f"{data.info.file_name} 存在源语言 {source_code} 中没有的 ID",
                    "确认该 ID 是否为新增内容，必要时同步到源语言文件",
                    data.info.sheet_name or data.info.file_name,
                    "普通问题",
                    "ID 缺失/多出",
                    source_code,
                    module=target_record.module,
                    source_file=source_data.info.file_name,
                    target_file=data.info.file_name,
                ))
                issue_index += 1

    for source_record in source_data.records.values():
        values = {0: source_record.text_id, 1: source_record.module, 2: source_record.text}
        metadata_by_id[source_record.text_id] = {
            "module": source_record.module,
            "source_row": source_record.row_number,
        }
        for offset, code in enumerate(enabled_codes, start=3):
            target_data = file_data.get(code)
            target_record = target_data.records.get(source_record.text_id) if target_data else None
            # Missing IDs are reported separately. Use source text as a neutral
            # placeholder to avoid duplicate "空翻译" noise for the same problem.
            values[offset] = target_record.text if target_record is not None else source_record.text
        row = LocalizationRow(text_id=source_record.text_id, row_number=source_record.row_number, values=values)
        row_records.append(row)
        records[source_record.text_id] = row

    diagnostics = [
        f"多文件字典模式：识别到 {len(files)} 份 dictionary_*.xlsx；源语言={source_code}/{source_data.info.file_name}；目标语言={', '.join(enabled_codes)}",
        "字典模式固定忽略“所属模块”列，只用于报告定位；目标空翻译仅在源文本非空时报告。",
        "字典大表模式已自动跳过低收益高成本规则：拼写疑似、术语一致、不同源文本译文重复；保留空翻译、残留中文、占位符、标签、特殊符号和数值检查。",
    ]
    for data in file_data.values():
        diagnostics.extend(data.diagnostics[:1])

    table = LocalizationTable(
        path=str(_dictionary_root(path)),
        sheet_name=f"多文件字典（源语言 {source_code}）",
        headers=headers,
        key_col_index=0,
        records=records,
        duplicate_ids={},
        invalid_rows=[],
        display_headers=display_headers,
        type_headers=[""] * len(headers),
        header_row=1,
        data_start_row=5,
        structure_type="多文件字典模式",
        diagnostics=diagnostics,
        row_records=row_records,
    )

    local_options = replace(
        options,
        source_language_column=source_code,
        enabled_language_columns=[source_code] + enabled_codes,
        table_structure="单 Sheet 多语言列",
        check_spelling=False,
        check_terms=False,
        check_duplicate_translation=False,
        empty_requires_source_text=True,
    )
    quality_issues = _check_single_table_quality(table, detected_columns, local_options, stop_event, progress_callback, 0.35, 0.96)
    issues.extend(quality_issues)

    for item in issues:
        text_id = item.item_id
        meta = metadata_by_id.get(text_id, {})
        item.module = item.module or str(meta.get("module", ""))
        item.source_file = item.source_file or source_data.info.file_name
        code = str(item.language or "").strip()
        if code in target_file_by_code:
            item.target_file = item.target_file or target_file_by_code[code]
            item.sheet_name = target_sheet_by_code.get(code, item.sheet_name)
            row_number = target_row_by_code.get(code, {}).get(text_id)
            if row_number:
                item.row_number = str(row_number)
        elif code == source_code:
            item.target_file = item.target_file or source_data.info.file_name
            item.sheet_name = source_data.info.sheet_name or source_data.info.file_name

    columns = [column for column in detected_columns]
    return _reindex(issues), table, columns


def _language_allowed_chinese(language: str, allowed_chinese: set[str]) -> bool:
    return _is_allowed_chinese_language(language, allowed_chinese)


def _newline_marker_count(text: str) -> int:
    value = "" if text is None else str(text)
    escaped = len(ESCAPED_NEWLINE_RE.findall(value))
    actual = value.replace("\\n", "").count("\n")
    return escaped + actual


def _visible_text_length(text: str) -> int:
    value = "" if text is None else str(text)
    for regex in (
        BRACE_PLACEHOLDER_RE,
        SQUARE_PLACEHOLDER_RE,
        ANGLE_PLACEHOLDER_RE,
        PRINTF_PLACEHOLDER_RE,
        POSITIONAL_PERCENT_PLACEHOLDER_RE,
        ANY_BRACKET_TAG_RE,
        ANY_ANGLE_TAG_RE,
        HEX_COLOR_RE,
    ):
        value = regex.sub("", value)
    value = re.sub(r"\\n|\s+", "", value)
    value = re.sub(r"[，。！？；：、,.!?;:()\[\]{}<>/%％+\\-]", "", value)
    return len(value)


def _text_truncation_risk(language: str, source_text: str, current_text: str) -> Tuple[bool, str]:
    if not source_text or not current_text:
        return False, ""
    source_len = _visible_text_length(source_text)
    current_len = _visible_text_length(current_text)
    if source_len <= 0 or current_len <= 0:
        return False, ""
    lang_code = _language_code_from_label(language)
    ratio = current_len / max(source_len, 1)
    # 先保守处理：短中文文本翻译明显膨胀、或任意文本超过较高倍率时提示人工确认。
    threshold = 4.2 if lang_code in {"DE", "FR", "ES", "TR", "ARB"} else 3.8
    if source_len <= 8 and current_len >= 30 and ratio >= 3.8:
        return True, f"源文本可见长度约 {source_len}，译文约 {current_len}，短文本扩展明显"
    if source_len <= 12 and current_len >= 42 and ratio >= threshold:
        return True, f"源文本可见长度约 {source_len}，译文约 {current_len}，短文本扩展明显"
    if source_len <= 24 and current_len >= 80 and ratio >= 3.2:
        return True, f"源文本可见长度约 {source_len}，译文约 {current_len}，可能超出 UI 容量"
    return False, ""


def _add_text_quality_issues(issues: List[LocalizationIssue], index: int, *, text_id: str, language: str, sheet_name: str, row_number: int | str, source_text: str, current_text: str, source_language: str, allowed_chinese: set[str], symbols: Sequence[str], options: LocalizationCheckOptions, enable_layout_heuristics: bool = False, source_tag_invalid: bool = False) -> int:
    if options.check_empty and _empty(current_text):
        if not options.empty_requires_source_text or str(source_text or "").strip():
            issues.append(_make_issue(index, "空翻译", text_id, language, row_number, source_text, current_text, f"{language} 翻译为空", f"补充 {language} 翻译文本，或确认该 ID 是否允许为空", sheet_name, "高风险错误", "空翻译", source_language))
            return index + 1
        return index

    unfinished_marker = _contains_unfinished_marker(current_text) if options.check_unfinished else ""
    if unfinished_marker:
        issues.append(_make_issue(index, "未完成/修改中占位文本", text_id, language, row_number, source_text, current_text, f"命中未完成占位文本：{unfinished_marker}", "补齐正式翻译；命中后已跳过数值、标签、占位符等派生检查，避免重复噪音", sheet_name, "高风险错误", "未完成占位文本", source_language))
        return index + 1

    if options.check_chinese and not _is_allowed_chinese_language(language, allowed_chinese) and CHINESE_RE.search(current_text or ""):
        lang_code = _language_code_from_label(language)
        if lang_code == "JP":
            jp_level, jp_reason = _jp_chinese_residual_level(current_text)
            if jp_level == "high":
                issues.append(_make_issue(index, "残留中文", text_id, language, row_number, source_text, current_text, f"JP 文本疑似整句中文或无假名中文片段：{jp_reason}", "确认是否为未翻译中文；正常日文汉字已兼容", sheet_name, "高风险错误", "残留中文", source_language))
                index += 1
            elif jp_level == "suspected":
                issues.append(_make_issue(index, "JP 日文汉字疑似残留", text_id, language, row_number, source_text, current_text, f"JP 文本包含疑似简中或占位中文片段：{jp_reason}", "人工确认是否为正常日文汉字或残留中文", sheet_name, "疑似问题", "残留中文", source_language, True))
                index += 1
        else:
            fragments = _extract_unallowed_chinese_fragments(current_text, allowed_chinese)
            if fragments:
                issues.append(_make_issue(index, "残留中文", text_id, language, row_number, source_text, current_text, f"非中文语言中检测到未命中白名单的中文片段：{', '.join(fragments[:8])}", "确认是否为未翻译文本；语言名称/项目名/专有名词可复用允许中文白名单配置", sheet_name, "高风险错误", "残留中文", source_language))
                index += 1

    if options.check_placeholders and (PLACEHOLDER_SIGNAL_RE.search(source_text or "") or PLACEHOLDER_SIGNAL_RE.search(current_text or "")):
        source_placeholder_invalid = bool(_placeholder_syntax_errors(source_text))
        target_placeholder_invalid = bool(_placeholder_syntax_errors(current_text))
        issue_type, level, remark, suggestion, suspected = _malformed_placeholder_issue(current_text, target=True)
        if issue_type:
            issues.append(_make_issue(index, issue_type, text_id, language, row_number, source_text, current_text, remark, suggestion, sheet_name, level, "占位符", source_language, suspected))
            index += 1
        if not source_placeholder_invalid and not target_placeholder_invalid:
            issue_type, level, remark, suggestion, suspected = _placeholder_issue(source_text, current_text)
            if issue_type:
                issues.append(_make_issue(index, issue_type, text_id, language, row_number, source_text, current_text, remark, suggestion, sheet_name, level, "占位符", source_language, suspected))
                index += 1

    if options.check_tags and (TAG_SIGNAL_RE.search(source_text or "") or TAG_SIGNAL_RE.search(current_text or "")):
        tag_error = _tag_error(current_text)
        source_tag_signature = [(kind, name, _normalize_tag_attrs(attrs)) for _pos, kind, name, attrs in _rich_tag_items(source_text)]
        target_tag_signature = [(kind, name, _normalize_tag_attrs(attrs)) for _pos, kind, name, attrs in _rich_tag_items(current_text)]
        report_target_tag_error = bool(tag_error and (not source_tag_invalid or source_tag_signature != target_tag_signature))
        if report_target_tag_error:
            issues.append(_make_issue(index, "目标文本标签结构异常", text_id, language, row_number, source_text, current_text, tag_error, "修正富文本标签括号、嵌套和闭合关系", sheet_name, "高风险错误", "标签词法预检查", source_language))
            index += 1
        if not source_tag_invalid and not tag_error:
            tag_issue_type, level, remark, suggestion, suspected = _tag_compare_issue(source_text, current_text)
            if tag_issue_type and tag_issue_type != "纯标签片段":
                issues.append(_make_issue(index, tag_issue_type, text_id, language, row_number, source_text, current_text, remark, suggestion, sheet_name, level, "标签", source_language, suspected))
                index += 1

    if options.check_symbols:
        source_newlines = _newline_marker_count(source_text)
        current_newlines = _newline_marker_count(current_text)
        if source_newlines != current_newlines:
            issues.append(_make_issue(index, "换行符数量差异", text_id, language, row_number, source_text, current_text, f"源文本换行标记数量为 {source_newlines}，译文为 {current_newlines}；该差异属于格式排版，不属于业务占位符", "人工确认目标语言排版；如无需保持相同行数可忽略，否则调整 \\n/实际换行", sheet_name, "疑似问题", "格式/换行符", source_language, True))
            index += 1
    if options.check_symbols:
        missing, _compatible = _missing_critical_symbols(source_text, current_text, symbols)
        if missing:
            issues.append(_make_issue(index, "特殊符号差异待确认", text_id, language, row_number, source_text, current_text, f"上下文结构中的符号存在差异：{', '.join(missing)}", "人工确认该符号是否属于 URL、路径、比例、日期或配置表达式", sheet_name, "需人工确认", "特殊符号上下文", source_language, True))
            index += 1

    if options.check_numbers and (NUMBER_SIGNAL_RE.search(source_text or "") or NUMBER_SIGNAL_RE.search(current_text or "")):
        number_issue_type, level, number_remark, suspected = _number_consistency_error(source_text, current_text)
        if number_issue_type:
            issues.append(_make_issue(index, number_issue_type, text_id, language, row_number, source_text, current_text, number_remark, "请确认译文中的业务数值是否与源文本一致", sheet_name, level, "数值", source_language, suspected))
            index += 1

    if options.check_spelling:
        bad, suggestion = _spelling_issue(language, current_text)
        if bad:
            issues.append(_make_issue(index, "英文疑似拼写错误", text_id, language, row_number, source_text, current_text, f"发现疑似拼写错误片段：{bad}；{suggestion}", "人工确认拼写；如为项目专有名词可加入白名单", sheet_name, "疑似问题", "拼写检查", source_language, True))
            index += 1

    if options.check_terms:
        group_name, found_terms = _term_issue(language, current_text)
        if group_name:
            issues.append(_make_issue(index, "术语一致性疑似问题", text_id, language, row_number, source_text, current_text, f"同一文本中混用术语：{group_name}；命中={found_terms}", "确认目标语言内功能名是否需要统一表达", sheet_name, "疑似问题", "术语一致性", source_language, True))
            index += 1

        if enable_layout_heuristics:
            truncation, truncation_reason = _text_truncation_risk(language, source_text, current_text)
            if truncation:
                issues.append(_make_issue(index, "文本截断风险", text_id, language, row_number, source_text, current_text, truncation_reason, "人工确认目标语言在实际 UI 容器中是否会被截断", sheet_name, "疑似问题", "长度/截断风险", source_language, True))
                index += 1

    return index


def _add_source_self_format_issues(issues: List[LocalizationIssue], index: int, *, text_id: str, language: str, sheet_name: str, row_number: int | str, source_text: str, options: LocalizationCheckOptions) -> int:
    source_label = _language_code_from_label(language) or language
    if options.check_placeholders and PLACEHOLDER_SIGNAL_RE.search(source_text or ""):
        issue_type, level, remark, suggestion, suspected = _malformed_placeholder_issue(source_text, target=False)
        if issue_type:
            issues.append(_make_issue(index, f"{source_label} 源文本占位符结构异常", text_id, language, row_number, source_text, source_text, remark, suggestion, sheet_name, "源文本问题", "源文本占位符", language, suspected, result_category="源文本问题", count_in_error_stats=False))
            index += 1
    if options.check_tags and TAG_SIGNAL_RE.search(source_text or ""):
        tag_error = _tag_error(source_text)
        if tag_error:
            issues.append(_make_issue(index, f"{source_label} 源文本标签结构异常", text_id, language, row_number, source_text, source_text, tag_error, "先修正源文本富文本标签括号、嵌套和闭合关系，再同步翻译", sheet_name, "源文本问题", "源文本标签词法预检查", language, result_category="源文本问题", count_in_error_stats=False))
            index += 1
    return index


def _add_english_copy_issue(issues: List[LocalizationIssue], index: int, *, text_id: str, language: str, sheet_name: str, row_number: int | str, source_text: str, current_text: str, english_text: str, source_language: str) -> int:
    lang_code = _language_code_from_label(language)
    if lang_code in {"", "EN", "CN", "TW"}:
        return index
    current_norm = _normalize_duplicate_text(current_text)
    english_norm = _normalize_duplicate_text(english_text)
    if not current_norm or len(current_norm) < 4 or current_norm != english_norm:
        return index
    issues.append(_make_issue(
        index,
        "待确认",
        text_id,
        language,
        row_number,
        source_text,
        current_text,
        "当前译文与英文列完全一致，疑似复制英文或尚未本地化",
        "人工确认该语言是否允许直接使用英文；若不允许，请补充本地化翻译",
        sheet_name,
        "疑似问题",
        "英文复制检查",
        source_language,
        True,
    ))
    return index + 1


def _add_same_source_translation_variant_issues(issues: List[LocalizationIssue], index: int, table: LocalizationTable, target_columns: Sequence[LocalizationColumn], source_col: int, source_language: str) -> int:
    grouped: Dict[Tuple[str, str], List[Tuple[int, str, str]]] = defaultdict(list)
    records_to_check = table.row_records or list(table.records.values())
    for record in records_to_check:
        source_text = record.values.get(source_col, "")
        source_key = _normalize_duplicate_text(source_text)
        if len(source_key) < 2:
            continue
        for column in target_columns:
            language = _localization_column_label(column)
            current_text = record.values.get(column.index, "")
            current_key = _normalize_duplicate_text(current_text)
            if len(current_key) < 4 or _contains_unfinished_marker(current_text):
                continue
            grouped[(language, source_key)].append((record.row_number, record.text_id, current_text))

    for (language, _source_key), items in grouped.items():
        variants: Dict[str, Tuple[int, str, str]] = {}
        for row_number, text_id, translation in items:
            variants.setdefault(_normalize_duplicate_text(translation), (row_number, text_id, translation))
        if len(variants) < 2:
            continue
        examples = list(variants.values())[:4]
        source_text = ""
        for record in records_to_check:
            if _normalize_duplicate_text(record.values.get(source_col, "")) == _source_key:
                source_text = record.values.get(source_col, "")
                break
        row_text = "；".join(str(item[0]) for item in examples)
        current_text = "；".join(f"{item[2]}（行{item[0]}）" for item in examples)
        related_ids = "；".join(item[1] for item in examples if item[1])
        issues.append(_make_issue(
            index,
            "术语不一致",
            related_ids,
            language,
            row_text,
            source_text,
            current_text,
            "相同源文本在同一目标语言中存在多种译法，可能造成术语或按钮文案不统一",
            "人工确认是否为正常上下文差异；若不是，请统一该源文本的目标译法",
            table.sheet_name,
            "疑似问题",
            "同源文本译文一致性",
            source_language,
            True,
        ))
        index += 1
    return index


def _add_duplicate_translation_issues_for_entries(issues: List[LocalizationIssue], index: int, entries: Sequence, *, source_language: str) -> int:
    grouped: Dict[Tuple[str, str], List] = defaultdict(list)
    for entry in entries:
        translation = getattr(entry, "translation", "")
        if not _is_translation_duplicate_candidate(translation):
            continue
        language = str(getattr(entry, "language", ""))
        key = (language, _normalize_duplicate_text(translation))
        grouped[key].append(entry)
    for (language, _norm_text), items in grouped.items():
        if len(items) < 2:
            continue
        distinct_sources = {str(getattr(item, "source_text", "")).strip() for item in items if str(getattr(item, "source_text", "")).strip()}
        if len(distinct_sources) < 2:
            continue
        source_list = list(distinct_sources)
        if len(source_list) == 2 and SequenceMatcher(None, _normalize_duplicate_text(source_list[0]), _normalize_duplicate_text(source_list[1])).ratio() >= 0.86:
            # 源文本高度相近时，重复译文多半是正常复用，避免刷屏；保留明显源意不同的重复译文异常。
            continue
        # 控制噪音：同一重复组最多输出 5 条，报告里给出关联 ID。
        related_ids = ", ".join(str(getattr(item, "text_id", "")) for item in items[:8])
        for item in items[:5]:
            issues.append(_make_issue(index, "不同源文本译文重复", getattr(item, "text_id", ""), language, getattr(item, "row_number", ""), getattr(item, "source_text", ""), getattr(item, "translation", ""), f"同一语言内不同源文本翻成完全相同译文；关联 ID：{related_ids}", "人工确认是否为正常复用文本；若源文本含义不同，需修正译文", getattr(item, "sheet_name", language), "疑似问题", "译文重复", source_language, True))
            index += 1
    return index


def _add_duplicate_translation_issues_for_single_sheet(issues: List[LocalizationIssue], index: int, table: LocalizationTable, target_columns: Sequence[LocalizationColumn], source_col: int, source_language: str) -> int:
    class EntryObj:
        pass
    entries = []
    records_to_check = table.row_records or list(table.records.values())
    for record in records_to_check:
        text_id = record.text_id
        source_text = record.values.get(source_col, "")
        for column in target_columns:
            col_index = column.index
            if col_index == table.key_col_index or col_index == source_col:
                continue
            header = _localization_column_label(column, table.headers[col_index] if col_index < len(table.headers) else f"列{col_index + 1}")
            translation = record.values.get(col_index, "")
            e = EntryObj()
            e.text_id = text_id
            e.source_text = source_text
            e.translation = translation
            e.language = header
            e.sheet_name = table.sheet_name
            e.row_number = record.row_number
            entries.append(e)
    return _add_duplicate_translation_issues_for_entries(issues, index, entries, source_language=source_language)

def _check_multi_sheet_quality(dataset: NormalizedLocalizationDataset, options: LocalizationCheckOptions, stop_event=None) -> Tuple[List[LocalizationIssue], LocalizationTable, List[LocalizationColumn]]:
    allowed_chinese = {_normalize_column_key(value) for value in _split_config_values(options.allowed_chinese_languages_text)}
    ignore_ids = {value for value in _split_config_values(options.ignore_ids_text)}
    symbols = _symbol_tokens(options.symbols_text)
    expected_languages = _expected_language_codes(options)
    max_length = safe_int(options.max_length, 0) if options.max_length else 0
    try:
        length_ratio = float(options.length_ratio or 0)
    except Exception:
        length_ratio = 0.0

    issues: List[LocalizationIssue] = []
    index = 1

    for message in dataset.header_errors:
        sheet_name = message.split("：", 1)[0] if "：" in message else ""
        issues.append(_make_issue(index, "表头异常", "", "", "", "", "", message, "检查该语言 Sheet 是否包含 ID、源文本和翻译列", sheet_name, "普通问题", "解析诊断"))
        index += 1

    if expected_languages:
        existing = {code.upper() for code in dataset.records_by_language.keys()}
        for code in expected_languages:
            if code and code not in existing:
                issues.append(_make_issue(index, "语言 Sheet 缺失", "", code, "", "", "", f"预期语言 {code} 缺少对应 Sheet", "补充该语言 Sheet，或从预期语言列表中移除", code, "普通问题", "结构识别"))
                index += 1

    if options.check_invalid_id:
        for sheet_name, row_number in dataset.invalid_rows:
            issues.append(_make_issue(index, "无效 ID", "", "", row_number, "", "", "主键 ID 为空但该行存在内容", "补充 ID 或删除无效行", sheet_name, "普通问题", "主键"))
            index += 1

    if options.check_duplicate_id:
        for language, dup_map in sorted(dataset.duplicate_ids.items()):
            for text_id, rows_ in sorted(dup_map.items()):
                issues.append(_make_issue(index, "重复 ID", text_id, language, ",".join(map(str, rows_)), "", "", "同一个 ID 在当前语言 Sheet 中出现多次", "确认是否需要保留唯一配置", language, "普通问题", "重复 ID"))
                index += 1

    base_language = dataset.base_language
    base_records = dataset.records_by_language.get(base_language, {}) if base_language else {}
    base_ids = set(base_records.keys())
    if options.check_missing_id and base_ids:
        for language, records in sorted(dataset.records_by_language.items(), key=lambda item: LANGUAGE_PRIORITY.index(item[0]) if item[0] in LANGUAGE_PRIORITY else 999):
            if language == base_language:
                continue
            ids = set(records.keys())
            for text_id in sorted(base_ids - ids):
                if text_id in ignore_ids:
                    continue
                base_entry = base_records.get(text_id)
                issues.append(_make_issue(index, "ID 缺失", text_id, language, "", base_entry.source_text if base_entry else "", "", f"{language} Sheet 缺少基准语言 {base_language} 中存在的 ID", "补充该语言对应 ID 的翻译行，或确认该 ID 是否无需翻译", language, "普通问题", "ID 缺失/多出", base_language))
                index += 1
            for text_id in sorted(ids - base_ids):
                if text_id in ignore_ids:
                    continue
                entry = records.get(text_id)
                issues.append(_make_issue(index, "ID 多出", text_id, language, entry.row_number if entry else "", entry.source_text if entry else "", entry.translation if entry else "", f"{language} Sheet 存在基准语言 {base_language} 中没有的 ID", "确认该 ID 是否为新增内容，必要时同步到其他语言 Sheet", entry.sheet_name if entry else language, "普通问题", "ID 缺失/多出", base_language))
                index += 1

    if options.check_source_consistency:
        for text_id in sorted(dataset.text_ids):
            if text_id in ignore_ids:
                continue
            source_map: Dict[str, List] = {}
            for language, records in dataset.records_by_language.items():
                entry = records.get(text_id)
                if entry is None:
                    continue
                key = (entry.source_text or "").strip()
                source_map.setdefault(key, []).append(entry)
            non_empty_sources = {k: v for k, v in source_map.items() if k}
            if len(non_empty_sources) > 1:
                canonical = next(iter(non_empty_sources.keys()))
                for source_text, entries in non_empty_sources.items():
                    if source_text == canonical:
                        continue
                    for entry in entries:
                        issues.append(_make_issue(index, "源文本不一致", text_id, entry.language, entry.row_number, canonical, entry.source_text, f"同一 ID 在不同语言 Sheet 中源文本不一致；基准源文={canonical}", "确认各语言 Sheet 的源文本是否同步到最新版本", entry.sheet_name, "普通问题", "源文本一致", base_language))
                        index += 1

    reported_source_empty: set[Tuple[str, str]] = set()
    reported_source_format: set[Tuple[str, str]] = set()
    for entry in dataset.entries:
        if stop_event is not None and getattr(stop_event, "is_set", lambda: False)():
            break
        if entry.text_id in ignore_ids:
            continue
        # Multi-language Sheet pages often repeat the same CN source in every
        # language Sheet. Deduplicate by ID + normalized source content so one
        # source defect is reported once, while different source revisions are
        # still kept separate.
        source_key = (entry.text_id, _normalize_duplicate_text(entry.source_text))
        if options.check_empty and _empty(entry.source_text):
            if source_key not in reported_source_empty:
                issues.append(_make_issue(
                    index,
                    f"{base_language}源文本为空",
                    entry.text_id,
                    base_language,
                    entry.row_number,
                    entry.source_text,
                    entry.source_text,
                    f"ID 已存在，但 {base_language} 源文本为空或仅包含空格/换行/不可见字符",
                    "先补充或确认源语言文本，再进行目标语言质量比对",
                    entry.sheet_name,
                    "源文本问题",
                    "源文本空文本",
                    base_language,
                ))
                index += 1
                reported_source_empty.add(source_key)
            continue
        source_tag_invalid = bool(options.check_tags and TAG_SIGNAL_RE.search(entry.source_text or "") and _tag_error(entry.source_text))
        if source_tag_invalid and source_key not in reported_source_format:
            index = _add_source_self_format_issues(
                issues,
                index,
                text_id=entry.text_id,
                language=base_language,
                sheet_name=entry.sheet_name,
                row_number=entry.row_number,
                source_text=entry.source_text,
                options=options,
            )
            reported_source_format.add(source_key)
        index = _add_text_quality_issues(
            issues, index,
            text_id=entry.text_id,
            language=entry.language,
            sheet_name=entry.sheet_name,
            row_number=entry.row_number,
            source_text=entry.source_text,
            current_text=entry.translation,
            source_language=base_language,
            allowed_chinese=allowed_chinese,
            symbols=symbols,
            options=options,
            source_tag_invalid=source_tag_invalid,
        )
        if options.check_length:
            text_len = len(entry.translation or "")
            source_len = len(entry.source_text or "")
            if max_length > 0 and text_len > max_length:
                issues.append(_make_issue(index, "文本过长", entry.text_id, entry.language, entry.row_number, entry.source_text, entry.translation, f"当前长度 {text_len}，超过最大长度 {max_length}", "精简文本或确认 UI 容量", entry.sheet_name, "普通问题", "长度", base_language))
                index += 1
            if length_ratio > 0 and source_len > 0 and text_len > source_len * length_ratio:
                issues.append(_make_issue(index, "文本过长", entry.text_id, entry.language, entry.row_number, entry.source_text, entry.translation, f"当前长度 {text_len}，超过源文本长度 {source_len} 的 {length_ratio:g} 倍", "检查该语言是否可能超框", entry.sheet_name, "普通问题", "长度", base_language))
                index += 1

    if options.check_duplicate_translation:
        index = _add_duplicate_translation_issues_for_entries(issues, index, dataset.entries, source_language=base_language)

    columns = [LocalizationColumn(name=info.language, header=info.sheet_name, index=i) for i, info in enumerate(dataset.sheet_infos) if not info.error]
    table = LocalizationTable(
        path=dataset.path,
        sheet_name="多 Sheet 语言页",
        headers=["ID", "源文本", "语言", "翻译", "Sheet", "行号"],
        key_col_index=0,
        records={},
        duplicate_ids={},
        invalid_rows=[row for _, row in dataset.invalid_rows],
        header_row=1,
        data_start_row=2,
        structure_type="多 Sheet 多语言页",
        diagnostics=[f"已识别语言页：{', '.join(dataset.records_by_language.keys())}", f"已识别文本 ID 数量：{len(dataset.text_ids)}"],
    )
    for idx, text_id in enumerate(sorted(dataset.text_ids), start=1):
        table.records[text_id] = LocalizationRow(text_id=text_id, row_number=idx, values={})
    return _reindex(issues), table, columns



def _build_no_language_columns_error(table: LocalizationTable, detected_columns: Sequence[LocalizationColumn]) -> str:
    candidates: List[str] = []
    excluded: List[str] = []
    for idx, header in enumerate(table.headers):
        display = table.display_headers[idx] if idx < len(table.display_headers) else ""
        code, name, confidence = _language_from_header(header, display)
        label = str(header or display or f"列{idx + 1}").replace("\n", " / ")
        if code:
            candidates.append(f"{label} -> {code}/{name}/{confidence}")
        elif not str(header).strip() and not str(display).strip():
            excluded.append(f"列{idx + 1}=空列")
        elif _is_auxiliary_header(header, display):
            excluded.append(f"{label}=辅助列")
    first_row_hint = ""
    first_headers = table.headers[:8]
    joined = " ".join(str(value) for value in first_headers)
    if any(token in _normalize_column_key(joined) for token in ["text_id", "source_text", "translation", "翻译", "原文"]):
        first_row_hint = "；检测到当前表头包含 Text ID / Translation / Source Text 等字段，建议确认该行是否为真实表头"
    return (
        "未识别到单 Sheet 多语言列结构。\n"
        f"当前识别结果：Sheet={table.sheet_name}；表头行=第 {table.header_row} 行；数据起始行=第 {table.data_start_row} 行；"
        f"主键列={table.headers[table.key_col_index] if 0 <= table.key_col_index < len(table.headers) else '未识别'}；"
        f"记录数={len(table.records)}。\n"
        f"候选语言字段：{'; '.join(candidates[:20]) if candidates else '未识别'}。\n"
        f"被排除字段：{'; '.join(excluded[:20]) if excluded else '无'}。\n"
        f"识别诊断：{'; '.join(table.diagnostics[:8]) if table.diagnostics else '无'}{first_row_hint}。\n"
        "建议：检查真实表头行、数据起始行、主键列；若是单行混合语言表头，请确认字段包含 EN Translation / CN Source Text 等语言标识。"
    )

def check_localization_file(
    path: str,
    options: LocalizationCheckOptions,
    stop_event=None,
    progress_callback: ProgressCallback | None = None,
) -> Tuple[List[LocalizationIssue], LocalizationTable, List[LocalizationColumn]]:
    _emit_progress(progress_callback, "detect_structure", "正在识别多语言表结构…", ratio=0.02)
    structure = _resolve_localization_structure(path, options, progress_callback)
    if structure == MULTI_FILE_DICTIONARY_STRUCTURE:
        return _check_dictionary_files_quality(path, options, stop_event, progress_callback)
    if structure == MULTI_TABLE_SHEETS_STRUCTURE:
        return _check_multi_table_sheets_quality(path, options, stop_event, progress_callback)
    if structure == MULTI_SHEET_STRUCTURE:
        _emit_progress(progress_callback, "read_multisheet", "正在读取多 Sheet 语言页…", ratio=0.08)
        dataset = parse_multi_sheet_localization(path, safe_int(options.header_row, 0), _enabled_language_codes(options))
        _emit_progress(progress_callback, "check_multisheet", f"正在检查多 Sheet 语言页：{len(dataset.entries)} 条翻译", ratio=0.35)
        return _check_multi_sheet_quality(dataset, options, stop_event)
    if structure == UNKNOWN_STRUCTURE:
        raise FileReadError("未能自动识别多语言表结构，请手动选择「单 Sheet 多语言列」或「多 Sheet 语言页」")

    _emit_progress(progress_callback, "read_single_sheet", "正在读取单 Sheet 翻译表…", ratio=0.08)
    table = parse_localization_table(options, path)
    detected_columns = detect_language_columns(table.headers, table.key_col_index, table.display_headers, table.type_headers)
    issues = _check_single_table_quality(table, detected_columns, options, stop_event, progress_callback, 0.18, 0.96)
    _emit_progress(progress_callback, "check_done", "多语言检查处理完成，正在渲染结果…", ratio=0.98)
    return _reindex(issues), table, detected_columns


def _reindex(issues: Iterable[LocalizationIssue]) -> List[LocalizationIssue]:
    items = list(issues)
    for index, item in enumerate(items, start=1):
        item.index = index
    return items


def localization_issues_to_tsv(issues: Iterable[LocalizationIssue]) -> str:
    headers = ["序号", "问题等级", "问题类型", "ID", "源语言列", "目标语言列", "Sheet", "行号", "源文本", "当前文本", "命中规则", "问题说明", "建议处理", "是否人工确认项", "是否计入错误统计", "规则置信度", "所属模块", "源文件", "目标文件"]
    rows = ["\t".join(headers)]
    for item in issues:
        rows.append("\t".join([
            str(item.index), item.issue_level, item.issue_type, _format_identifier(item.item_id), item.source_language, item.language, item.sheet_name, item.row_number,
            item.source_text.replace("\n", "\\n"), item.current_text.replace("\n", "\\n"), item.rule_name, item.remark, item.suggestion,
            "是" if item.requires_manual_confirm else "否", "是" if item.count_in_error_stats else "否", item.rule_confidence,
            item.module, item.source_file, item.target_file,
        ]))
    return "\n".join(rows)


def build_single_sheet_preview(table: LocalizationTable, columns: Sequence[LocalizationColumn]) -> str:
    def col_label(index: int) -> str:
        if index < 0:
            return "未识别"
        n = index + 1
        s = ""
        while n:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        header = table.headers[index] if index < len(table.headers) else ""
        display = table.display_headers[index] if index < len(table.display_headers) else ""
        if display and display != header:
            return f"{s}列 / {header}（{display}）"
        return f"{s}列 / {header or '空表头'}"

    language_text = "、".join(
        f"{col.header}->{col.name}" + (f"（展示名：{col.display_name}）" if col.display_name else "")
        for col in columns
    ) or "未识别"
    excluded = "、".join(f"{name}：{reason}" for name, reason in table.excluded_columns.items()) or "无"
    sample_ids = list(table.records.keys())[:5]
    lines = [
        f"已识别结构：{table.structure_type}",
        f"当前 Sheet：{table.sheet_name}",
        f"表头所在行：第 {table.header_row} 行",
        f"数据起始行：第 {table.data_start_row} 行",
        f"主键列：{col_label(table.key_col_index)}",
        f"有效记录数：{len(table.records)}",
        f"识别到的语言列：{language_text}",
        f"被排除的辅助列：{excluded}",
        f"前 5 个 ID：{', '.join(sample_ids) if sample_ids else '无'}",
    ]
    if table.diagnostics:
        lines.append("识别诊断：" + "；".join(table.diagnostics[:8]))
    return "\n".join(lines)


# ===== Excel 导出 =====

def export_localization_issues_to_excel(issues: Iterable[LocalizationIssue], output_path: str) -> None:
    """Export a categorized, filter-ready localization quality report."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    items = list(issues)
    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    wb = Workbook(write_only=True)

    def set_widths(ws, widths: Sequence[int]):
        # write_only 模式下设置列宽可能被部分环境忽略，但不会影响内容有效性。
        for index, width in enumerate(widths, start=1):
            try:
                ws.column_dimensions[get_column_letter(index)].width = width
            except Exception:
                pass

    def safe_text(value):
        text = "" if value is None else str(value)
        # Excel 单元格最多 32767 字符；同时过滤不可见控制字符，避免保存失败。
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        return text if len(text) <= 32000 else text[:32000] + "...【已截断】"

    def append_safe_row(ws, values: Sequence[object]):
        cells = []
        for value in values:
            cell = WriteOnlyCell(ws, value=value)
            # Localization text may legitimately be "#NAME?" or begin with
            # "=". Force these values to string cells so Excel/openpyxl does
            # not reinterpret user text as an error or formula.
            if isinstance(value, str) and (value.startswith("=") or value in {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}):
                cell.data_type = "s"
            cells.append(cell)
        ws.append(cells)

    def write_header(ws, headers: Sequence[str]):
        try:
            # write_only worksheets must configure panes before the first row.
            ws.freeze_panes = "A2"
        except Exception:
            pass
        cells = []
        for value in headers:
            cell = WriteOnlyCell(ws, value=value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cells.append(cell)
        ws.append(cells)

    def finish_sheet(ws, column_count: int, row_count: int):
        try:
            ws.auto_filter.ref = f"A1:{get_column_letter(max(1, column_count))}{max(1, row_count)}"
        except Exception:
            pass

    category_groups: Dict[str, List[LocalizationIssue]] = defaultdict(list)
    for item in items:
        category = item.result_category or _result_category(item.issue_level, item.issue_type, item.rule_name)
        category_groups[category].append(item)

    grouped_languages: Dict[Tuple[str, str, str, str, str, str], set[str]] = defaultdict(set)
    for item in items:
        group_key = (
            item.sheet_name,
            item.row_number,
            _format_identifier(item.item_id),
            item.issue_type,
            item.rule_name,
            item.remark,
        )
        if item.language:
            grouped_languages[group_key].add(item.language)

    # 1. 汇总说明
    ws = wb.create_sheet("汇总说明")
    set_widths(ws, [24, 120])
    write_header(ws, ["指标", "数量/内容"])
    clear_items = category_groups.get("明确问题", [])
    source_items = category_groups.get("源文本问题", [])
    manual_items = category_groups.get("需人工确认", [])
    coverage_items = category_groups.get("语言覆盖范围", [])
    ignored_items = category_groups.get("忽略项或白名单", [])
    error_total = sum(1 for item in clear_items if item.count_in_error_stats)
    for row in [
        ["全部检查记录", len(items)],
        ["错误总数（高置信度明确问题）", error_total],
        ["明确问题", len(clear_items)],
        ["源文本问题", len(source_items)],
        ["需人工确认", len(manual_items)],
        ["语言覆盖记录", len(coverage_items)],
        ["忽略项或白名单命中", len(ignored_items)],
        ["不计入错误统计数", sum(1 for item in items if not item.count_in_error_stats)],
        ["明确问题类型", "；".join(f"{k}={v}" for k, v in Counter(item.issue_type for item in clear_items).most_common())],
        ["需人工确认类型", "；".join(f"{k}={v}" for k, v in Counter(item.issue_type for item in manual_items).most_common())],
        ["统计口径", "错误总数只统计明确问题且是否计入错误统计=是；源文本、人工确认、语言覆盖、白名单不计入目标语言错误总数。"],
    ]:
        append_safe_row(ws, row)
    finish_sheet(ws, 2, 12)

    detail_headers = ["序号", "结果分类", "Sheet 名称", "Excel 行号", "主键 ID", "源语言列", "目标语言列", "受影响语言", "源文本", "目标文本", "问题类型", "问题等级", "命中规则", "具体差异/命中原因", "处理建议", "是否人工确认项", "是否计入错误统计", "规则置信度", "所属模块", "源文件", "目标文件"]
    detail_widths = [8, 16, 20, 12, 20, 14, 14, 24, 46, 46, 20, 14, 20, 50, 40, 14, 16, 12, 22, 28, 28]

    def issue_row(item: LocalizationIssue):
        group_key = (item.sheet_name, item.row_number, _format_identifier(item.item_id), item.issue_type, item.rule_name, item.remark)
        affected = "、".join(sorted(grouped_languages.get(group_key, set()))) or item.affected_languages or item.language
        return [
            item.index, item.result_category, item.sheet_name, item.row_number, _format_identifier(item.item_id), item.source_language, item.language,
            safe_text(affected), safe_text(item.source_text), safe_text(item.current_text), item.issue_type, item.issue_level,
            item.rule_name, safe_text(item.remark), safe_text(item.suggestion), "是" if item.requires_manual_confirm else "否",
            "是" if item.count_in_error_stats else "否", item.rule_confidence,
            safe_text(item.module), safe_text(item.source_file), safe_text(item.target_file),
        ]

    def write_issue_sheet(name: str, sheet_items: Sequence[LocalizationIssue]) -> None:
        sheet = wb.create_sheet(name)
        set_widths(sheet, detail_widths)
        write_header(sheet, detail_headers)
        for item in sheet_items:
            append_safe_row(sheet, issue_row(item))
        finish_sheet(sheet, len(detail_headers), len(sheet_items) + 1)

    write_issue_sheet("明确问题", clear_items)
    write_issue_sheet("源文本问题", source_items)
    write_issue_sheet("需人工确认", manual_items)

    # 语言覆盖范围使用适合项目管理查看的独立字段。
    ws = wb.create_sheet("语言覆盖范围")
    coverage_headers = ["序号", "Sheet", "语言", "总数据行", "已填写", "空白", "检查状态", "说明"]
    set_widths(ws, [8, 22, 18, 14, 12, 12, 30, 80])
    write_header(ws, coverage_headers)
    for idx, item in enumerate(coverage_items, start=1):
        append_safe_row(ws, [idx, item.sheet_name, item.language, item.coverage_total, item.coverage_filled, item.coverage_blank, item.coverage_status, safe_text(item.remark)])
    finish_sheet(ws, len(coverage_headers), len(coverage_items) + 1)

    write_issue_sheet("忽略项或白名单命中", ignored_items)
    write_issue_sheet("全部明细", items)

    # 规则说明
    ws = wb.create_sheet("规则说明")
    set_widths(ws, [24, 120])
    write_header(ws, ["项目", "说明"])
    rule_row_count = 1
    for row in [
        ["结构识别", "支持单行/双行/三行表头；检测到 #tid、cn、en、tw、de、fr、id、ru 等程序字段时，优先使用该行为真实字段行。"],
        ["辅助列排除", "status/state/翻译状态/勿改/date/modified/batch/trans_round/change_log/备注/note/comment 等列默认排除，不作为语言列。"],
        ["语言列识别", "同时使用程序字段名和展示名映射，支持 cn/en/tw/de/fr/id/ru/jp/kr/th/it/pt/tr/vn/arb/es。"],
        ["占位符", "按完整结构识别 printf、%1、[%s1]、${name}、{name}、#NAME 等；名称和数量比较，顺序变化不判错，换行符不作为业务占位符。"],
        ["特殊符号/换行", "斜线、方括号和百分号仅在 URL、路径、日期、比例、占位符、注册标签或配置语法上下文中检查；换行差异单独归入人工确认。"],
        ["数值", "比较标准化后的数值集合和次数；兼容等级前缀、范围连接符、数字语序、中文序数、本地化小数和日期单元格；不确定的本地化表达降级为人工确认。"],
        ["标签", "先执行注册标签词法扫描，再分别比较标签结构与属性；color/size/font 属性差异为人工确认，url/ref/link 功能属性变化保持高风险。"],
        ["空文本", "同时检查源语言和目标语言空文本；空格、换行、零宽字符及其他不可见字符按空文本处理。"],
        ["语言覆盖", "整列为空的语言只生成一条覆盖记录，不逐行生成空翻译；活跃语言列中的零散空白仍逐行检查。"],
        ["问题分类", "明确问题、源文本问题、需人工确认、语言覆盖范围、忽略项/白名单分别输出；错误总数只统计高置信度明确问题。"],
    ]:
        append_safe_row(ws, row)
        rule_row_count += 1
    finish_sheet(ws, 2, rule_row_count)

    wb.save(path)
