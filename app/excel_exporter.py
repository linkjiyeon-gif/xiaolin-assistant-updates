from __future__ import annotations

from pathlib import Path
from typing import Iterable

from .comparer import DiffItem


def export_diffs_to_excel(diffs: Iterable[DiffItem], output_path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "差异结果"

    items = list(diffs)
    is_translation = any(
        getattr(item, "item_id", "") or getattr(item, "language", "") or getattr(item, "source_row", "") or getattr(item, "config_row", "")
        for item in items
    )

    if is_translation:
        headers = ["序号", "差异类型", "ID", "语言", "源文件行号", "配置表行号", "源文件文本", "配置表文本", "备注"]
        widths = [8, 16, 22, 14, 14, 14, 46, 46, 36]
    else:
        headers = ["序号", "差异类型", "文档内容", "文件内容", "所在位置", "备注"]
        widths = [8, 14, 45, 45, 28, 36]

    ws.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in items:
        if is_translation:
            ws.append([
                item.index,
                item.diff_type,
                getattr(item, "item_id", ""),
                getattr(item, "language", ""),
                getattr(item, "source_row", ""),
                getattr(item, "config_row", ""),
                item.document_content,
                item.file_content,
                item.remark,
            ])
        else:
            ws.append([
                item.index,
                item.diff_type,
                item.document_content,
                item.file_content,
                item.location,
                item.remark,
            ])

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def export_value_compare_diffs_to_excel(diffs: Iterable[DiffItem], output_path: str) -> None:
    """Export structured value/config compare result into QA-friendly categorized sheets."""
    items = list(diffs)
    specialized_types = {
        "奖励多配", "奖励漏配", "数量不一致", "pcid不一致", "价格不一致",
        "持续时间不一致", "积分/分数不一致", "限购次数不一致",
        "配置表缺少对应商品", "参考表无法识别奖励",
    }
    if any(item.diff_type in specialized_types for item in items):
        # 商业化专项结果沿用已验证的单 Sheet 六列报告，避免被通用分类导出器
        # 过滤掉“奖励多配/数量不一致/pcid不一致”等专项差异。
        export_diffs_to_excel(items, output_path)
        return

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    def _safe(value: object) -> str:
        if value is None:
            return ""
        return str(value)

    def _count(types: set[str]) -> int:
        return sum(1 for d in items if d.diff_type in types)

    def _split_field_value(text: str) -> tuple[str, str]:
        text = _safe(text).strip()
        if ":" in text:
            field, value = text.split(":", 1)
            return field.strip(), value.strip()
        if "：" in text:
            field, value = text.split("：", 1)
            return field.strip(), value.strip()
        return "", text

    def _location_part(location: str, keyword: str) -> str:
        location = _safe(location)
        # Common format: "参考:xxx | 配置:yyy | 主键:zzz". Keep parser tolerant.
        for sep in [" | ", "｜"]:
            for part in location.split(sep):
                if part.strip().startswith(keyword):
                    return part.split(":", 1)[1].strip() if ":" in part else part.strip()
        return ""

    def _key_from_location(location: str) -> str:
        return _location_part(location, "主键") or _location_part(location, "key") or _safe(location)

    def _ref_location(item: DiffItem) -> str:
        return _location_part(item.location, "参考") or _safe(item.location)

    def _cfg_location(item: DiffItem) -> str:
        return _location_part(item.location, "配置") or _location_part(item.location, "候选")

    def _parse_filter_summary() -> dict[str, str]:
        data: dict[str, str] = {}
        for d in items:
            if d.diff_type not in {"比对摘要", "过滤风险"}:
                continue
            for text in [d.document_content, d.file_content, d.location, d.remark]:
                text = _safe(text)
                for part in text.replace("，", "|").split("|"):
                    if ":" in part:
                        k, v = part.split(":", 1)
                        k = k.strip()
                        v = v.strip()
                        if k and v and len(k) <= 30:
                            data[k] = v
                    elif "：" in part:
                        k, v = part.split("：", 1)
                        k = k.strip()
                        v = v.strip()
                        if k and v and len(k) <= 30:
                            data[k] = v
        return data

    meta = _parse_filter_summary()

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    warn_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")

    def _make_sheet(name: str, headers: list[str], widths: list[int]):
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
        return ws

    def _finish_sheet(ws):
        if ws.max_row == 1:
            ws.append(["无"] + [""] * (ws.max_column - 1))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.auto_filter.ref = ws.dimensions

    # 1. 总览
    overview = _make_sheet("总览", ["项目", "内容"], [28, 100])
    overview_rows = [
        ("参考表文件", meta.get("参考表文件") or meta.get("参考表") or "详见比对摘要/解析诊断"),
        ("配置表文件", meta.get("配置表文件") or meta.get("配置表") or "详见比对摘要/解析诊断"),
        ("参考表 Sheet", meta.get("参考表 Sheet") or meta.get("Sheet") or "详见比对摘要"),
        ("参考表分区", meta.get("参考表分区") or meta.get("分区") or "详见解析诊断"),
        ("配置表过滤条件", meta.get("过滤条件") or meta.get("当前过滤条件") or "详见过滤风险/比对摘要"),
        ("参考表有效数据条数", meta.get("参考表有效数据条数") or meta.get("参考表有效记录") or "详见比对摘要"),
        ("配置表参与比对条数", meta.get("配置表参与比对条数") or meta.get("配置表有效记录") or "详见比对摘要"),
        ("确认差异数量", str(_count({"确认字段差异", "字段差异", "值不一致"}))),
        ("待确认数量", str(_count({"待确认字段差异", "待确认匹配项", "待确认", "高可信匹配"}))),
        ("疑似漏配数量", str(_count({"疑似漏配", "漏配"}))),
        ("疑似多配数量", str(_count({"疑似多配", "多配"}))),
        ("未映射字段数量", str(_count({"未映射字段"}))),
        ("解析风险数量", str(_count({"解析诊断", "过滤风险", "参考表解析风险", "配置表解析风险"}))),
    ]
    for row in overview_rows:
        overview.append(list(row))
    _finish_sheet(overview)

    # 2. 确认差异
    confirm = _make_sheet(
        "确认差异",
        ["问题类型", "参考表位置", "配置表位置", "主键", "字段", "参考值", "配置值", "差异说明", "建议处理"],
        [18, 34, 34, 30, 20, 44, 44, 54, 36],
    )
    for item in items:
        if item.diff_type not in {"确认字段差异", "字段差异", "值不一致"}:
            continue
        ref_field, ref_value = _split_field_value(item.document_content)
        cfg_field, cfg_value = _split_field_value(item.file_content)
        field = ref_field or cfg_field
        confirm.append([
            item.diff_type,
            _ref_location(item),
            _cfg_location(item),
            _key_from_location(item.location),
            field,
            ref_value,
            cfg_value,
            item.remark,
            "按字段口径确认后修复配置；如为策划口径变更，回写参考表。",
        ])
    _finish_sheet(confirm)

    # 3. 待确认匹配项（包含低可信匹配与待确认字段差异，避免误报为确认差异）
    pending = _make_sheet(
        "待确认匹配项",
        ["参考表位置", "候选配置", "匹配字段", "匹配分数", "待确认原因"],
        [36, 48, 34, 16, 70],
    )
    for item in items:
        if item.diff_type not in {"待确认匹配项", "待确认字段差异", "待确认", "高可信匹配"}:
            continue
        field, ref_value = _split_field_value(item.document_content)
        _, cfg_value = _split_field_value(item.file_content)
        match_field = field or item.diff_type
        score = ""
        for token in ["分数", "score", "匹配分"]:
            if token in item.remark:
                score = item.remark
                break
        pending.append([
            _ref_location(item),
            _cfg_location(item) or cfg_value or item.file_content,
            match_field,
            score,
            item.remark or f"参考值：{ref_value}；配置值：{cfg_value}",
        ])
    _finish_sheet(pending)

    # 4. 疑似漏配
    missing = _make_sheet(
        "疑似漏配",
        ["参考表位置", "参考表关键字段", "疑似漏配原因", "建议确认方式"],
        [36, 58, 70, 42],
    )
    for item in items:
        if item.diff_type not in {"疑似漏配", "漏配"}:
            continue
        missing.append([
            _ref_location(item),
            item.document_content,
            item.remark,
            "确认业务范围过滤条件、联合主键、系列名/档位/价格/pcid/reward 是否可对应。",
        ])
    _finish_sheet(missing)

    # 5. 疑似多配
    extra = _make_sheet(
        "疑似多配",
        ["配置表位置", "配置表关键字段", "疑似多配原因", "建议确认方式"],
        [36, 58, 70, 42],
    )
    for item in items:
        if item.diff_type not in {"疑似多配", "多配"}:
            continue
        extra.append([
            _cfg_location(item) or item.location,
            item.file_content,
            item.remark,
            "确认配置表过滤范围是否过宽，或该配置是否属于参考表未覆盖业务。",
        ])
    _finish_sheet(extra)

    # 6. 未映射字段
    unmapped = _make_sheet(
        "未映射字段",
        ["来源文件", "字段名", "示例值", "未映射原因", "建议映射字段"],
        [24, 28, 44, 58, 34],
    )
    for item in items:
        if item.diff_type != "未映射字段":
            continue
        field, sample = _split_field_value(item.document_content)
        source = "参考表" if "参考" in item.location or "参考" in item.remark else "配置表"
        unmapped.append([
            source,
            field or item.document_content,
            sample or item.file_content,
            item.remark,
            item.file_content,
        ])
    _finish_sheet(unmapped)

    # 7. 解析诊断
    diag = _make_sheet(
        "解析诊断",
        ["文件", "Sheet", "行号", "诊断类型", "诊断说明", "处理结果"],
        [24, 24, 14, 22, 70, 48],
    )
    for item in items:
        if item.diff_type not in {"解析诊断", "参考表解析风险", "配置表解析风险", "比对摘要", "已忽略项/规则说明"}:
            continue
        sheet = _location_part(item.location, "Sheet")
        row_no = _location_part(item.location, "行") or _location_part(item.location, "row")
        diag.append([
            _ref_location(item) or _cfg_location(item),
            sheet,
            row_no,
            item.diff_type,
            item.document_content or item.remark,
            item.file_content or item.remark,
        ])
    _finish_sheet(diag)

    # 8. 过滤风险
    risk = _make_sheet(
        "过滤风险",
        ["配置表字段", "当前过滤条件", "风险说明", "建议处理"],
        [24, 54, 70, 44],
    )
    for item in items:
        if item.diff_type != "过滤风险":
            continue
        field, condition = _split_field_value(item.document_content)
        risk.append([
            field or item.location,
            condition or item.document_content,
            item.remark or item.file_content,
            "缩小 type/father_id/ID/名称关键词等过滤条件后重新比对。",
        ])
    _finish_sheet(risk)

    # 9. 字段格式差异但语义一致 / 默认值差异说明（额外 Sheet，避免非确认项丢失）
    semantic = _make_sheet(
        "格式语义一致",
        ["类型", "参考表位置", "配置表位置", "字段", "参考值", "配置值", "说明"],
        [24, 34, 34, 20, 42, 42, 70],
    )
    for item in items:
        if item.diff_type not in {"字段格式差异但语义一致", "默认值差异但不影响配置"}:
            continue
        ref_field, ref_value = _split_field_value(item.document_content)
        cfg_field, cfg_value = _split_field_value(item.file_content)
        semantic.append([
            item.diff_type,
            _ref_location(item),
            _cfg_location(item),
            ref_field or cfg_field,
            ref_value,
            cfg_value,
            item.remark,
        ])
    _finish_sheet(semantic)

    # Visual emphasis for overview warning-like rows.
    for row in overview.iter_rows(min_row=2):
        if row[0].value and ("风险" in str(row[0].value) or "待确认" in str(row[0].value)):
            for cell in row:
                cell.fill = warn_fill

    wb.save(path)
