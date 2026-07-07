from __future__ import annotations

import csv
import json
import os
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Tuple

from .constants import TEXT_EXTENSIONS


@dataclass
class ParsedContent:
    path: str
    ext: str
    text: str
    lines: List[str]
    fields: Dict[str, str] = field(default_factory=dict)
    status: str = "读取成功"


class FileReadError(Exception):
    pass


KEY_VALUE_PATTERN = re.compile(r"^\s*([^:=：\t]{1,120})\s*[:=：\t]\s*(.*?)\s*$")


def read_file(path: str) -> ParsedContent:
    if not path:
        raise FileReadError("文件路径为空")

    file_path = Path(path)
    if not file_path.exists():
        raise FileReadError(f"文件不存在：{path}")

    ext = file_path.suffix.lower()

    try:
        if ext in TEXT_EXTENSIONS:
            text = read_text_with_fallback(file_path)
            fields = extract_key_values_from_text(text)
        elif ext == ".json":
            text, fields = read_json(file_path)
        elif ext == ".csv":
            text, fields = read_csv(file_path)
        elif ext == ".xlsx":
            text, fields = read_xlsx(file_path)
        elif ext == ".xml":
            text, fields = read_xml(file_path)
        elif ext in {".yaml", ".yml"}:
            text, fields = read_yaml(file_path)
        elif ext == ".docx":
            text, fields = read_docx(file_path)
        elif ext == ".pdf":
            text, fields = read_pdf(file_path)
        else:
            text = read_text_with_fallback(file_path)
            fields = extract_key_values_from_text(text)
    except FileReadError:
        raise
    except Exception as exc:
        raise FileReadError(f"读取失败：{exc}") from exc

    lines = split_lines(text)
    return ParsedContent(path=str(file_path), ext=ext, text=text, lines=lines, fields=fields)


def read_text_with_fallback(path: Path) -> str:
    encodings = ["utf-8-sig", "utf-8", "gb18030", "gbk", "cp936", "latin-1"]
    last_error = None
    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise FileReadError(f"文本编码无法识别：{last_error}")


def split_lines(text: str) -> List[str]:
    return [line.rstrip("\n\r") for line in text.splitlines()]


def extract_key_values_from_text(text: str) -> Dict[str, str]:
    fields: Dict[str, str] = {}
    for index, line in enumerate(split_lines(text), start=1):
        match = KEY_VALUE_PATTERN.match(line)
        if not match:
            continue
        key = clean_key(match.group(1))
        value = match.group(2).strip()
        if key:
            fields[key] = value
    return fields


def clean_key(key: str) -> str:
    return re.sub(r"\s+", " ", str(key).strip())


def flatten_data(data, prefix: str = "") -> Dict[str, str]:
    fields: Dict[str, str] = {}

    if isinstance(data, dict):
        for key, value in data.items():
            next_key = f"{prefix}.{key}" if prefix else str(key)
            fields.update(flatten_data(value, next_key))
    elif isinstance(data, list):
        for index, value in enumerate(data):
            next_key = f"{prefix}[{index}]" if prefix else f"[{index}]"
            fields.update(flatten_data(value, next_key))
    else:
        if prefix:
            fields[prefix] = "" if data is None else str(data)

    return fields


def fields_to_text(fields: Dict[str, str]) -> str:
    return "\n".join(f"{key}: {value}" for key, value in fields.items())


def read_json(path: Path) -> Tuple[str, Dict[str, str]]:
    raw = read_text_with_fallback(path)
    data = json.loads(raw)
    fields = flatten_data(data)
    text = json.dumps(data, ensure_ascii=False, indent=2)
    return text, fields


def read_yaml(path: Path) -> Tuple[str, Dict[str, str]]:
    raw = read_text_with_fallback(path)
    try:
        import yaml
    except ImportError:
        return raw, extract_key_values_from_text(raw)

    data = yaml.safe_load(raw)
    if data is None:
        return raw, {}
    fields = flatten_data(data)
    fields.update(extract_key_values_from_text(raw))
    return raw, fields


def read_csv(path: Path) -> Tuple[str, Dict[str, str]]:
    raw = read_text_with_fallback(path)
    sample = raw[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel

    rows = list(csv.reader(raw.splitlines(), dialect))
    text_lines = ["\t".join(row) for row in rows]
    fields: Dict[str, str] = {}

    if rows:
        headers = [clean_key(cell) for cell in rows[0]]
        for row_index, row in enumerate(rows[1:], start=2):
            for col_index, value in enumerate(row):
                col_name = headers[col_index] if col_index < len(headers) and headers[col_index] else f"col_{col_index + 1}"
                fields[f"row_{row_index}.{col_name}"] = value
                if col_name not in fields:
                    fields[col_name] = value

    return "\n".join(text_lines), fields


def read_xlsx(path: Path) -> Tuple[str, Dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FileReadError("缺少依赖 openpyxl，请先安装 requirements.txt") from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    text_lines: List[str] = []
    fields: Dict[str, str] = {}

    for ws in wb.worksheets:
        text_lines.append(f"[Sheet] {ws.title}")
        try:
            ws.reset_dimensions()
        except Exception:
            pass
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [clean_key(cell) if cell is not None else "" for cell in rows[0]]
        for row_index, row in enumerate(rows, start=1):
            values = ["" if value is None else str(value) for value in row]
            text_lines.append("\t".join(values))
            for col_index, value in enumerate(values, start=1):
                if value == "":
                    continue
                cell_key = f"{ws.title}!R{row_index}C{col_index}"
                fields[cell_key] = value
                if row_index > 1:
                    header = headers[col_index - 1] if col_index - 1 < len(headers) and headers[col_index - 1] else f"col_{col_index}"
                    fields[f"{ws.title}.row_{row_index}.{header}"] = value

    wb.close()
    return "\n".join(text_lines), fields


def read_xml(path: Path) -> Tuple[str, Dict[str, str]]:
    raw = read_text_with_fallback(path)
    root = ET.fromstring(raw)
    fields: Dict[str, str] = {}

    def walk(node, prefix: str):
        tag = strip_namespace(node.tag)
        current = f"{prefix}.{tag}" if prefix else tag
        text_value = (node.text or "").strip()
        if text_value:
            fields[current] = text_value
        for attr_key, attr_value in node.attrib.items():
            fields[f"{current}@{attr_key}"] = attr_value
        for child in node:
            walk(child, current)

    walk(root, "")
    fields.update(extract_key_values_from_text(raw))
    return raw, fields


def strip_namespace(tag: str) -> str:
    return tag.split("}", 1)[-1] if "}" in tag else tag


def read_docx(path: Path) -> Tuple[str, Dict[str, str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise FileReadError("缺少依赖 python-docx，请先安装 requirements.txt") from exc

    document = Document(path)
    text_lines: List[str] = []
    fields: Dict[str, str] = {}

    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            text_lines.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        text_lines.append(f"[Table {table_index}]")
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            text_lines.append("\t".join(values))
            if len(values) >= 2 and values[0]:
                fields[clean_key(values[0])] = values[1]
                fields[f"table_{table_index}.row_{row_index}.{clean_key(values[0])}"] = values[1]
            for col_index, value in enumerate(values, start=1):
                if value:
                    fields[f"table_{table_index}.R{row_index}C{col_index}"] = value

    text = "\n".join(text_lines)
    fields.update(extract_key_values_from_text(text))
    return text, fields


def read_pdf(path: Path) -> Tuple[str, Dict[str, str]]:
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise FileReadError("缺少依赖 PyPDF2，请先安装 requirements.txt") from exc

    reader = PdfReader(str(path))
    text_lines: List[str] = []
    for page_index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            text_lines.append(f"[Page {page_index}]")
            text_lines.append(page_text.strip())

    text = "\n".join(text_lines)
    if not text.strip():
        raise FileReadError("PDF 未读取到可复制文本，扫描版 PDF 暂不支持")
    fields = extract_key_values_from_text(text)
    return text, fields


def get_file_info(path: str) -> str:
    if not path:
        return "未选择"
    file_path = Path(path)
    if not file_path.exists():
        return "文件不存在"
    size_kb = os.path.getsize(path) / 1024
    return f"{file_path.suffix.lower() or '无扩展名'} / {size_kb:.1f} KB"
