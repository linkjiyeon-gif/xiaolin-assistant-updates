from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from .translation_compare import safe_int
from .config_rule_templates import TYPE_TOKENS


DEFAULT_KEY_COLUMNS = ["ID", "id", "Key", "key", "config_id", "item_id", "task_id", "reward_id"]
BOOL_VALUES = {"0", "1", "true", "false", "TRUE", "FALSE", "True", "False"}
TIME_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d"]

KEY_INDEX_FIELD_NAMES = {"#id", "id", "#key", "key", "no", "index", "idx", "uid", "guid"}
REWARD_FIELD_EXACT_NAMES = {
    "reward", "rewards", "content", "treasure", "mail_reward", "drop_reward",
    "item_reward", "mall_reward", "extra_reward", "mall_extrareward",
    "gift_reward", "award", "awards",
}
COMPOUND_FIELD_EXAMPLES = "reward、content、treasure、restrict、condition"


def normalize_config_field_name(field_name: str) -> str:
    return str(field_name or "").strip().lower()


def is_key_index_field(field_name: str) -> bool:
    return normalize_config_field_name(field_name) in KEY_INDEX_FIELD_NAMES


def is_reward_like_field(field_name: str) -> bool:
    name = normalize_config_field_name(field_name)
    return bool(name and (name in REWARD_FIELD_EXACT_NAMES or "reward" in name))


def resolve_header_name(headers: Sequence[str], field_name: str) -> str:
    field = str(field_name or "").strip()
    if not field:
        return ""
    if field in headers:
        return field
    lower_map = {str(h).strip().lower(): h for h in headers}
    return lower_map.get(field.lower(), field)


def find_reward_like_field(headers: Sequence[str]) -> str:
    safe_headers = [h for h in headers if str(h or "").strip() and not is_key_index_field(h)]
    for header in safe_headers:
        if normalize_config_field_name(header) == "reward":
            return header
    for header in safe_headers:
        if "reward" in normalize_config_field_name(header):
            return header
    for header in safe_headers:
        if normalize_config_field_name(header) in {"content", "treasure", "award", "awards"}:
            return header
    return ""


@dataclass
class ConfigRecord:
    row_number: int
    values: Dict[str, str]


@dataclass
class ConfigTable:
    file_path: str
    file_name: str
    sheet_name: str
    headers: List[str]
    records: List[ConfigRecord]
    key_column: str
    field_types: Dict[str, str] = field(default_factory=dict)
    field_comments: Dict[str, str] = field(default_factory=dict)
    header_row: int = 1
    data_start_row: int = 2
    type_row: int = 0
    description_row: int = 0


@dataclass
class ValidationIssue:
    index: int
    issue_type: str
    file_name: str
    sheet: str
    row_number: str
    field_name: str
    field_type: str
    item_id: str
    current_value: str
    remark: str
    suggestion: str


@dataclass
class ConfigValidationOptions:
    header_row: int = 1
    data_start_row: int = 2
    key_column: str = ""
    sheet_name: str = ""
    description_row: int = 1
    type_row: int = 2
    check_type_fields: bool = False
    check_compound_rules: bool = False
    compound_rules_text: str = ""
    required_fields_text: str = ""
    numeric_fields_text: str = ""
    range_rules_text: str = ""
    boolean_fields_text: str = ""
    enum_rules_text: str = ""
    time_fields_text: str = ""
    time_range_rules_text: str = ""
    reward_fields_text: str = ""
    reference_rules_text: str = ""
    interval_rules_text: str = ""
    required_columns_text: str = ""
    check_duplicate_key: bool = True
    check_empty_key: bool = True
    check_required_fields: bool = True
    check_numeric_fields: bool = True
    check_range: bool = True
    check_boolean_fields: bool = True
    check_enum_fields: bool = True
    check_time_fields: bool = True
    check_time_range: bool = True
    check_reward_fields: bool = True
    check_reference_ids: bool = True
    check_interval: bool = True
    check_empty_rows: bool = True
    check_required_columns: bool = True
    ui_limit: int = 3000


class ConfigValidationError(Exception):
    pass


def split_names(text: str) -> List[str]:
    if not text:
        return []
    parts = re.split(r"[,，;；\n\r\t]+", text)
    return [p.strip() for p in parts if p.strip()]


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    return _cell_to_text(value)


def _dedupe_headers(headers: List[str]) -> List[str]:
    result: List[str] = []
    seen: Dict[str, int] = {}
    for index, header in enumerate(headers, start=1):
        name = header.strip() if header else f"列{index}"
        if name in seen:
            seen[name] += 1
            name = f"{name}__{seen[name]}"
        else:
            seen[name] = 1
        result.append(name)
    return result



def _build_header_meta(headers: Sequence[str], rows: Sequence[Sequence[Any]], row_number: int) -> Dict[str, str]:
    """Build per-field metadata from a 1-based row number."""
    result: Dict[str, str] = {}
    if not row_number or row_number < 1 or len(rows) < row_number:
        return result
    row = rows[row_number - 1]
    for index, header in enumerate(headers):
        result[header] = _cell_to_text(row[index]) if index < len(row) else ""
    return result


def _looks_like_type_row(row: Sequence[Any]) -> int:
    score = 0
    for cell in row:
        text = _cell_to_text(cell).lower()
        if text in TYPE_TOKENS:
            score += 1
    return score


def _looks_like_field_row(row: Sequence[Any]) -> int:
    score = 0
    for cell in row:
        text = _cell_to_text(cell).strip()
        low = text.lower()
        if not text:
            continue
        if text.startswith("#") or text.startswith("$") or text.startswith("^"):
            score += 2
        if low in {"id", "key", "#id", "#key", "reward", "content", "restrict", "condition"}:
            score += 3
        elif re.fullmatch(r"[#$^]?[a-zA-Z_][a-zA-Z0-9_]*", text):
            score += 1
    return score


def read_config_raw_rows(path: str, sheet_name: str = "") -> Tuple[List[List[str]], str]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".csv":
        return _read_csv_rows(p), "CSV"
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(p, read_only=True, data_only=True)
        except Exception as exc:
            raise ConfigValidationError(f"读取 xlsx 失败：{exc}") from exc
        try:
            if sheet_name and sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            elif sheet_name:
                raise ConfigValidationError(f"找不到 Sheet：{sheet_name}")
            else:
                ws = wb[wb.sheetnames[0]]
            return [[_cell_to_text(v) for v in row] for row in ws.iter_rows(values_only=True)], ws.title
        finally:
            try:
                wb.close()
            except Exception:
                pass
    raise ConfigValidationError(f"不支持自动识别的配置表格式：{ext}")


def detect_config_structure(path: str, sheet_name: str = "") -> Dict[str, Any]:
    rows, sheet = read_config_raw_rows(path, sheet_name)
    if not rows:
        raise ConfigValidationError("配置表为空，无法识别结构")
    max_scan = min(8, len(rows))
    type_scores = [(i + 1, _looks_like_type_row(rows[i])) for i in range(max_scan)]
    type_row = max(type_scores, key=lambda x: x[1])[0] if type_scores else 0
    if max((s for _, s in type_scores), default=0) < 2:
        type_row = 0
    field_scores = [(i + 1, _looks_like_field_row(rows[i])) for i in range(max_scan)]
    field_row = max(field_scores, key=lambda x: x[1])[0] if field_scores else 1
    if max((s for _, s in field_scores), default=0) < 2:
        field_row = type_row + 1 if type_row else 1
    data_start_row = field_row + 1
    description_row = 1 if field_row > 1 else 0
    headers = _dedupe_headers([_normalize_header(v) for v in rows[field_row - 1]]) if len(rows) >= field_row else []
    key_column = detect_key_column(headers, "")
    field_types = _build_header_meta(headers, rows, type_row)
    field_comments = _build_header_meta(headers, rows, description_row)
    data_rows = max(0, len(rows) - data_start_row + 1)
    empty_cells = 0
    if headers:
        for row in rows[data_start_row - 1:]:
            for idx in range(len(headers)):
                value = _cell_to_text(row[idx]) if idx < len(row) else ""
                if value == "":
                    empty_cells += 1
    return {
        "sheet_name": sheet,
        "description_row": description_row,
        "type_row": type_row,
        "header_row": field_row,
        "data_start_row": data_start_row,
        "key_column": key_column,
        "headers": headers,
        "field_types": field_types,
        "field_comments": field_comments,
        "data_rows": data_rows,
        "type_field_count": sum(1 for v in field_types.values() if v),
        "empty_cell_count": empty_cells,
    }

def detect_key_column(headers: Sequence[str], preferred: str = "") -> str:
    if preferred and preferred in headers:
        return preferred
    lower_map = {h.lower(): h for h in headers}
    if preferred and preferred.lower() in lower_map:
        return lower_map[preferred.lower()]
    for cand in DEFAULT_KEY_COLUMNS:
        if cand in headers:
            return cand
        if cand.lower() in lower_map:
            return lower_map[cand.lower()]
    return headers[0] if headers else ""


def parse_config_table(path: str, header_row: int = 1, data_start_row: int = 2, key_column: str = "", sheet_name: str = "", type_row: int = 0, description_row: int = 0) -> ConfigTable:
    p = Path(path)
    if not p.exists():
        raise ConfigValidationError(f"文件不存在：{path}")
    ext = p.suffix.lower()
    header_row = max(1, int(header_row or 1))
    data_start_row = max(1, int(data_start_row or header_row + 1))
    if ext == ".xlsx":
        return _parse_xlsx(p, header_row, data_start_row, key_column, sheet_name, type_row, description_row)
    if ext == ".csv":
        return _parse_csv(p, header_row, data_start_row, key_column, type_row, description_row)
    if ext == ".json":
        return _parse_json(p, key_column)
    raise ConfigValidationError(f"不支持的配置表格式：{ext}")


def _parse_xlsx(path: Path, header_row: int, data_start_row: int, key_column: str, sheet_name: str, type_row: int = 0, description_row: int = 0) -> ConfigTable:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ConfigValidationError(f"读取 xlsx 失败：{exc}") from exc
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif sheet_name:
            raise ConfigValidationError(f"找不到 Sheet：{sheet_name}")
        else:
            ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < header_row:
            raise ConfigValidationError("表头行超出文件范围")
        headers = _dedupe_headers([_normalize_header(v) for v in rows[header_row - 1]])
        key_col = detect_key_column(headers, key_column)
        records: List[ConfigRecord] = []
        for idx, row in enumerate(rows[data_start_row - 1 :], start=data_start_row):
            values = {headers[i]: _cell_to_text(row[i]) if i < len(row) else "" for i in range(len(headers))}
            records.append(ConfigRecord(row_number=idx, values=values))
        field_types = _build_header_meta(headers, rows, type_row)
        field_comments = _build_header_meta(headers, rows, description_row)
        return ConfigTable(str(path), path.name, ws.title, headers, records, key_col, field_types, field_comments, header_row, data_start_row, type_row, description_row)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _read_csv_rows(path: Path) -> List[List[str]]:
    encodings = ["utf-8-sig", "utf-8", "gbk", "gb18030"]
    last_exc: Optional[Exception] = None
    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    # 不把 | 当作 CSV 列分隔符：配置表 reward/restrict 等字段常用 | 作为单元格内部分隔符
                    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
                except Exception:
                    dialect = csv.excel
                return [list(row) for row in csv.reader(f, dialect)]
        except UnicodeDecodeError as exc:
            last_exc = exc
            continue
        except Exception as exc:
            last_exc = exc
            break
    raise ConfigValidationError(f"读取 csv 失败：{last_exc}")


def _parse_csv(path: Path, header_row: int, data_start_row: int, key_column: str, type_row: int = 0, description_row: int = 0) -> ConfigTable:
    rows = _read_csv_rows(path)
    if len(rows) < header_row:
        raise ConfigValidationError("表头行超出文件范围")
    headers = _dedupe_headers([_normalize_header(v) for v in rows[header_row - 1]])
    key_col = detect_key_column(headers, key_column)
    records: List[ConfigRecord] = []
    for idx, row in enumerate(rows[data_start_row - 1 :], start=data_start_row):
        values = {headers[i]: _cell_to_text(row[i]) if i < len(row) else "" for i in range(len(headers))}
        records.append(ConfigRecord(row_number=idx, values=values))
    field_types = _build_header_meta(headers, rows, type_row)
    field_comments = _build_header_meta(headers, rows, description_row)
    return ConfigTable(str(path), path.name, "CSV", headers, records, key_col, field_types, field_comments, header_row, data_start_row, type_row, description_row)


def _flatten_dict(data: Dict[str, Any], prefix: str = "") -> Dict[str, str]:
    result: Dict[str, str] = {}
    for key, value in data.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            result.update(_flatten_dict(value, name))
        elif isinstance(value, list):
            result[name] = json.dumps(value, ensure_ascii=False)
        else:
            result[name] = _cell_to_text(value)
    return result


def _parse_json(path: Path, key_column: str) -> ConfigTable:
    try:
        with path.open("r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except UnicodeDecodeError:
        with path.open("r", encoding="gbk") as f:
            data = json.load(f)
    except Exception as exc:
        raise ConfigValidationError(f"读取 json 失败：{exc}") from exc

    if isinstance(data, dict):
        if "data" in data and isinstance(data["data"], list):
            data = data["data"]
        elif "items" in data and isinstance(data["items"], list):
            data = data["items"]
        else:
            data = [data]
    if not isinstance(data, list):
        raise ConfigValidationError("JSON 根节点需要是对象或对象数组")
    flat_rows: List[Dict[str, str]] = []
    for item in data:
        if isinstance(item, dict):
            flat_rows.append(_flatten_dict(item))
        else:
            flat_rows.append({"value": _cell_to_text(item)})
    headers: List[str] = []
    seen = set()
    for row in flat_rows:
        for key in row.keys():
            if key not in seen:
                headers.append(key)
                seen.add(key)
    key_col = detect_key_column(headers, key_column)
    records = [ConfigRecord(row_number=i + 1, values={h: row.get(h, "") for h in headers}) for i, row in enumerate(flat_rows)]
    return ConfigTable(str(path), path.name, "JSON", headers, records, key_col, {}, {}, 1, 1, 0, 0)


def _is_number(value: str) -> bool:
    value = value.strip()
    if value == "":
        return True
    return re.fullmatch(r"[-+]?\d+(\.\d+)?", value) is not None


def _parse_range_rules(text: str) -> Dict[str, Tuple[Optional[float], Optional[float]]]:
    rules: Dict[str, Tuple[Optional[float], Optional[float]]] = {}
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        parts = re.split(r"[:=,，|]+", line)
        if len(parts) < 2:
            continue
        field = parts[0].strip()
        min_v = None
        max_v = None
        try:
            min_v = float(parts[1].strip()) if len(parts) >= 2 and parts[1].strip() != "" else None
        except Exception:
            min_v = None
        try:
            max_v = float(parts[2].strip()) if len(parts) >= 3 and parts[2].strip() != "" else None
        except Exception:
            max_v = None
        if field:
            rules[field] = (min_v, max_v)
    return rules


def _parse_enum_rules(text: str) -> Dict[str, set]:
    rules: Dict[str, set] = {}
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if "=" in line:
            field, values = line.split("=", 1)
        elif ":" in line:
            field, values = line.split(":", 1)
        else:
            continue
        allowed = {v.strip() for v in re.split(r"[,，;；|]", values) if v.strip()}
        if field.strip() and allowed:
            rules[field.strip()] = allowed
    return rules


def _parse_pair_rules(text: str) -> List[Tuple[str, str]]:
    rules: List[Tuple[str, str]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        parts = re.split(r"[,，;；|]+", line)
        if len(parts) >= 2 and parts[0].strip() and parts[1].strip():
            rules.append((parts[0].strip(), parts[1].strip()))
    return rules


def _parse_reference_rules(text: str) -> List[Dict[str, Any]]:
    rules = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if "=" not in line:
            continue
        field, rest = line.split("=", 1)
        parts = [p.strip() for p in rest.split("|")]
        if not field.strip() or not parts or not parts[0]:
            continue
        rules.append({
            "field": field.strip(),
            "path": parts[0],
            "key_column": parts[1] if len(parts) > 1 else "",
            "sheet_name": parts[2] if len(parts) > 2 else "",
            "header_row": safe_int(parts[3], 1) if len(parts) > 3 else 1,
            "data_start_row": safe_int(parts[4], 2) if len(parts) > 4 else 2,
        })
    return rules


def _parse_time(value: str) -> Optional[datetime]:
    v = value.strip()
    if v == "":
        return None
    if re.fullmatch(r"\d{10}", v):
        try:
            return datetime.fromtimestamp(int(v))
        except Exception:
            return None
    if re.fullmatch(r"\d{13}", v):
        try:
            return datetime.fromtimestamp(int(v) / 1000)
        except Exception:
            return None
    for fmt in TIME_FORMATS:
        try:
            return datetime.strptime(v, fmt)
        except Exception:
            continue
    return None


def _is_valid_time(value: str) -> bool:
    return value.strip() == "" or _parse_time(value) is not None


def _reward_error(value: str) -> str:
    v = value.strip()
    if not v:
        return ""
    # normalize common item separators into groups; group separator is semicolon.
    groups = [g.strip() for g in re.split(r"[;；]", v) if g.strip()]
    for group in groups:
        if ":" in group:
            parts = [p.strip() for p in group.split(":", 1)]
        elif "|" in group:
            parts = [p.strip() for p in group.split("|", 1)]
        elif "," in group or "，" in group:
            parts = [p.strip() for p in re.split(r"[,，]", group, maxsplit=1)]
        else:
            return f"奖励项缺少分隔符：{group}"
        if len(parts) < 2 or not parts[0] or not parts[1]:
            return f"奖励 ID 或数量为空：{group}"
        if not _is_number(parts[1]) or float(parts[1]) < 0:
            return f"奖励数量不是合法非负数字：{group}"
    return ""


def _is_empty_row(record: ConfigRecord) -> bool:
    return all(str(v).strip() == "" for v in record.values.values())


def _is_empty_value(value: str) -> bool:
    return value.strip() in {"", "None", "none", "NULL", "null"}


def _is_integer(value: str, allow_negative: bool = True) -> bool:
    v = value.strip()
    if not v:
        return False
    pattern = r"[-+]?\d+" if allow_negative else r"\d+"
    return re.fullmatch(pattern, v) is not None


def _type_value_error(field_type: str, value: str) -> str:
    ftype = (field_type or "").strip().lower()
    v = value.strip()
    if not ftype:
        return ""
    if ftype == "uint":
        if _is_empty_value(v):
            return "uint 类型字段不能为空"
        if not _is_integer(v, allow_negative=False):
            return "uint 类型字段必须为非负整数，不能配置为空、小数、负数或文本"
    if ftype == "int":
        if _is_empty_value(v):
            return "int 类型字段不能为空"
        if not _is_integer(v, allow_negative=True):
            return "int 类型字段必须为整数"
    if ftype == "float":
        if _is_empty_value(v):
            return "float 类型字段不能为空"
        if not _is_number(v):
            return "float 类型字段必须为合法整数或小数"
    if ftype == "bool":
        if _is_empty_value(v):
            return "bool 类型字段不能为空"
        if v not in BOOL_VALUES:
            return "bool 类型字段只支持 0、1、true、false"
    if ftype in {"uint[]", "int[]"}:
        if _is_empty_value(v):
            return ""
        parts = [p.strip() for p in re.split(r"[|,，;；]", v)]
        for part in parts:
            if part == "":
                return f"{ftype} 数组字段存在空元素"
            if ftype == "uint[]" and not _is_integer(part, allow_negative=False):
                return f"uint[] 数组字段每个元素必须为非负整数，异常元素：{part}"
            if ftype == "int[]" and not _is_integer(part, allow_negative=True):
                return f"int[] 数组字段每个元素必须为整数，异常元素：{part}"
    return ""


def _parse_bool_text(text: str, default: bool = False) -> bool:
    if text is None:
        return default
    v = str(text).strip().lower()
    if v in {"1", "true", "yes", "y", "是", "允许", "可", "开启"}:
        return True
    if v in {"0", "false", "no", "n", "否", "不", "禁止", "关闭"}:
        return False
    return default


def _parse_compound_rule_line(line: str) -> Optional[Dict[str, Any]]:
    raw = line.strip()
    if not raw:
        return None
    # 简写：reward:4 或 reward=4
    m = re.fullmatch(r"([#\$\^]?[\w\u4e00-\u9fff]+)\s*[:=]\s*(\d+)", raw)
    if m:
        return {"field": m.group(1), "allow_empty": True, "group_sep": "|", "param_sep": "*", "count": int(m.group(2)), "element_type": "整数", "allow_zero": True, "allow_empty_param": False, "forbid_values": []}
    data: Dict[str, str] = {}
    for part in re.split(r"[;；]\s*", raw):
        if not part.strip():
            continue
        if "=" in part:
            k, v = part.split("=", 1)
        elif "：" in part:
            k, v = part.split("：", 1)
        elif ":" in part:
            k, v = part.split(":", 1)
        else:
            continue
        data[k.strip().lower()] = v.strip()
    field = data.get("字段") or data.get("field") or data.get("字段名")
    if not field:
        return None
    count_text = data.get("数量") or data.get("参数数量") or data.get("count") or data.get("每组参数数量") or "0"
    try:
        count = int(count_text)
    except Exception:
        count = 0
    if count <= 0:
        return None
    forbid = data.get("禁止值") or data.get("forbid") or ""
    forbid_values = [v.strip() for v in re.split(r"[,，]", forbid) if v.strip()]
    group_sep = data.get("一级") or data.get("一级分隔符") or data.get("group") or data.get("group_sep") or "|"
    param_sep = data.get("二级") or data.get("二级分隔符") or data.get("sep") or data.get("param_sep") or "*"
    if group_sep in {"无", "空", "none", "None", "NULL", "null"}:
        group_sep = ""
    return {
        "field": field,
        "allow_empty": _parse_bool_text(data.get("允许空") or data.get("允许为空") or data.get("allow_empty"), True),
        "group_sep": group_sep,
        "param_sep": param_sep,
        "count": count,
        "element_type": data.get("类型") or data.get("元素类型") or data.get("type") or "整数",
        "allow_zero": _parse_bool_text(data.get("允许0") or data.get("允许零") or data.get("allow_zero"), True),
        "allow_empty_param": _parse_bool_text(data.get("允许空参数") or data.get("allow_empty_param"), False),
        "forbid_values": forbid_values,
    }


def _parse_compound_rules(text: str) -> List[Dict[str, Any]]:
    rules: List[Dict[str, Any]] = []
    for line in (text or "").splitlines():
        rule = _parse_compound_rule_line(line)
        if rule:
            rules.append(rule)
    return rules


def _compound_value_error(value: str, rule: Dict[str, Any]) -> str:
    v = value.strip()
    if _is_empty_value(v):
        return "" if rule.get("allow_empty", True) else "字段不能为空"
    if v in set(rule.get("forbid_values") or []):
        return f"字段不允许填写单值 {v}"
    group_sep = rule.get("group_sep", "|")
    param_sep = rule.get("param_sep", "*")
    count = int(rule.get("count") or 0)
    groups = [v] if not group_sep else v.split(group_sep)
    for group in groups:
        group_text = group.strip()
        if group_text == "":
            return f"复合字段存在空组：{value}"
        parts = [group_text] if not param_sep else group_text.split(param_sep)
        if len(parts) != count:
            if is_reward_like_field(str(rule.get("field", ""))):
                return f"每组奖励参数应为 {count} 个，当前仅识别到 {len(parts)} 个，异常组：{group_text}"
            return f"每组参数应为 {count} 个，当前仅识别到 {len(parts)} 个，异常组：{group_text}"
        for part in parts:
            part_text = part.strip()
            if part_text == "" and not rule.get("allow_empty_param", False):
                return f"参数不能为空，异常组：{group_text}"
            if part_text == "":
                continue
            element_type = str(rule.get("element_type", "整数")).lower()
            if element_type in {"整数", "int", "uint", "数字"}:
                allow_negative = element_type == "int"
                if not _is_integer(part_text, allow_negative=allow_negative):
                    return f"参数必须为整数，异常参数：{part_text}"
                if not rule.get("allow_zero", True) and int(part_text) == 0:
                    return f"参数不允许为 0，异常组：{group_text}"
    return ""



def validate_config_table(path: str, options: ConfigValidationOptions, stop_event=None) -> Tuple[List[ValidationIssue], ConfigTable]:
    table = parse_config_table(path, options.header_row, options.data_start_row, options.key_column, options.sheet_name, options.type_row, options.description_row)
    issues: List[ValidationIssue] = []

    def stopped() -> bool:
        return bool(stop_event is not None and getattr(stop_event, "is_set", lambda: False)())

    def add(issue_type: str, record: Optional[ConfigRecord], field_name: str, value: str, remark: str, suggestion: str):
        issues.append(ValidationIssue(
            index=len(issues) + 1,
            issue_type=issue_type,
            file_name=table.file_name,
            sheet=table.sheet_name,
            row_number="" if record is None else str(record.row_number),
            field_name=field_name,
            field_type=table.field_types.get(field_name, ""),
            item_id="" if record is None else record.values.get(table.key_column, ""),
            current_value=value,
            remark=remark,
            suggestion=suggestion,
        ))

    required_fields = split_names(options.required_fields_text)
    numeric_fields = split_names(options.numeric_fields_text)
    boolean_fields = split_names(options.boolean_fields_text)
    time_fields = split_names(options.time_fields_text)
    reward_fields = split_names(options.reward_fields_text)
    required_columns = split_names(options.required_columns_text)
    range_rules = _parse_range_rules(options.range_rules_text)
    enum_rules = _parse_enum_rules(options.enum_rules_text)
    time_range_rules = _parse_pair_rules(options.time_range_rules_text)
    interval_rules = _parse_pair_rules(options.interval_rules_text)
    reference_rules = _parse_reference_rules(options.reference_rules_text)
    compound_rules = _parse_compound_rules(options.compound_rules_text)

    if options.check_compound_rules:
        if not compound_rules:
            raise ConfigValidationError(f"请先选择要校验的复合字段，例如 {COMPOUND_FIELD_EXAMPLES}")
        for rule in compound_rules:
            field_name = str(rule.get("field", "")).strip()
            if not field_name:
                raise ConfigValidationError(f"请先选择要校验的复合字段，例如 {COMPOUND_FIELD_EXAMPLES}")
            actual_field = resolve_header_name(table.headers, field_name)
            if is_key_index_field(field_name) or is_key_index_field(actual_field):
                raise ConfigValidationError(f"该字段是主键 / 索引字段，不适用于复合字段参数个数校验，请选择 {COMPOUND_FIELD_EXAMPLES} 等复合字段")

    if options.check_required_columns:
        for field_name in required_columns:
            if field_name not in table.headers:
                add("字段缺失", None, field_name, "", f"配置表缺少字段：{field_name}", "补充字段或检查表头行设置")

    if not table.key_column or table.key_column not in table.headers:
        add("字段缺失", None, options.key_column or "主键列", "", "无法识别主键列", "填写正确的主键列名或检查表头行")
        return issues, table

    if options.check_duplicate_key or options.check_empty_key:
        seen: Dict[str, List[ConfigRecord]] = {}
        for record in table.records:
            key = record.values.get(table.key_column, "").strip()
            if options.check_empty_key and key == "":
                add("主键为空", record, table.key_column, key, "主键 ID 为空", "填写唯一主键 ID")
            if key:
                seen.setdefault(key, []).append(record)
        if options.check_duplicate_key:
            for key, records in seen.items():
                if len(records) > 1:
                    for record in records:
                        add("主键重复", record, table.key_column, key, f"主键 ID 重复：{key}", "保持主键唯一")

    ref_sets: Dict[str, set] = {}
    if options.check_reference_ids:
        for rule in reference_rules:
            try:
                ref_table = parse_config_table(rule["path"], rule["header_row"], rule["data_start_row"], rule["key_column"], rule["sheet_name"])
                ref_sets[rule["field"]] = {r.values.get(ref_table.key_column, "").strip() for r in ref_table.records if r.values.get(ref_table.key_column, "").strip()}
            except Exception as exc:
                add("引用 ID 不存在", None, rule["field"], rule["path"], f"参考表读取失败：{exc}", "检查引用规则路径、主键列和 Sheet")

    for record in table.records:
        if stopped():
            break
        if options.check_empty_rows and _is_empty_row(record):
            add("空行", record, "", "", "数据区域存在整行为空", "删除空行或确认数据起始行")
            continue

        if options.check_type_fields:
            for field_name, field_type in table.field_types.items():
                if field_name in table.headers:
                    value = record.values.get(field_name, "")
                    err = _type_value_error(field_type, value)
                    if err:
                        add("类型行校验错误", record, field_name, value, f"{err}，当前字段 {field_name} 的值为：{value or '空'}", "按字段类型补充合法值，或确认类型行是否配置正确")

        if options.check_compound_rules:
            for rule in compound_rules:
                field_name = resolve_header_name(table.headers, str(rule.get("field", "")))
                if field_name in table.headers and not is_key_index_field(field_name):
                    rule_for_check = dict(rule)
                    rule_for_check["field"] = field_name
                    value = record.values.get(field_name, "")
                    err = _compound_value_error(value, rule_for_check)
                    if err:
                        count = rule.get("count", "")
                        if is_reward_like_field(field_name) and str(count) == "4":
                            suggestion = f"补齐 {field_name} 参数，确认格式是否为 type*id*num*value"
                        else:
                            suggestion = f"按规则补齐 {field_name} 字段格式：每组 {count} 个参数，一级分隔符 {rule.get('group_sep') or '无'}，二级分隔符 {rule.get('param_sep') or '无'}"
                        add("参数个数错误" if "参数" in err else "复合字段格式错误", record, field_name, value, f"{field_name} {err}", suggestion)

        if options.check_required_fields:
            for field_name in required_fields:
                if field_name in table.headers and record.values.get(field_name, "").strip() == "":
                    add("必填字段为空", record, field_name, "", f"必填字段为空：{field_name}", "补充字段值")

        if options.check_numeric_fields:
            for field_name in numeric_fields:
                if field_name in table.headers:
                    value = record.values.get(field_name, "")
                    if not _is_number(value):
                        add("数值格式错误", record, field_name, value, f"字段 {field_name} 不是合法数字", "改为整数、小数或负数格式")

        if options.check_range:
            for field_name, (min_v, max_v) in range_rules.items():
                if field_name in table.headers:
                    value = record.values.get(field_name, "")
                    if value.strip() and _is_number(value):
                        num = float(value)
                        if min_v is not None and num < min_v:
                            add("数值超范围", record, field_name, value, f"字段 {field_name} 小于最小值 {min_v}", "调整到允许范围内")
                        if max_v is not None and num > max_v:
                            add("数值超范围", record, field_name, value, f"字段 {field_name} 大于最大值 {max_v}", "调整到允许范围内")

        if options.check_boolean_fields:
            for field_name in boolean_fields:
                if field_name in table.headers:
                    value = record.values.get(field_name, "")
                    if value.strip() and value.strip() not in BOOL_VALUES:
                        add("布尔值非法", record, field_name, value, f"字段 {field_name} 不是合法布尔值", "使用 0/1/true/false")

        if options.check_enum_fields:
            for field_name, allowed in enum_rules.items():
                if field_name in table.headers:
                    value = record.values.get(field_name, "")
                    if value.strip() and value.strip() not in allowed:
                        add("枚举值非法", record, field_name, value, f"字段 {field_name} 不在允许值 {sorted(allowed)} 内", "改为允许枚举值")

        if options.check_time_fields:
            for field_name in time_fields:
                if field_name in table.headers:
                    value = record.values.get(field_name, "")
                    if not _is_valid_time(value):
                        add("时间格式错误", record, field_name, value, f"字段 {field_name} 时间格式错误", "使用 yyyy-MM-dd HH:mm:ss、yyyy/MM/dd HH:mm:ss、yyyy-MM-dd 或时间戳")

        if options.check_time_range:
            for start_field, end_field in time_range_rules:
                if start_field in table.headers and end_field in table.headers:
                    start_text = record.values.get(start_field, "")
                    end_text = record.values.get(end_field, "")
                    start_dt = _parse_time(start_text)
                    end_dt = _parse_time(end_text)
                    if start_text.strip() and end_text.strip() and start_dt and end_dt and start_dt > end_dt:
                        add("时间范围错误", record, f"{start_field}/{end_field}", f"{start_text} > {end_text}", "开始时间晚于结束时间", "调整开始或结束时间")

        if options.check_reward_fields:
            for field_name in reward_fields:
                if field_name in table.headers:
                    value = record.values.get(field_name, "")
                    err = _reward_error(value)
                    if err:
                        add("奖励格式错误", record, field_name, value, err, "检查奖励格式和数量")

        if options.check_reference_ids:
            for field_name, ref_ids in ref_sets.items():
                if field_name in table.headers:
                    value = record.values.get(field_name, "").strip()
                    if value and value not in ref_ids:
                        add("引用 ID 不存在", record, field_name, value, f"引用 ID 在参考表中不存在：{value}", "检查引用 ID 或参考表")

        if options.check_interval:
            for min_field, max_field in interval_rules:
                if min_field in table.headers and max_field in table.headers:
                    min_text = record.values.get(min_field, "")
                    max_text = record.values.get(max_field, "")
                    if min_text.strip() and max_text.strip() and _is_number(min_text) and _is_number(max_text):
                        if float(min_text) > float(max_text):
                            add("区间配置错误", record, f"{min_field}/{max_field}", f"{min_text} > {max_text}", "min 字段大于 max 字段", "调整区间上下限")

    return issues, table


def validation_issues_to_tsv(issues: Sequence[ValidationIssue]) -> str:
    headers = ["序号", "行号", "主键 ID", "字段名", "字段类型", "当前值", "问题类型", "问题说明", "建议处理", "文件名", "Sheet"]
    rows = ["\t".join(headers)]
    for item in issues:
        rows.append("\t".join([
            str(item.index), item.row_number, item.item_id, item.field_name, item.field_type,
            item.current_value, item.issue_type, item.remark, item.suggestion, item.file_name, item.sheet,
        ]))
    return "\n".join(rows)


def export_validation_issues_to_excel(issues: Sequence[ValidationIssue], output_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "配置表校验结果"
    headers = ["序号", "行号", "主键 ID", "字段名", "字段类型", "当前值", "问题类型", "问题说明", "建议处理", "文件名", "Sheet"]
    ws.append(headers)
    for item in issues:
        ws.append([
            item.index, item.row_number, item.item_id, item.field_name, item.field_type,
            item.current_value, item.issue_type, item.remark, item.suggestion, item.file_name, item.sheet,
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [8, 10, 18, 22, 14, 30, 18, 52, 42, 24, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(output_path)


def options_to_dict(options: ConfigValidationOptions) -> Dict[str, Any]:
    return dict(options.__dict__)


def options_from_dict(data: Dict[str, Any]) -> ConfigValidationOptions:
    valid = {field.name for field in ConfigValidationOptions.__dataclass_fields__.values()}
    return ConfigValidationOptions(**{k: v for k, v in data.items() if k in valid})
