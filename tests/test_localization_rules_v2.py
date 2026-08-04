from __future__ import annotations

import collections
import tempfile
import threading
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.localization_checker import (
    LocalizationColumn,
    LocalizationCheckOptions,
    check_localization_file,
    export_localization_issues_to_excel,
    recommend_target_languages_from_filename,
)


class LocalizationRulesV2Tests(unittest.TestCase):
    def _run(self, rows, *, aggregate=True):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / "rules.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Text"
        sheet.append(["#tid", "CN", "EN", "DE"])
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        issues, _table, _columns = check_localization_file(
            str(path),
            LocalizationCheckOptions(
                table_structure="单 Sheet 多语言列",
                source_language_column="CN",
                enabled_language_columns=["EN", "DE"],
                check_chinese=False,
                check_unfinished=False,
                check_spelling=False,
                check_terms=False,
                check_duplicate_translation=False,
                aggregate_empty_language_columns=aggregate,
                ui_limit=100000,
            ),
            threading.Event(),
        )
        return path, issues

    def test_numeric_equivalence_range_order_percent_and_real_change(self):
        _path, issues = self._run([
            ["level", "城堡3级", "Castle St. 3", "Schloss Niv. 3"],
            ["range", "开放1-30关", "Stages 1-30", "Stufen 1–30"],
            ["top", "地区排名前3", "Top-3 in the region", "Top-3 der Region"],
            ["order", "持续8小时提升50%", "Increase 50 percent for 8 hours", "8 Stunden, 50 Prozent"],
            ["decimal", "减少10.5%", "Reduce by 10.5%", "Um 10,5% reduzieren"],
            ["changed", "目标2", "Goal 3", "Ziel 2"],
        ], aggregate=False)
        by_id = collections.defaultdict(list)
        for item in issues:
            by_id[item.item_id].append(item)
        for text_id in ("level", "range", "top", "order", "decimal"):
            self.assertFalse(any(item.rule_name == "数值" for item in by_id[text_id]), text_id)
        self.assertIn("数值内容不一致", [item.issue_type for item in by_id["changed"]])
        self.assertFalse(any(item.rule_name == "数值" and item.language.startswith("DE") for item in by_id["changed"]))

    def test_date_cells_and_serial_text_are_not_clear_numeric_errors(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / "dates.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Text"
        sheet.append(["#tid", "CN", "EN", "DE"])
        sheet.append(["date", "三月二十九日", datetime(2026, 3, 29), "46110"])
        sheet["C2"].number_format = "mmmm d"
        workbook.save(path)
        issues, _table, _columns = check_localization_file(
            str(path),
            LocalizationCheckOptions(
                table_structure="单 Sheet 多语言列",
                source_language_column="CN",
                enabled_language_columns=["EN", "DE"],
                check_chinese=False,
                check_unfinished=False,
                check_spelling=False,
                check_terms=False,
                check_duplicate_translation=False,
                aggregate_empty_language_columns=False,
            ),
            threading.Event(),
        )
        self.assertFalse(any(item.language.startswith("EN") and item.rule_name == "数值" for item in issues))
        de_numbers = [item for item in issues if item.language.startswith("DE") and item.rule_name == "数值"]
        self.assertEqual(1, len(de_numbers))
        self.assertEqual("需人工确认", de_numbers[0].result_category)
        self.assertFalse(de_numbers[0].count_in_error_stats)

    def test_context_symbols_tags_source_dedupe_and_coverage(self):
        path, issues = self._run([
            ["bracket", "活动名称", "[I Scream]", "[Eis]"],
            ["slash", "他/她", "he or she", "er oder sie"],
            ["percent", "完成50%", "Complete 50 percent", "Zu 50 Prozent abschließen"],
            ["tag_attr", "[color=#f53d39]文本[/color]", "[color=#CFBFA3]Text[/color]", "[color=#f53d39]Text[/color]"],
            ["target_bad", "[b]文本[/b]", "[color=#ffff", "[b]Text[/b]"],
            ["source_bad", "[color=#d16316]源文本", "Text", "Text"],
            ["partial_empty", "中文", "", "Deutsch"],
        ])
        by_id = collections.defaultdict(list)
        for item in issues:
            by_id[item.item_id].append(item)
        self.assertFalse(any("特殊符号" in item.issue_type or "标签" in item.issue_type for item in by_id["bracket"]))
        self.assertFalse(any("特殊符号" in item.issue_type for item in by_id["slash"]))
        self.assertFalse(any("特殊符号" in item.issue_type or item.rule_name == "数值" for item in by_id["percent"]))
        attr = [item for item in by_id["tag_attr"] if item.issue_type == "标签属性差异"]
        self.assertEqual(1, len(attr))
        self.assertEqual("需人工确认", attr[0].result_category)
        self.assertIn("目标文本标签结构异常", [item.issue_type for item in by_id["target_bad"]])
        source_problems = [item for item in by_id["source_bad"] if item.result_category == "源文本问题"]
        self.assertEqual(1, len(source_problems))
        self.assertFalse(any(item.language.startswith("EN") or item.language.startswith("DE") for item in source_problems))
        self.assertIn("空翻译", [item.issue_type for item in by_id["partial_empty"]])

        coverage = [item for item in issues if item.result_category == "语言覆盖范围"]
        self.assertEqual(2, len(coverage))
        self.assertTrue(all(not item.count_in_error_stats for item in coverage))

        report_path = path.with_name("report.xlsx")
        export_localization_issues_to_excel(issues, str(report_path))
        report = load_workbook(report_path, read_only=False, data_only=True)
        try:
            self.assertEqual(
                ["汇总说明", "明确问题", "源文本问题", "需人工确认", "译文复用分析", "语言覆盖范围", "忽略项或白名单命中", "全部明细", "规则说明"],
                report.sheetnames,
            )
            for sheet_name in report.sheetnames:
                self.assertTrue(report[sheet_name].auto_filter.ref)
        finally:
            report.close()

    def test_filename_language_recommendation_uses_standalone_codes(self):
        columns = [LocalizationColumn(name=code, header=code, code=code, index=index) for index, code in enumerate(["CN", "EN", "ID", "IT", "TR", "DE"])]
        self.assertEqual(
            ["IT", "TR", "ID"],
            recommend_target_languages_from_filename("TP-story-IT_TR_ID.xlsx", columns),
        )
        self.assertEqual([], recommend_target_languages_from_filename("story_ID.xlsx", columns))

    def test_embedded_headers_todo_boundaries_control_rows_and_invalid_id(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / "TP-story-IT_TR_ID.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Text"
        sheet.append(["ID", "类型", "人物", "备注", "CN", "EN Final", "Indonesian", "Italian", "Turkish"])
        sheet.append(["ID12", "事件类型", "类型子项", "动作", "对白文字", "", "", "", ""])
        sheet.append(["control", "转场动画", "", "", "", "", "", "", ""])
        sheet.append(["source_missing", "旁白", "", "", "", "Reference", "Terjemahan", "Traduzione", "Çeviri"])
        sheet.append([None, "旁白", "", "正式内容", "中文", "English", "Indonesia", "Italiano", "Türkçe"])
        sheet.append(["empty", "旁白", "", "", "中文", "English", "", "", ""])
        sheet.append(["natural", "旁白", "", "", "方法", "Method", "todo", "metodo", "todos"])
        sheet.append(["marker", "旁白", "", "", "待处理", "TODO", "TODO", "Traduzione", "Çeviri"])
        workbook.save(path)

        issues, table, _columns = check_localization_file(
            str(path),
            LocalizationCheckOptions(
                table_structure="单 Sheet 多语言列",
                source_language_column="CN",
                enabled_language_columns=["ID", "IT", "TR"],
                check_chinese=False,
                check_spelling=False,
                check_terms=False,
                check_duplicate_translation=False,
                aggregate_empty_language_columns=False,
            ),
            threading.Event(),
        )
        self.assertNotIn("ID12", table.records)
        by_id = collections.defaultdict(list)
        for item in issues:
            by_id[item.item_id].append(item)
        self.assertFalse(by_id["control"])
        self.assertEqual(1, len([item for item in by_id["source_missing"] if item.result_category == "源文本问题"]))
        self.assertEqual(1, len([item for item in issues if item.issue_type == "无效 ID"]))
        self.assertEqual(3, len([item for item in by_id["empty"] if item.issue_type == "空翻译"]))
        self.assertFalse(any(item.issue_type == "未完成/修改中占位文本" for item in by_id["natural"]))
        self.assertEqual(1, len([item for item in by_id["marker"] if item.issue_type == "未完成/修改中占位文本"]))

    def test_short_date_number_words_months_reversed_newline_and_e_tag(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / "rules-IT_TR_ID.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Text"
        sheet.append(["ID", "类型", "CN", "Indonesian", "Italian", "Turkish"])
        sheet.append(["date", "标题", datetime(2022, 1, 1), "Bintang Jatuh", "Stella caduta", "Düşen Yıldız"])
        sheet["C2"].number_format = "m\\-d"
        sheet.append(["days", "旁白", "五天之后", "lima hari kemudian", "cinque giorni dopo", "Beş gün sonra"])
        sheet.append(["month", "标题", "7月份", "Juli", "Luglio", "Temmuz"])
        sheet.append(["idiom", "旁白", "那只是零星的希望", "Harapan yang samar", "Una speranza vaga", "Belirsiz bir umut"])
        sheet.append(["newline", "旁白", "n/奥兹大王将会到来", "\\nOz akan datang", "\\nOz arriverà", "\\nOz gelecek"])
        sheet.append(["tag_good", "旁白", "[e]文本[/e]", "[e]Teks[/e]", "[e]Testo[/e]", "[e]Metin[/e]"])
        sheet.append(["tag_source_bad", "旁白", "[e]文本[e]", "Text", "Testo", "Metin"])
        sheet.append(["tag_target_bad", "旁白", "[e]文本[/e]", "[e]Teks[/e]", "[e]Testo[e]", "[e]Metin[/e]"])
        sheet.append(["bracket", "旁白", "活动名称", "[Masa Intim]", "[Momenti intimi]", "[Samimi Anlar]"])
        workbook.save(path)

        issues, _table, _columns = check_localization_file(
            str(path),
            LocalizationCheckOptions(
                table_structure="单 Sheet 多语言列",
                source_language_column="CN",
                enabled_language_columns=["ID", "IT", "TR"],
                check_chinese=False,
                check_spelling=False,
                check_terms=False,
                check_duplicate_translation=False,
                aggregate_empty_language_columns=False,
            ),
            threading.Event(),
        )
        by_id = collections.defaultdict(list)
        for item in issues:
            by_id[item.item_id].append(item)
        date_issues = [item for item in by_id["date"] if item.issue_type == "源文本单元格格式异常"]
        self.assertEqual(1, len(date_issues))
        self.assertFalse(any(item.rule_name == "数值" for item in by_id["date"]))
        for text_id in ("days", "month", "idiom"):
            self.assertFalse(any(item.rule_name == "数值" for item in by_id[text_id]), text_id)
        self.assertEqual(1, len([item for item in by_id["newline"] if item.issue_type == "源文本换行格式符疑似写反"]))
        self.assertFalse(any(item.issue_type == "换行符数量差异" for item in by_id["newline"]))
        self.assertFalse(any("标签" in item.issue_type for item in by_id["tag_good"]))
        self.assertEqual(1, len([item for item in by_id["tag_source_bad"] if item.result_category == "源文本问题"]))
        self.assertIn("目标文本标签结构异常", [item.issue_type for item in by_id["tag_target_bad"]])
        self.assertFalse(any("标签" in item.issue_type for item in by_id["bracket"]))

    def test_duplicate_translation_is_separate_optional_analysis(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / "reuse.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ID", "CN", "Italian"])
        sheet.append(["a", "第一段完全不同的中文源文", "Questa è la stessa traduzione molto lunga"])
        sheet.append(["b", "另一段含义不同的中文内容", "Questa è la stessa traduzione molto lunga"])
        workbook.save(path)
        base = dict(
            table_structure="单 Sheet 多语言列",
            source_language_column="CN",
            enabled_language_columns=["IT"],
            check_chinese=False,
            check_spelling=False,
            check_terms=False,
            aggregate_empty_language_columns=False,
        )
        issues_off, _table, _columns = check_localization_file(str(path), LocalizationCheckOptions(**base), threading.Event())
        self.assertFalse(any(item.rule_name == "译文重复" for item in issues_off))
        issues_on, _table, _columns = check_localization_file(
            str(path), LocalizationCheckOptions(**base, check_duplicate_translation=True), threading.Event()
        )
        reuse = [item for item in issues_on if item.rule_name == "译文重复"]
        self.assertTrue(reuse)
        self.assertTrue(all(item.result_category == "译文复用分析" and not item.count_in_error_stats for item in reuse))


if __name__ == "__main__":
    unittest.main()
