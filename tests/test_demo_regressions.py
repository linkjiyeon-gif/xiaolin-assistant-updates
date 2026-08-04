from __future__ import annotations

import collections
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.config_validator import (  # noqa: E402
    ConfigValidationOptions,
    detect_config_structure,
    validate_config_table,
)
from app.localization_checker import (  # noqa: E402
    LocalizationCheckOptions,
    check_localization_file,
    export_localization_issues_to_excel,
)
from app.value_config_compare import (  # noqa: E402
    ValueCompareOptions,
    compare_value_reference_to_config,
)


DEMO_ROOT_TEXT = os.environ.get("QA_TOOLBOX_DEMO_ROOT", "").strip()
DEMO_ROOT = Path(DEMO_ROOT_TEXT) if DEMO_ROOT_TEXT else None


def _require_demo_root() -> Path:
    if DEMO_ROOT is None or not DEMO_ROOT.exists():
        raise unittest.SkipTest("请设置 QA_TOOLBOX_DEMO_ROOT 指向 QA测试工具盒_分享会演示素材")
    return DEMO_ROOT


def _report_rows(path: Path, sheet_index: int = 0):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.worksheets[sheet_index].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()


class DemoRegressionTests(unittest.TestCase):
    def test_commercial_compare_matches_verified_reports(self):
        root = _require_demo_root() / "商业化配置比对范例"
        mall = root / "mall.csv"
        cases = [
            ("常规礼包.xlsx", "常规礼包差异比对结果.xlsx"),
            ("月卡周卡.xlsx", "月卡周卡差异比对结果.xlsx"),
        ]
        for source_name, report_name in cases:
            with self.subTest(source=source_name):
                diffs, _report = compare_value_reference_to_config(
                    str(mall),
                    str(root / source_name),
                    ValueCompareOptions(template_name="mall", enable_general_compare=True),
                )
                expected = {
                    (str(row[1] or ""), str(row[2] or ""), str(row[3] or ""), str(row[4] or ""), str(row[5] or ""))
                    for row in _report_rows(root / report_name)
                }
                actual = {
                    (item.diff_type, item.document_content, item.file_content, item.location, item.remark)
                    for item in diffs
                }
                self.assertEqual(expected, actual)

    def test_config_validation_matches_verified_reports(self):
        root = _require_demo_root() / "配置表校验范例"
        cases = [
            ("shop_basic.csv", "shop_basic配置表校验报告.xlsx", True),
            ("chessboard_stage.csv", "chessboard_stage配置表校验报告.xlsx", False),
        ]
        compound_rule = "field=restrict;allow_empty=true;group=none;sep=|;count=2;type=uint;allow_zero=true;forbid=0"
        for source_name, report_name, use_compound in cases:
            with self.subTest(source=source_name):
                source = root / source_name
                structure = detect_config_structure(str(source))
                options = ConfigValidationOptions(
                    header_row=structure["header_row"],
                    data_start_row=structure["data_start_row"],
                    key_column=structure["key_column"],
                    sheet_name=structure.get("sheet_name", ""),
                    description_row=structure.get("description_row", 0),
                    type_row=structure.get("type_row", 0),
                    check_type_fields=True,
                    check_compound_rules=use_compound,
                    compound_rules_text=compound_rule if use_compound else "",
                    ui_limit=100000,
                )
                issues, _table = validate_config_table(str(source), options)
                expected = {
                    tuple(str(value or "") for value in row[1:9])
                    for row in _report_rows(root / report_name)
                }
                actual = {
                    (
                        str(item.row_number or ""),
                        str(item.item_id or ""),
                        str(item.field_name or ""),
                        str(item.field_type or ""),
                        str(item.current_value or ""),
                        str(item.issue_type or ""),
                        str(item.remark or ""),
                        str(item.suggestion or ""),
                    )
                    for item in issues
                }
                self.assertEqual(expected, actual)

    def test_localization_checks_match_verified_reports(self):
        root = _require_demo_root() / "多语言检查范例"
        cases = [
            ("单Sheet语言.xlsx", "单Sheet语言文本质量检查报告.xlsx"),
            ("多Sheet语言.xlsx", "多Sheet语言文本质量检查报告.xlsx"),
        ]
        for source_name, report_name in cases:
            with self.subTest(source=source_name):
                issues, _table, _columns = check_localization_file(
                    str(root / source_name),
                    LocalizationCheckOptions(ui_limit=100000, aggregate_empty_language_columns=False),
                    threading.Event(),
                )
                expected_rows = _report_rows(root / report_name, sheet_index=1)
                def normalize_historical_issue(values):
                    values = list(values)
                    if values[4] == "数值数量不一致":
                        # The rule still detects the same historical rows, but
                        # the new policy deliberately downgrades uncertain
                        # localization conversions to manual confirmation.
                        values[5] = "数值复核"
                        values[6] = "数值数量差异"
                    return tuple(values)

                expected = {
                    normalize_historical_issue((
                        str(row[1] or ""),
                        str(row[2] or ""),
                        str(row[3] or ""),
                        str(row[5] or ""),
                        str(row[8] or ""),
                        str(row[9] or ""),
                        str(row[11] or ""),
                    ))
                    for row in expected_rows
                    if str(row[8] or "") != "数值顺序变化"
                    # Confirmed historical false positives: lowercase natural-language
                    # words such as Spanish/Portuguese ``todo`` are not TODO markers.
                    and not (
                        str(row[8] or "") == "未完成/修改中占位文本"
                        and str(row[11] or "").endswith("：todo")
                    )
                    # Translation reuse is now an opt-in low-confidence analysis and
                    # therefore is intentionally absent from the default result set.
                    and str(row[8] or "") != "不同源文本译文重复"
                }
                actual = {
                    normalize_historical_issue((
                        str(item.sheet_name or ""),
                        str(item.row_number or ""),
                        str(item.item_id or ""),
                        str(item.language or ""),
                        str(item.issue_type or ""),
                        str(item.issue_level or ""),
                        str(item.remark or ""),
                    ))
                    for item in issues
                    if item.issue_type != "换行符数量差异" and item.issue_level != "源文本问题"
                }
                self.assertEqual(expected, actual)
                self.assertEqual(
                    collections.Counter(
                        row[8]
                        for row in expected_rows
                        if str(row[8] or "") != "数值顺序变化"
                        and not (
                            str(row[8] or "") == "未完成/修改中占位文本"
                            and str(row[11] or "").endswith("：todo")
                        )
                        and str(row[8] or "") != "不同源文本译文重复"
                    ),
                    collections.Counter(
                        item.issue_type
                        for item in issues
                        if item.issue_type != "换行符数量差异" and item.issue_level != "源文本问题"
                    ),
                )

    def test_localization_empty_language_column_is_aggregated_by_default(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "empty_language_column.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Text"
            sheet.append(["#tid", "CN", "EN", "DE"])
            for row_index in range(1, 6):
                sheet.append([f"id_{row_index}", f"中文{row_index}", "", f"DE {row_index}"])
            workbook.save(path)

            issues, _table, _columns = check_localization_file(
                str(path),
                LocalizationCheckOptions(
                    table_structure="单 Sheet 多语言列",
                    source_language_column="CN",
                    enabled_language_columns=["EN", "DE"],
                    check_empty=True,
                    check_chinese=False,
                    check_placeholders=False,
                    check_tags=False,
                    check_symbols=False,
                    check_numbers=False,
                    check_unfinished=False,
                    check_spelling=False,
                    check_terms=False,
                    check_duplicate_translation=False,
                ),
                threading.Event(),
            )

        coverage = [item for item in issues if item.result_category == "语言覆盖范围"]
        self.assertEqual(2, len(coverage))
        empty_column = next(item for item in coverage if item.language == "EN")
        self.assertEqual("语言覆盖状态", empty_column.issue_type)
        self.assertEqual("整列", empty_column.row_number)
        self.assertEqual("整列空翻译", empty_column.rule_name)
        self.assertFalse(empty_column.count_in_error_stats)
        self.assertEqual(5, empty_column.coverage_blank)

    def test_localization_dictionary_files_support_switchable_source_language(self):
        def make_dictionary(path: Path, rows: list[tuple[str, str, str]]) -> None:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = path.stem
            sheet.append(["", "ID", "", "Contents"])
            sheet.append(["", "string", "", "string"])
            sheet.append(["注释1", "#唯一功能缩写+日期+2位编号", "所属模块", ""])
            sheet.append(["注释2", "0", "3", "0"])
            for text_id, module, contents in rows:
                sheet.append(["", text_id, module, contents])
            workbook.save(path)

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            make_dictionary(
                root / "dictionary_ChineseSimplified.xlsx",
                [
                    ("ID_001", "UI", "开始游戏"),
                    ("ID_002", "UI", ""),
                    ("ID_003", "Battle", "数量[%s1]"),
                    ("ID_004", "Battle", "播放"),
                ],
            )
            make_dictionary(
                root / "dictionary_English.xlsx",
                [
                    ("ID_001", "UI", ""),
                    ("ID_002", "UI", ""),
                    ("ID_003", "Battle", "Count[%s1]"),
                    ("ID_004", "Battle", "Play"),
                ],
            )
            make_dictionary(
                root / "dictionary_French.xlsx",
                [
                    ("ID_001", "UI", "Démarrer"),
                    ("ID_002", "UI", ""),
                    ("ID_003", "Battle", "Nombre[%s1]"),
                    ("ID_004", "Battle", ""),
                ],
            )

            cn_source_issues, _table, columns = check_localization_file(
                str(root),
                LocalizationCheckOptions(
                    table_structure="多文件字典模式",
                    source_language_column="CN",
                    enabled_language_columns=["EN", "FR"],
                    check_empty=True,
                    check_chinese=False,
                    check_placeholders=False,
                    check_tags=False,
                    check_symbols=False,
                    check_numbers=False,
                    check_unfinished=False,
                    check_spelling=False,
                    check_terms=False,
                    check_duplicate_translation=False,
                    aggregate_empty_language_columns=False,
                ),
                threading.Event(),
            )
            empty_by_language_and_id = {(item.language, item.item_id): item for item in cn_source_issues if item.issue_type == "空翻译"}

            self.assertEqual(["CN", "EN", "FR"], [column.header for column in columns])
            self.assertIn(("EN", "ID_001"), empty_by_language_and_id)
            self.assertNotIn(("EN", "ID_002"), empty_by_language_and_id)
            self.assertIn(("FR", "ID_004"), empty_by_language_and_id)
            self.assertEqual("UI", empty_by_language_and_id[("EN", "ID_001")].module)
            self.assertEqual("dictionary_ChineseSimplified.xlsx", empty_by_language_and_id[("EN", "ID_001")].source_file)
            self.assertEqual("dictionary_English.xlsx", empty_by_language_and_id[("EN", "ID_001")].target_file)

            en_source_issues, _table, columns = check_localization_file(
                str(root / "dictionary_English.xlsx"),
                LocalizationCheckOptions(
                    table_structure="多文件字典模式",
                    source_language_column="EN",
                    enabled_language_columns=["CN", "FR"],
                    check_empty=True,
                    check_chinese=False,
                    check_placeholders=False,
                    check_tags=False,
                    check_symbols=False,
                    check_numbers=False,
                    check_unfinished=False,
                    check_spelling=False,
                    check_terms=False,
                    check_duplicate_translation=False,
                    aggregate_empty_language_columns=False,
                ),
                threading.Event(),
            )
            empty_by_language_and_id = {(item.language, item.item_id): item for item in en_source_issues if item.issue_type == "空翻译"}

            self.assertEqual(["EN", "CN", "FR"], [column.header for column in columns])
            self.assertIn(("FR", "ID_004"), empty_by_language_and_id)
            self.assertNotIn(("CN", "ID_001"), empty_by_language_and_id)
            self.assertEqual("dictionary_English.xlsx", empty_by_language_and_id[("FR", "ID_004")].source_file)
            self.assertEqual("dictionary_French.xlsx", empty_by_language_and_id[("FR", "ID_004")].target_file)

    def test_localization_numeric_id_does_not_export_trailing_decimal(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "numeric_id.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Text"
            sheet.append(["#tid", "CN", "EN"])
            sheet.append([103659.0, "中文", ""])
            workbook.save(path)

            issues, _table, _columns = check_localization_file(
                str(path),
                LocalizationCheckOptions(
                    table_structure="单 Sheet 多语言列",
                    source_language_column="CN",
                    enabled_language_columns=["EN"],
                    check_empty=True,
                    check_chinese=False,
                    check_placeholders=False,
                    check_tags=False,
                    check_symbols=False,
                    check_numbers=False,
                    check_unfinished=False,
                    check_spelling=False,
                    check_terms=False,
                    check_duplicate_translation=False,
                    aggregate_empty_language_columns=False,
                ),
                threading.Event(),
            )
            empty_issue = next(item for item in issues if item.issue_type == "空翻译")
            self.assertEqual("103659", empty_issue.item_id)

            report_path = Path(temp_dir) / "report.xlsx"
            export_localization_issues_to_excel(issues, str(report_path))
            report = load_workbook(report_path, read_only=True, data_only=True)
            try:
                detail = report["全部明细"]
                exported_ids = [str(row[4]) for row in detail.iter_rows(min_row=2, values_only=True) if row[10] == "空翻译"]
                self.assertEqual(["103659"], exported_ids)
            finally:
                report.close()

    def test_localization_quality_rule_boundaries_and_source_classification(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "quality_boundaries.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Text"
            sheet.append(["#tid", "CN", "EN"])
            sheet.append(["language_name", "简体中文", "简体中文"])
            sheet.append(["percent_word", "每[%s1]秒恢复", "Restore every [%s1] seconds ]%for"])
            sheet.append(["newline_format", "第一行\\n第二行", "Line one and line two"])
            sheet.append(["bracket_prose", "冰淇淋", "[I Scream]"])
            sheet.append(["source_bad_tag", "文字[/b]", "Text"])
            sheet.append(["target_missing_tag", "[b]重点[/b]", "Important"])
            sheet.append(["source_empty", "\u200b", "Contact Us"])
            sheet.append(["hash_missing", "完成 #NAME 任务", "Complete the task"])
            sheet.append(["hash_same", "完成 #NAME 任务", "Complete the #NAME task"])
            sheet.append(["ordinal_equal", "第二章节", "Chapter 2"])
            sheet.append(["number_changed", "目标2", "Goal 3"])
            sheet.append(["percent_m_added", "信件内容", "Letter content %m"])
            sheet.append(["bracket_with_placeholder", "内容：[活动] 时间：%s", "Event time: %s"])
            workbook.save(path)

            issues, _table, _columns = check_localization_file(
                str(path),
                LocalizationCheckOptions(
                    table_structure="单 Sheet 多语言列",
                    source_language_column="CN",
                    enabled_language_columns=["EN"],
                    check_empty=True,
                    check_chinese=True,
                    check_placeholders=True,
                    check_tags=True,
                    check_symbols=True,
                    check_numbers=True,
                    check_unfinished=False,
                    check_spelling=False,
                    check_terms=False,
                    check_duplicate_translation=False,
                    aggregate_empty_language_columns=False,
                ),
                threading.Event(),
            )

        by_id = collections.defaultdict(list)
        for issue in issues:
            by_id[issue.item_id].append(issue)

        self.assertFalse(any(item.issue_type == "残留中文" for item in by_id["language_name"]))
        self.assertFalse(any(item.issue_type.startswith("占位符") for item in by_id["percent_word"]))
        self.assertEqual(["换行符数量差异"], [item.issue_type for item in by_id["newline_format"]])
        self.assertFalse(any("标签" in item.issue_type for item in by_id["bracket_prose"]))
        self.assertEqual(["CN 源文本标签结构异常"], [item.issue_type for item in by_id["source_bad_tag"]])
        self.assertIn("标签缺失", [item.issue_type for item in by_id["target_missing_tag"]])
        self.assertEqual(["CN源文本为空"], [item.issue_type for item in by_id["source_empty"]])
        self.assertIn("占位符缺失", [item.issue_type for item in by_id["hash_missing"]])
        self.assertFalse(any(item.issue_type.startswith("占位符") for item in by_id["hash_same"]))
        self.assertFalse(any(item.issue_type.startswith("数值") for item in by_id["ordinal_equal"]))
        self.assertIn("数值内容不一致", [item.issue_type for item in by_id["number_changed"]])
        self.assertIn("占位符新增", [item.issue_type for item in by_id["percent_m_added"]])
        self.assertFalse(any("特殊符号" in item.issue_type for item in by_id["bracket_with_placeholder"]))


if __name__ == "__main__":
    unittest.main()
